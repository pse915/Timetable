import io
import base64
import json
import math
import os
import html as html_lib
from contextlib import contextmanager
from functools import lru_cache
import uuid
import threading
import hashlib
import secrets as py_secrets
import time
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo
from collections import defaultdict
import pandas as pd
import streamlit as st
import requests
import re
import gspread
from gspread.exceptions import WorksheetNotFound
from google.oauth2.service_account import Credentials
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
st.set_page_config(page_title="시간표·결보강 관리", page_icon="📘", layout="wide", initial_sidebar_state="collapsed")
st.markdown("""
<style>
:root{--ui-bg:#ffffff;--ui-surface:#ffffff;--ui-soft:#f5f5f7;--ui-soft-2:#fafafc;--ui-text:#1d1d1f;--ui-text-2:#3a3a3c;--ui-muted:#6e6e73;--ui-muted-2:#86868b;--ui-line:#d2d2d7;--ui-line-soft:#e5e5ea;--ui-accent:#0066cc;--ui-accent-hover:#0071e3;--ui-focus:rgba(0,102,204,.20);--ui-success:#1b7f3a;--ui-danger:#c62828;--ui-radius-sm:8px;--ui-radius-md:12px;--ui-radius-lg:16px;--ui-pill:999px;}html,body,[data-testid="stAppViewContainer"],[data-testid="stApp"],[data-testid="stMain"]{background:var(--ui-bg)!important;color:var(--ui-text)!important}body{-webkit-font-smoothing:antialiased;text-rendering:optimizeLegibility}html,body,[data-testid="stAppViewContainer"],[data-testid="stAppViewContainer"] *{font-family:var(--app-font,"SF Pro Text","SF Pro Display",-apple-system,BlinkMacSystemFont,"Segoe UI","Apple SD Gothic Neo","Noto Sans KR",sans-serif)!important}body,body *{font-synthesis:auto}table,thead,tbody,tfoot,tr,th,td,caption{font-family:var(--app-font,"SF Pro Text",system-ui,sans-serif)!important}[data-baseweb] *,[data-testid="stPopover"] *,[data-testid="stDialog"] *,[data-testid="stExpander"] *,[data-testid="stSidebar"] *,[role="menu"] *,[role="option"] *,[role="listbox"] *,[role="tooltip"] *{font-family:var(--app-font,"SF Pro Text",system-ui,sans-serif)!important}[data-testid="stDataFrame"],[data-testid="stDataEditor"],[data-testid="stDataFrame"] *,[data-testid="stDataEditor"] *{--gdg-font-family:var(--app-font,"SF Pro Text",system-ui,sans-serif)!important;}[data-testid="stDataFrame"],[data-testid="stDataEditor"]{font-family:var(--app-font,"SF Pro Text",system-ui,sans-serif)!important;}[data-testid="stDataFrame"] [role="gridcell"],[data-testid="stDataFrame"] [role="columnheader"],[data-testid="stDataEditor"] [role="gridcell"],[data-testid="stDataEditor"] [role="columnheader"]{font-family:var(--app-font,"SF Pro Text",system-ui,sans-serif)!important}[data-testid="stAppViewContainer"] svg text,.js-plotly-plot text,.plotly .legendtext,.plotly .gtitle,.plotly .xtick text,.plotly .ytick text{font-family:var(--app-font,"SF Pro Text",system-ui,sans-serif)!important}[class*="material-symbols"],[data-testid="stIconMaterial"],[class*="MaterialSymbols"],[data-testid="stIconMaterial"] *,[aria-label="More options"] span{font-family:"Material Symbols Rounded","Material Symbols Outlined",sans-serif!important;font-style:normal!important}[data-testid="stHeader"]{background:rgba(255,255,255,.88)!important;border-bottom:1px solid var(--ui-line-soft)!important;backdrop-filter:saturate(150%) blur(16px)}[data-testid="stDecoration"]{display:none!important}.block-container{width:100%!important;max-width:none!important;padding:0 clamp(16px,2.2vw,42px) 36px!important}.app-top-safe-space{height:30px;width:100%}.app-topbar{position:relative;z-index:3;display:flex;align-items:center;gap:14px;min-height:54px;padding:4px 0 8px;margin:0 0 14px;border-bottom:1px solid var(--ui-line-soft)}.app-identity{white-space:nowrap;color:var(--ui-muted);font-size:12px;line-height:1.25}.app-identity strong{color:var(--ui-text);font-weight:600}.app-topbar .stButton>button{min-height:34px!important;padding:6px 13px!important;border-radius:var(--ui-pill)!important;font-size:12px!important}.app-topbar [data-testid="stRadio"] [role="radiogroup"]{display:flex!important;align-items:center!important;gap:2px!important;width:100%!important;padding:2px!important;background:var(--ui-soft)!important;border:1px solid var(--ui-line-soft)!important;border-radius:var(--ui-pill)!important;overflow-x:auto!important;scrollbar-width:none}.app-topbar [data-testid="stRadio"] [role="radiogroup"]::-webkit-scrollbar{display:none}.app-topbar [data-testid="stRadio"] [role="radio"]{flex:0 0 auto!important;min-height:34px!important;padding:0 14px!important;border-radius:var(--ui-pill)!important;color:var(--ui-muted)!important;font-size:13px!important;font-weight:500!important;background:transparent!important;border:0!important}.app-topbar [data-testid="stRadio"] [role="radio"][aria-checked="true"]{background:var(--ui-surface)!important;color:var(--ui-text)!important;font-weight:600!important;box-shadow:0 1px 3px rgba(0,0,0,.08)!important}.app-topbar [data-testid="stRadio"] [role="radio"]>div:first-child{display:none!important}.app-topbar [data-testid="stSelectbox"]>div>div{min-height:34px!important;border-radius:var(--ui-pill)!important;background:var(--ui-soft)!important;border-color:var(--ui-line-soft)!important}.work-page-head{display:flex;align-items:flex-end;justify-content:space-between;gap:24px;padding:8px 2px 16px;margin:0 0 18px;border-bottom:1px solid var(--ui-line-soft)}.work-page-head h1{margin:0!important;font-size:30px!important;line-height:1.12!important;font-weight:600!important;letter-spacing:-.04em!important;color:var(--ui-text)!important}.work-page-head p{margin:6px 0 0!important;font-size:13px!important;line-height:1.45!important;color:var(--ui-muted)!important}.work-page-meta{font-size:12px;color:var(--ui-muted);white-space:nowrap;padding-bottom:3px}.work-section{margin:0 0 24px}.work-section-title{font-size:14px;font-weight:600;color:var(--ui-text);margin:0 0 8px 2px}.work-section-note{font-size:12px;color:var(--ui-muted);margin:-3px 0 10px 2px}[data-testid="stVerticalBlockBorderWrapper"]{background:var(--ui-surface)!important;border:1px solid var(--ui-line-soft)!important;border-radius:var(--ui-radius-lg)!important;box-shadow:none!important}[data-testid="stExpander"]{background:var(--ui-surface)!important;border:1px solid var(--ui-line-soft)!important;border-radius:var(--ui-radius-md)!important;box-shadow:none!important;overflow:hidden!important}[data-testid="stExpander"] summary{min-height:42px!important;padding:6px 12px!important;color:var(--ui-text)!important;font-size:13px!important;font-weight:600!important}[data-testid="stExpander"] summary:hover{background:var(--ui-soft)!important}h1,h2,h3,h4,h5,h6{color:var(--ui-text)!important}h2{font-size:22px!important;font-weight:600!important;letter-spacing:-.025em!important}h3{font-size:18px!important;font-weight:600!important;letter-spacing:-.02em!important}h4{font-size:15px!important;font-weight:600!important}[data-testid="stCaptionContainer"],.stCaption{color:var(--ui-muted)!important;font-size:12px!important;line-height:1.45!important}.stButton>button,.stDownloadButton>button,.stFormSubmitButton>button{min-height:40px!important;padding:7px 16px!important;border:1px solid var(--ui-line)!important;border-radius:var(--ui-pill)!important;background:var(--ui-surface)!important;color:var(--ui-text)!important;box-shadow:none!important;transform:none!important;transition:background .12s ease,border-color .12s ease!important;font-size:13px!important;font-weight:500!important}.stButton>button:hover,.stDownloadButton>button:hover,.stFormSubmitButton>button:hover{background:var(--ui-soft)!important;border-color:var(--ui-line)!important;box-shadow:none!important}.stButton>button[kind="primary"],.stFormSubmitButton>button[kind="primary"]{background:var(--ui-accent)!important;border-color:var(--ui-accent)!important;color:#fff!important;font-weight:600!important}.stButton>button[kind="primary"]:hover,.stFormSubmitButton>button[kind="primary"]:hover{background:var(--ui-accent-hover)!important;border-color:var(--ui-accent-hover)!important}.stButton>button:focus-visible,.stDownloadButton>button:focus-visible,.stFormSubmitButton>button:focus-visible{outline:3px solid var(--ui-focus)!important;outline-offset:1px}[data-baseweb="input"],[data-baseweb="textarea"],[data-baseweb="select"]>div,[data-testid="stDateInput"]>div>div{min-height:40px!important;background:var(--ui-surface)!important;color:var(--ui-text)!important;border:1px solid var(--ui-line)!important;border-radius:var(--ui-radius-sm)!important;box-shadow:none!important}[data-baseweb="input"] input,[data-baseweb="textarea"] textarea,[data-baseweb="select"] input{font-family:var(--app-font,"SF Pro Text",system-ui,sans-serif)!important;color:var(--ui-text)!important;background:transparent!important;font-size:13px!important}input::placeholder,textarea::placeholder{color:var(--ui-muted-2)!important}[data-baseweb="input"]:focus-within,[data-baseweb="textarea"]:focus-within,[data-baseweb="select"]:focus-within{border-color:var(--ui-accent)!important;box-shadow:0 0 0 3px var(--ui-focus)!important}[data-testid="stRadio"]:not(.app-topbar [data-testid="stRadio"]) [role="radiogroup"]{display:flex!important;gap:6px!important;padding:0!important;background:transparent!important;border:0!important;overflow:visible!important;flex-wrap:wrap}[data-testid="stRadio"]:not(.app-topbar [data-testid="stRadio"]) [role="radio"]{min-height:36px!important;padding:6px 13px!important;border:1px solid var(--ui-line)!important;border-radius:var(--ui-pill)!important;background:var(--ui-surface)!important;color:var(--ui-text-2)!important;font-size:13px!important;font-weight:500!important}[data-testid="stRadio"]:not(.app-topbar [data-testid="stRadio"]) [role="radio"][aria-checked="true"]{background:var(--ui-soft)!important;border-color:var(--ui-line)!important;color:var(--ui-text)!important;font-weight:600!important}.swap-result-summary{margin:8px 0 10px;padding:8px 12px;border:1px solid var(--ui-line-soft);border-radius:10px;background:var(--ui-soft);font-size:12px;color:var(--ui-muted)}.swap-result-summary strong{color:var(--ui-text);font-weight:600}.matrix-toolbar{display:flex;align-items:center;justify-content:space-between;gap:12px;margin:0 0 9px}.matrix-title{font-size:14px;font-weight:600;color:var(--ui-text)}.matrix-subtitle{font-size:12px;color:var(--ui-muted)}.matrix-legend{display:flex;gap:10px;flex-wrap:wrap;font-size:11px;color:var(--ui-muted);margin:0 0 8px 2px}.matrix-legend span{white-space:nowrap}[data-testid="stDataFrame"]{border:1px solid var(--ui-line)!important;border-radius:var(--ui-radius-md)!important;overflow:hidden!important;background:var(--ui-surface)!important;box-shadow:none!important}[data-testid="stDataFrame"] [role="columnheader"]{background:var(--ui-soft)!important;color:var(--ui-text-2)!important;font-size:12px!important;font-weight:600!important;border-right:1px solid var(--ui-line-soft)!important;border-bottom:1px solid var(--ui-line-soft)!important}[data-testid="stDataFrame"] [role="gridcell"]{background:var(--ui-surface)!important;color:var(--ui-text)!important;font-size:13px!important;border-right:1px solid var(--ui-line-soft)!important;border-bottom:1px solid var(--ui-line-soft)!important}.changed-teacher-block{margin:18px 0 6px}.changed-teacher-head{display:flex;align-items:center;justify-content:space-between;gap:16px;padding:0 2px 8px;border-bottom:1px solid var(--ui-line-soft)}.changed-teacher-name{font-size:16px;font-weight:650;color:var(--ui-text);letter-spacing:-.02em}.changed-teacher-index{display:inline-flex;align-items:center;justify-content:center;width:28px;height:28px;margin-right:7px;border-radius:8px;background:var(--ui-soft);color:var(--ui-muted);font-size:11px;font-weight:650}.changed-teacher-subject{margin-left:7px;color:var(--ui-muted);font-size:12px}.changed-teacher-status{font-size:11px;color:var(--ui-muted);white-space:nowrap}.changed-teacher-badge{display:inline-flex;align-items:center;margin-left:8px;padding:3px 7px;border-radius:999px;background:#fff3cd;color:#7a5b00;font-size:10px;font-weight:650;border:1px solid #f0d98a}.changed-teacher-divider{height:1px;background:var(--ui-line-soft);margin:24px 0 8px}.streamlit-header-safe-space{height:18px;width:100%;display:block;flex:0 0 auto}.changed-teacher-selector{display:flex;align-items:center;gap:8px;margin:0 0 10px}.changed-teacher-chip{display:inline-flex;align-items:center;padding:7px 10px;border:1px solid var(--ui-line-soft);border-radius:10px;background:var(--ui-soft-2);color:var(--ui-text);font-size:12px}.apple-note,.work-note{padding:10px 12px;margin:0 0 12px;background:var(--ui-soft-2);border:1px solid var(--ui-line-soft);border-radius:10px;color:var(--ui-muted);font-size:12px;line-height:1.5}.apple-note strong,.work-note strong{color:var(--ui-text)}.sandbox-title{margin:0 0 4px!important;font-size:24px!important;line-height:1.2!important;font-weight:600!important;letter-spacing:-.045em!important;color:var(--ui-text)!important}.swap-group-head{display:flex;align-items:center;justify-content:space-between;margin:14px 0 4px;padding-top:4px}.swap-group-title{font-size:13px;font-weight:600;color:var(--ui-text)}.swap-group-count{font-size:11px;color:var(--ui-muted);padding:3px 7px;border:1px solid var(--ui-line-soft);border-radius:999px;background:var(--ui-soft)}[data-testid="stDialog"]>div>div{max-height:calc(100vh - 5rem)!important;overflow-y:auto!important;background:var(--ui-surface)!important;color:var(--ui-text)!important;border:1px solid var(--ui-line)!important;border-radius:18px!important;box-shadow:0 18px 48px rgba(0,0,0,.14)!important}[data-testid="stDialog"] [data-testid="stVerticalBlockBorderWrapper"]{border:0!important;background:transparent!important}[data-testid="stDialog"] [data-testid="stSlider"]{margin:0!important;padding:3px 0!important}[data-testid="stDialog"] [data-testid="stSlider"]>label{display:none!important}[data-testid="stDialog"] [data-testid="stSlider"] [data-baseweb="slider"]{min-height:34px!important;margin:0!important}[data-baseweb="tab-list"]{gap:4px!important;border-bottom:1px solid var(--ui-line-soft)!important}[data-baseweb="tab"]{font-family:var(--app-font,"SF Pro Text",system-ui,sans-serif)!important;color:var(--ui-muted)!important}[data-baseweb="tab"][aria-selected="true"]{color:var(--ui-text)!important}a{color:var(--ui-accent)!important}hr,[data-testid="stDivider"]{border-color:var(--ui-line-soft)!important}[data-testid="stMetric"]{padding:8px 0!important}.ui-kpi-row{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:10px;margin:0 0 18px}.ui-kpi{padding:13px 15px;border:1px solid var(--ui-line-soft);border-radius:12px;background:var(--ui-surface)}.ui-kpi-label{font-size:11px;color:var(--ui-muted)}.ui-kpi-value{margin-top:4px;font-size:21px;font-weight:600;color:var(--ui-text)}@media(max-width:900px){.block-container{padding-left:12px!important;padding-right:12px!important}.app-top-safe-space{height:24px}.work-page-head{display:block}.work-page-meta{margin-top:8px}.ui-kpi-row{grid-template-columns:1fr}.app-topbar{overflow:hidden}}@media(prefers-color-scheme:dark){:root{--ui-bg:#000;--ui-surface:#1c1c1e;--ui-soft:#2c2c2e;--ui-soft-2:#232326;--ui-text:#f5f5f7;--ui-text-2:#e5e5ea;--ui-muted:#a1a1a6;--ui-muted-2:#8e8e93;--ui-line:#48484a;--ui-line-soft:#38383a;--ui-accent:#2997ff;--ui-accent-hover:#47a6ff;--ui-focus:rgba(41,151,255,.30)}[data-testid="stHeader"]{background:rgba(0,0,0,.76)!important}.stButton>button,.stDownloadButton>button,.stFormSubmitButton>button,[data-baseweb="input"],[data-baseweb="textarea"],[data-baseweb="select"]>div,[data-testid="stDateInput"]>div>div{background:var(--ui-surface)!important;color:var(--ui-text)!important;border-color:var(--ui-line)!important}[data-testid="stDataFrame"]{background:var(--ui-surface)!important;border-color:var(--ui-line)!important}[data-testid="stDataFrame"] [role="columnheader"]{background:var(--ui-soft)!important;color:var(--ui-text-2)!important}[data-testid="stDataFrame"] [role="gridcell"]{background:var(--ui-surface)!important;color:var(--ui-text)!important}[data-testid="stExpander"] summary:hover{background:var(--ui-soft)!important}[data-testid="stDialog"]>div>div{background:var(--ui-surface)!important;color:var(--ui-text)!important;border-color:var(--ui-line)!important}.ui-kpi{background:var(--ui-surface)!important;border-color:var(--ui-line-soft)!important}}
</style>
""", unsafe_allow_html=True)
st.markdown(r"""
<style>
:root{--ui-bg:#fff;--ui-surface:#fff;--ui-soft:#f5f5f7;--ui-soft-2:#fafafa;--ui-text:#1d1d1f;--ui-text-2:#424245;--ui-muted:#6e6e73;--ui-muted-2:#86868b;--ui-line:#d2d2d7;--ui-line-soft:#e8e8ed;--ui-accent:#0066cc;--ui-accent-hover:#0077ed;}.block-container{padding-top:5.5rem!important;padding-bottom:56px!important}.app-top-safe-space{height:18px!important}.app-topbar{min-height:58px!important;padding:0 0 10px!important;margin-bottom:26px!important;gap:18px!important;border-bottom:1px solid #e8e8ed!important}.app-identity{font-size:11px!important;color:#86868b!important;letter-spacing:-.01em!important}.app-identity strong{font-size:13px!important;color:#1d1d1f!important}.app-topbar .stButton>button{height:36px!important;min-height:36px!important;border-radius:10px!important;padding:5px 12px!important}.app-topbar [data-testid="stRadio"] [role="radiogroup"]{background:transparent!important;border:0!important;padding:0!important;gap:4px!important}.app-topbar [data-testid="stRadio"] [role="radio"]{height:36px!important;min-height:36px!important;padding:0 13px!important;border-radius:9px!important;font-size:12px!important;color:#6e6e73!important}.app-topbar [data-testid="stRadio"] [role="radio"][aria-checked="true"]{background:#f5f5f7!important;color:#1d1d1f!important;box-shadow:none!important}.work-page-head{padding:0 2px 20px!important;margin-bottom:22px!important;border-bottom:0!important;align-items:center!important}.work-page-head h1{font-size:32px!important;line-height:1.1!important;letter-spacing:-.045em!important}.work-page-head p{font-size:13px!important;margin-top:7px!important;color:#86868b!important}.work-page-meta{font-size:11px!important;color:#86868b!important;padding:7px 10px!important;border:1px solid #e8e8ed!important;border-radius:999px!important;background:#fafafa!important}.work-section{margin-bottom:30px!important}.work-section-title{font-size:15px!important;margin-bottom:5px!important}.work-section-note{font-size:12px!important;color:#86868b!important}[data-testid="stDateInput"] label,[data-testid="stSelectbox"] label,[data-testid="stRadio"] label,[data-testid="stTextInput"] label,[data-testid="stNumberInput"] label{font-size:11px!important;color:#6e6e73!important;font-weight:500!important;margin-bottom:5px!important}[data-baseweb="input"],[data-baseweb="textarea"],[data-baseweb="select"]>div,[data-testid="stDateInput"]>div>div{border-color:#d2d2d7!important;border-radius:9px!important;min-height:38px!important}[data-testid="stRadio"]:not(.app-topbar [data-testid="stRadio"]) [role="radiogroup"]{gap:5px!important}[data-testid="stRadio"]:not(.app-topbar [data-testid="stRadio"]) [role="radio"]{min-height:36px!important;border-radius:9px!important;padding:5px 12px!important}.stButton>button,.stDownloadButton>button,.stFormSubmitButton>button{min-height:38px!important;border-radius:9px!important;font-size:12px!important;padding:6px 14px!important}.stButton>button[kind="primary"],.stFormSubmitButton>button[kind="primary"]{border-radius:9px!important}.matrix-toolbar{margin:0 0 8px!important}.matrix-title{font-size:15px!important;font-weight:600!important;letter-spacing:-.02em!important}.matrix-subtitle{font-size:11px!important;color:#86868b!important}.matrix-legend{gap:14px!important;margin:0 0 9px 2px!important;font-size:11px!important;color:#86868b!important}[data-testid="stDataFrame"]{border:1px solid #d2d2d7!important;border-radius:10px!important;box-shadow:none!important}[data-testid="stDataFrame"] [role="columnheader"]{background:#f5f5f7!important;color:#424245!important;font-size:11px!important;font-weight:600!important;border-color:#e8e8ed!important}[data-testid="stDataFrame"] [role="gridcell"]{font-size:12px!important;border-color:#eeeeef!important}.swap-result-summary{margin:12px 0!important;padding:10px 13px!important;border-radius:9px!important;background:#f5f5f7!important;border-color:#e8e8ed!important}.swap-group-head{margin:20px 0 6px!important;padding:0!important}.swap-group-title{font-size:13px!important}.swap-group-count{border:0!important;background:#f5f5f7!important;border-radius:999px!important}[data-testid="stDialog"]>div>div{border:1px solid #d2d2d7!important;border-radius:16px!important;box-shadow:0 16px 50px rgba(0,0,0,.12)!important}[data-testid="stDialog"] h1,[data-testid="stDialog"] h2,[data-testid="stDialog"] h3{letter-spacing:-.03em!important}[data-testid="stDialog"] .stButton>button{min-height:40px!important}[data-testid="stAlert"]{border-radius:9px!important;font-size:12px!important}.apple-note,.work-note{padding:10px 12px!important;border-radius:9px!important;background:#fafafa!important}.sandbox-title{font-size:28px!important;letter-spacing:-.045em!important;margin-bottom:7px!important}@media(max-width:1400px){.block-container{padding-left:22px!important;padding-right:22px!important}.app-topbar{gap:10px!important}.app-topbar [data-testid="stRadio"] [role="radio"]{padding:0 10px!important}}@media(max-width:900px){.block-container{padding-top:5rem!important;padding-left:14px!important;padding-right:14px!important}.work-page-head{display:block!important}.work-page-meta{display:inline-block;margin-top:10px}}[data-testid="st-key-top_quick_refresh"],[data-testid="st-key-top_quick_load"],[data-testid="st-key-top_quick_save"]{margin:0!important;}[data-testid="st-key-top_quick_refresh"] button,[data-testid="st-key-top_quick_load"] button,[data-testid="st-key-top_quick_save"] button{min-height:38px!important;height:38px!important;padding:5px 9px!important;border-radius:11px!important;border:1px solid #dedee3!important;background:#fff!important;color:#242426!important;box-shadow:0 1px 2px rgba(0,0,0,.04),inset 0 1px 0 rgba(255,255,255,.65)!important;font-size:11px!important;font-weight:600!important;letter-spacing:-.03em!important;white-space:nowrap!important;}[data-testid="st-key-top_quick_refresh"] button:hover,[data-testid="st-key-top_quick_load"] button:hover{background:#f6f6f8!important;border-color:#cfcfd5!important;transform:translateY(-1px);}[data-testid="st-key-top_quick_save"] button,[data-testid="st-key-top_quick_save"] button[kind="primary"]{background:#1d1d1f!important;border-color:#1d1d1f!important;color:#fff!important;box-shadow:0 2px 5px rgba(0,0,0,.12)!important;}[data-testid="st-key-top_quick_save"] button:hover{background:#2d2d30!important;border-color:#2d2d30!important;transform:translateY(-1px);}@media(max-width:1100px){.top-quick-step-hide{display:none!important}}
</style>
""", unsafe_allow_html=True)
SCHOOL_NAME = "서라벌여자중학교"
SCHOOL_YEAR = "2026"
APP_VERSION = "5.0-Apple-ProductUX-Renewal"
GITHUB_FONT_BASE = "https://cdn.jsdelivr.net/gh/pse915/Timetable@main"
GITHUB_FONT_RAW_BASE = "https://raw.githubusercontent.com/pse915/Timetable/main"
GITHUB_FONT_FAMILY = "Pse Noto Sans KR"
HAKYO_FONT_FAMILY = "Hakgyoansim Wooju R"
UI_FONT_OPTIONS = {
    "시스템 기본 (Apple / Windows)": '"SF Pro Text", "SF Pro Display", -apple-system, BlinkMacSystemFont, "Segoe UI", "Apple SD Gothic Neo", "Noto Sans KR", sans-serif',
    "Pretendard": '"Pretendard", "Apple SD Gothic Neo", "Noto Sans KR", "Segoe UI", sans-serif',
    "Noto Sans KR": '"Noto Sans KR", "Apple SD Gothic Neo", "Malgun Gothic", "Segoe UI", sans-serif',
    "Noto Sans KR · GitHub (Light + Bold)": f'"{GITHUB_FONT_FAMILY}", "Noto Sans KR", "Apple SD Gothic Neo", "Malgun Gothic", sans-serif',
    "학교안심 우주체 · GitHub": f'"{HAKYO_FONT_FAMILY}", "Apple SD Gothic Neo", "Malgun Gothic", "Noto Sans KR", sans-serif',
    "Inter": '"Inter", "Pretendard", "Noto Sans KR", "Segoe UI", sans-serif',
}
UI_FONT_DEFAULT = "시스템 기본 (Apple / Windows)"
def render_font_runtime_css(selected_font: str):
    selected_font = selected_font if selected_font in UI_FONT_OPTIONS else UI_FONT_DEFAULT
    if selected_font == "Noto Sans KR · GitHub (Light + Bold)":
        light_url = f"{GITHUB_FONT_BASE}/NotoSansKR-Light.ttf"
        bold_url = f"{GITHUB_FONT_BASE}/NotoSansKR-Bold.ttf"
        light_raw = f"{GITHUB_FONT_RAW_BASE}/NotoSansKR-Light.ttf"
        bold_raw = f"{GITHUB_FONT_RAW_BASE}/NotoSansKR-Bold.ttf"
        st.markdown(
            f"""
<style id="github-noto-sans-kr-runtime-font">
@font-face {{ font-family: '{GITHUB_FONT_FAMILY}'; font-style: normal; font-weight: 300; font-display: swap;
  src: url('{light_url}') format('truetype'), url('{light_raw}') format('truetype'); }}
@font-face {{ font-family: '{GITHUB_FONT_FAMILY}'; font-style: normal; font-weight: 700; font-display: swap;
  src: url('{bold_url}') format('truetype'), url('{bold_raw}') format('truetype'); }}
:root {{ --app-font: '{GITHUB_FONT_FAMILY}', 'Noto Sans KR', 'Apple SD Gothic Neo', 'Malgun Gothic', sans-serif; }}
p,label,h1,h2,h3,h4,h5,h6,button,input,textarea,select,[data-testid="stWidgetLabel"],[data-testid="stCaptionContainer"],.stMarkdown,.stCaption,[role="radio"],[role="tab"],[data-baseweb="select"],[data-baseweb="input"] input,[data-baseweb="textarea"] textarea {{ font-family: var(--app-font) !important; }}
[data-testid="stDataFrame"],[data-testid="stDataEditor"],[data-testid="stDataFrame"] *,[data-testid="stDataEditor"] * {{ --gdg-font-family: var(--app-font) !important; font-family: var(--app-font) !important; }}
table,thead,tbody,tfoot,tr,th,td,caption {{ font-family: var(--app-font) !important; }}
</style>
""",
            unsafe_allow_html=True,
        )
        return
    if selected_font == "학교안심 우주체 · GitHub":
        hakyo_url = f"{GITHUB_FONT_BASE}/HakgyoansimWoojuR.ttf"
        hakyo_raw = f"{GITHUB_FONT_RAW_BASE}/HakgyoansimWoojuR.ttf"
        st.markdown(
            f"""
<style id="github-hakgyoansim-wooju-runtime-font">
@font-face {{ font-family: '{HAKYO_FONT_FAMILY}'; font-style: normal; font-weight: 400; font-display: swap;
  src: url('{hakyo_url}') format('truetype'), url('{hakyo_raw}') format('truetype'); }}
:root {{ --app-font: '{HAKYO_FONT_FAMILY}', 'Apple SD Gothic Neo', 'Malgun Gothic', 'Noto Sans KR', sans-serif; }}
p,label,h1,h2,h3,h4,h5,h6,button,input,textarea,select,[data-testid="stWidgetLabel"],[data-testid="stCaptionContainer"],.stMarkdown,.stCaption,[role="radio"],[role="tab"],[data-baseweb="select"],[data-baseweb="input"] input,[data-baseweb="textarea"] textarea {{ font-family: var(--app-font) !important; }}
[data-testid="stDataFrame"],[data-testid="stDataEditor"],[data-testid="stDataFrame"] *,[data-testid="stDataEditor"] * {{ --gdg-font-family: var(--app-font) !important; font-family: var(--app-font) !important; }}
table,thead,tbody,tfoot,tr,th,td,caption {{ font-family: var(--app-font) !important; }}
</style>
""",
            unsafe_allow_html=True,
        )
        return
    st.markdown(
        f"""<style id="runtime-app-font-global">
:root {{ --app-font: {UI_FONT_OPTIONS[selected_font]}; }}
html,body,[data-testid="stAppViewContainer"],[data-testid="stAppViewContainer"] *{{font-family:var(--app-font)!important}}
body table,body table *,[data-testid="stDataFrame"],[data-testid="stDataFrame"] *,[data-testid="stDataEditor"],[data-testid="stDataEditor"] *{{font-family:var(--app-font)!important;--gdg-font-family:var(--app-font)!important}}
[class*="material-symbols"],[data-testid="stIconMaterial"],[data-testid="stIconMaterial"] *{{font-family:"Material Symbols Rounded","Material Symbols Outlined",sans-serif!important}}
</style>""",
        unsafe_allow_html=True,
    )
render_font_runtime_css(st.session_state.get("ui_font", UI_FONT_DEFAULT))
DAYS = ["월", "화", "수", "목", "금"]
PERIODS_PER_DAY = {"월": 6, "화": 7, "수": 7, "목": 7, "금": 6}
MAX_PERIOD = 7
WEEKDAY_KR = {0: "월", 1: "화", 2: "수", 3: "목", 4: "금", 5: "토", 6: "일"}
SCHOOL_WEEKDAYS = (0, 1, 2, 3, 4)
_ISO_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
KST = ZoneInfo("Asia/Seoul")
_GSHEET_NETWORK_SEMAPHORE = threading.BoundedSemaphore(2)
def _today_kst() -> date:
    return datetime.now(KST).date()
def _is_current_week(ref_date: date, today: date | None = None) -> bool:
    today = today or _today_kst()
    return (ref_date - timedelta(days=ref_date.weekday())) == (today - timedelta(days=today.weekday()))
def _hide_past_week_slots(matrix: pd.DataFrame, ref_date: date, *, hide_past=True) -> pd.DataFrame:
    if not hide_past or matrix is None or matrix.empty or not _is_current_week(ref_date):
        return matrix
    today = _today_kst()
    monday = ref_date - timedelta(days=ref_date.weekday())
    out = matrix.copy()
    for i, day in enumerate(DAYS):
        day_date = monday + timedelta(days=i)
        if day_date >= today:
            continue
        for p in range(1, PERIODS_PER_DAY.get(day, MAX_PERIOD) + 1):
            col = f"{day}{p}"
            if col in out.columns:
                out[col] = ""
    return out
TIMETABLE_SHEET_ID = "1jZhTHyJ8vKXn6tkoFXfY_f52-pj6eQTdVvRCo3cCmBA"
WORK_SHEET_ID = "1g1B1cyZG_tfRn3AD1NZzr30YxYNYFewJeZYdos2obpU"
MAX_HISTORY = 5
ABSENCE_REASONS = ["병가", "연가", "출장", "공가", "조퇴", "외출", "연수", "특별휴가", "기타"]
MAX_LOGIN_ATTEMPTS = 5
SUB_COST = 10000
SUBJECT_GROUP = {
    "국어1": "국어", "국어2": "국어", "사회": "사회", "사회1": "사회", "사회2": "사회", "사회3": "사회",
    "역사": "역사", "도덕1": "도덕", "도덕2": "도덕", "수학": "수학", "수학1": "수학", "수학2": "수학",
    "과학": "과학", "과학1": "과학", "과학2": "과학", "기가": "기술가정",
    "체육1": "체육", "체육2": "체육", "체육3": "체육", "스포": "스포츠",
    "음악": "음악", "음악1": "음악", "음악2": "음악", "미술": "미술",
    "영어": "영어", "영어1": "영어", "영어2": "영어", "영회": "영어",
    "한문": "한문", "일본어": "일본어", "정보": "정보", "진동": "진로활동",
}
ROLE_MASTER = "마스터"
ROLE_EDU = "교육과정부"
ROLE_OFFICE = "교무계원"
ROLE_TEACHER = "일반교사"
ROLE_GUEST = "게스트"
MASTER_ID = "pse915"
VALID_ROLES = frozenset({ROLE_MASTER, ROLE_EDU, ROLE_OFFICE, ROLE_TEACHER})
MAX_ID_LENGTH = 64
MAX_NAME_LENGTH = 80
MAX_EMAIL_LENGTH = 254
MAX_MEMO_LENGTH = 1000
MAX_TAB_TEXT_LENGTH = 4000
AUTH_RECHECK_SECONDS = 60.0
_FORMULA_PREFIX_RE = re.compile(r"^[=+@]")
_NUMERIC_TEXT_RE = re.compile(r"^-?(?:\d+(?:\.\d*)?|\.\d+)$")
_CONTROL_CHAR_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
def _clean_user_text(value, max_length=MAX_MEMO_LENGTH):
    if value is None:
        return ""
    text = _CONTROL_CHAR_RE.sub("", str(value)).strip()
    return text[:max_length]
def _safe_sheet_text(value):
    if not isinstance(value, str):
        return value
    text = _CONTROL_CHAR_RE.sub("", value)
    if _FORMULA_PREFIX_RE.match(text):
        return "'" + text
    if text.startswith("-") and not _NUMERIC_TEXT_RE.fullmatch(text):
        return "'" + text
    return text
def _validated_role(role, fallback=ROLE_TEACHER):
    role = _clean_user_text(role, 30)
    return role if role in VALID_ROLES else fallback
def _now_text(fmt="%Y-%m-%d %H:%M:%S"):
    return datetime.now(KST).strftime(fmt)
ALL_TABS = [
    "시간표 조회",
    "결강·보강",
    "시간표 맞교환 & 변경 추천",
    "시간표 변경 테스트용",
    "변경된 교사 주간표",
    "시간강사 관리",
    "통계",
    "📋 복무 관리 & 판단",
    "🛠️ 다중 출장·전체 조정 추천",
    "🔑 아이디·권한 관리",
    "📑 회원별 탭 권한 관리"
]
OFFICE_TABS = [
    "교무호봉획정",
]
ALL_APP_TABS = ALL_TABS + OFFICE_TABS
NAV_LABELS = {
    "시간표 조회": "시간표",
    "시간강사 관리": "시간강사",
    "결강·보강": "결강·보강",
    "시간표 맞교환 & 변경 추천": "맞교환",
    "통계": "통계",
    "시간표 변경 테스트용": "테스트",
    "변경된 교사 주간표": "변경 교사",
    "📋 복무 관리 & 판단": "복무",
    "🛠️ 다중 출장·전체 조정 추천": "다중 조정",
    "🔑 아이디·권한 관리": "아이디",
    "📑 회원별 탭 권한 관리": "탭 권한",
    "교무호봉획정": "호봉획정",
}
DEFAULT_TABS = {
    ROLE_MASTER: ALL_APP_TABS,
    ROLE_EDU: ALL_APP_TABS,
    ROLE_OFFICE: OFFICE_TABS,
    ROLE_TEACHER: [
        "시간표 조회", "시간강사 관리", "결강·보강",
        "시간표 맞교환 & 변경 추천", "통계",
        "시간표 변경 테스트용", "변경된 교사 주간표", "📋 복무 관리 & 판단",
        "교무호봉획정"
    ],
    ROLE_GUEST: []
}
def safe_int(val, default=0):
    try:
        if pd.isna(val) or val is None or str(val).strip() in ("", "nan", "None"):
            return default
        return int(float(str(val).strip()))
    except Exception:
        return default
def normalize_date_str(d_str):
    if d_str is None or not d_str:
        return ""
    try:
        if pd.isna(d_str):
            return ""
    except (TypeError, ValueError):
        pass
    text = str(d_str).strip()
    if text.lower() in ("", "nan", "none"):
        return ""
    if _ISO_DATE_RE.fullmatch(text):
        try:
            date.fromisoformat(text)
            return text
        except ValueError:
            pass
    try:
        return pd.to_datetime(text, errors="coerce").strftime("%Y-%m-%d")
    except Exception:
        return text
@lru_cache(maxsize=512)
def subject_group(subject: str) -> str:
    if not isinstance(subject, str) or not subject.strip():
        return ""
    s = subject.strip()
    return SUBJECT_GROUP.get(s, s.rstrip("0123456789"))
@lru_cache(maxsize=512)
def grade_of(class_name: str) -> str:
    if isinstance(class_name, str) and "-" in class_name:
        return class_name.split("-", 1)[0]
    return ""
def format_periods(periods):
    normalized = [safe_int(p) for p in periods]
    periods = sorted({p for p in normalized if p >= 0})
    if not periods:
        return ""
    if 0 in periods:
        return "전체"
    if len(periods) == 1:
        return f"{periods[0]}교시"
    ranges = []
    start = prev = periods[0]
    for p in periods[1:]:
        if p == prev + 1:
            prev = p
        else:
            ranges.append(f"{start}~{prev}교시" if start != prev else f"{start}교시")
            start = prev = p
    ranges.append(f"{start}~{prev}교시" if start != prev else f"{start}교시")
    return ", ".join(ranges)
def _month_calendar_df(month_anchor: date, selected_dates=None, range_start=None, range_end=None):
    selected_dates = set(selected_dates or [])
    first = month_anchor.replace(day=1)
    next_month = (first.replace(day=28) + timedelta(days=4)).replace(day=1)
    last = next_month - timedelta(days=1)
    grid_start = first - timedelta(days=first.weekday())
    grid_end = last + timedelta(days=(6 - last.weekday()))
    weekdays = ["월", "화", "수", "목", "금"]
    rows = []
    cur = grid_start
    while cur <= grid_end:
        row = {}
        for i, wd in enumerate(weekdays):
            d = cur + timedelta(days=i)
            if d.month != first.month:
                row[wd] = ""
            else:
                mark = ""
                if d in selected_dates:
                    mark += "● "
                if range_start and d == range_start:
                    mark += "시작 "
                if range_end and d == range_end:
                    mark += "종료 "
                if range_start and range_end and range_start < d < range_end:
                    mark += "■ "
                row[wd] = f"{mark}{d.day:02d}"
        rows.append(row)
        cur += timedelta(days=7)
    return pd.DataFrame(rows, columns=weekdays)
def _calendar_set_day(state_key, picked):
    st.session_state[state_key] = picked

def _calendar_prev_month(month_key):
    m = st.session_state[month_key]
    st.session_state[month_key] = (m.replace(day=1) - timedelta(days=1)).replace(day=1)

def _calendar_next_month(month_key):
    m = st.session_state[month_key]
    st.session_state[month_key] = (m.replace(day=28) + timedelta(days=4)).replace(day=1)

def _calendar_today(month_key, selected_key):
    today = _today_kst()
    if today.weekday() >= 5:
        today = today - timedelta(days=today.weekday() - 4)
    st.session_state[month_key] = today.replace(day=1)
    st.session_state[selected_key] = today

def calendar_picker(label, value=None, key="calendar", help_text=None, rerun_scope=None):
    value = value or _today_kst()
    if value.weekday() >= 5:
        value = value - timedelta(days=value.weekday() - 4)
    month_key = f"_{key}_month"
    selected_key = f"_{key}_selected"
    if month_key not in st.session_state:
        st.session_state[month_key] = value.replace(day=1)
    if selected_key not in st.session_state:
        st.session_state[selected_key] = value
    elif st.session_state[selected_key].weekday() >= 5:
        selected = st.session_state[selected_key]
        st.session_state[selected_key] = selected - timedelta(days=selected.weekday() - 4)
    st.markdown(f"**{label}**")
    nav1, nav2, nav3 = st.columns([1, 4, 1])
    with nav1:
        st.button("◀", key=f"{key}_prev", width="stretch", on_click=_calendar_prev_month, args=(month_key,))
    with nav2:
        st.markdown(f"<div style='text-align:center;font-weight:700;font-size:1.05rem'>{st.session_state[month_key].year}년 {st.session_state[month_key].month}월</div>", unsafe_allow_html=True)
    with nav3:
        st.button("▶", key=f"{key}_next", width="stretch", on_click=_calendar_next_month, args=(month_key,))
    quick1, quick2 = st.columns([1, 5])
    with quick1:
        st.button("오늘", key=f"{key}_today", width="stretch", on_click=_calendar_today, args=(month_key, selected_key))
    with quick2:
        st.caption(f"선택: **{st.session_state[selected_key]:%Y-%m-%d}** · 평일(월~금)만 표시됩니다. 날짜 셀을 클릭하세요.")
    weekdays = ["월", "화", "수", "목", "금"]
    month_anchor = st.session_state[month_key]
    first = month_anchor.replace(day=1)
    next_month = (first.replace(day=28) + timedelta(days=4)).replace(day=1)
    last = next_month - timedelta(days=1)
    grid_start = first - timedelta(days=first.weekday())
    grid_end = last + timedelta(days=6 - last.weekday())
    style_id = re.sub(r"[^A-Za-z0-9_-]", "_", str(key))
    st.markdown(f"""<style id="calendar-grid-{style_id}">
[data-testid^="st-key-{key}_cell_"] button{{min-height:34px!important;height:34px!important;padding:3px 6px!important;border:1px solid #e5e5ea!important;border-radius:7px!important;background:#fff!important;color:#1d1d1f!important;font-size:12px!important;font-weight:500!important;box-shadow:none!important}}
[data-testid^="st-key-{key}_cell_"] button:hover{{background:#f5f5f7!important;border-color:#d2d2d7!important}}
[data-testid^="st-key-{key}_cell_"] button[kind="primary"]{{background:#f5f5f7!important;border-color:#d2d2d7!important;color:#1d1d1f!important;font-weight:700!important}}
[data-testid^="st-key-{key}_blank_"] button{{visibility:hidden!important;height:34px!important;min-height:34px!important;padding:0!important;border:0!important;background:transparent!important;pointer-events:none!important}}
.calendar-weekday-{style_id}{{text-align:center;font-size:11px;font-weight:600;color:#6e6e73;padding:2px 0 6px}}
</style>""", unsafe_allow_html=True)
    header_cols = st.columns(5)
    for i, wd in enumerate(weekdays):
        with header_cols[i]:
            st.markdown(f'<div class="calendar-weekday-{style_id}">{wd}</div>', unsafe_allow_html=True)
    cur = grid_start
    selected = st.session_state[selected_key]
    while cur <= grid_end:
        row_cols = st.columns(5)
        for i in range(5):
            d = cur + timedelta(days=i)
            with row_cols[i]:
                if d.month != first.month:
                    st.button(" ", key=f"{key}_blank_{d:%Y%m%d}", disabled=True, width="stretch")
                    continue
                active = d == selected
                st.button(
                    f"✓ {d.day:02d}" if active else f"{d.day:02d}",
                    key=f"{key}_cell_{d:%Y%m%d}",
                    width="stretch",
                    type="primary" if active else "secondary",
                    on_click=_calendar_set_day,
                    args=(selected_key, d),
                )
        cur += timedelta(days=7)
    if help_text:
        st.caption(help_text)
    return st.session_state[selected_key]
def calendar_range_picker(start_value=None, end_value=None, key="calendar_range", help_text=None):
    start_value = start_value or _today_kst()
    end_value = end_value or start_value
    c1, c2 = st.columns(2)
    with c1:
        start_date = calendar_picker("시작일", start_value, f"{key}_start")
    with c2:
        end_date = calendar_picker("종료일", end_value, f"{key}_end")
    if start_date > end_date:
        st.warning("시작일이 종료일보다 늦습니다. 종료일 달력에서 시작일 이후 날짜를 선택하세요.")
    else:
        st.caption(f"선택 기간: **{start_date:%Y-%m-%d} ~ {end_date:%Y-%m-%d}**")
    if help_text:
        st.caption(help_text)
    return start_date, end_date
def period_matrix_picker(label, key, selected=None, allow_all=True, rerun_scope=None, single=False, state_key=None, allowed_periods=None):
    allowed = None if allowed_periods is None else {safe_int(x) for x in allowed_periods if 1 <= safe_int(x) <= 7}
    initial = set(safe_int(x) for x in (selected or []))
    if allowed is not None:
        initial.intersection_update(allowed)
    if state_key:
        saved = safe_int(st.session_state.get(state_key, 0))
        if single and saved > 0:
            initial = {saved} if allowed is None or saved in allowed else set()
            if saved > 0 and not initial:
                st.session_state[state_key] = 0
    selected = initial
    st.markdown(f"**{label}**")
    if allowed is not None:
        st.caption("선택한 교사가 해당 날짜에 수업이 없는 교시만 표시됩니다.")
    cols = st.columns(7)
    for p, col in enumerate(cols, 1):
        with col:
            if allowed is not None and p not in allowed:
                st.empty()
                continue
            active = p in selected
            if st.button(f"{'✓ ' if active else ''}{p}교시", key=f"{key}_{p}", width="stretch", type="primary" if active else "secondary"):
                if single:
                    selected = {p}
                    if state_key:
                        st.session_state[state_key] = p
                elif active:
                    selected.remove(p)
                else:
                    selected.add(p)
                (st.rerun(scope=rerun_scope) if rerun_scope else st.rerun())
    if allow_all:
        all_active = 0 in selected
        if st.button(f"{'✓ ' if all_active else ''}하루 전체", key=f"{key}_all", width="stretch", type="primary" if all_active else "secondary"):
            if all_active:
                selected.clear()
                if state_key:
                    st.session_state[state_key] = 0
            else:
                selected = {0}
                if state_key:
                    st.session_state[state_key] = 0
            (st.rerun(scope=rerun_scope) if rerun_scope else st.rerun())
    if single and state_key:
        saved = safe_int(st.session_state.get(state_key, 0))
        return [saved] if 1 <= saved <= 7 and (allowed is None or saved in allowed) else []
    if 0 in selected:
        return [0]
    return sorted(p for p in selected if 1 <= p <= 7 and (allowed is None or p in allowed))

def _teacher_empty_periods_on_date(teacher, on_date, version=0, use_test=False):
    teacher = str(teacher or '').strip()
    norm = normalize_date_str(on_date)
    if not teacher or not norm:
        return []
    try:
        d = date.fromisoformat(norm)
    except ValueError:
        return []
    max_period = PERIODS_PER_DAY.get(WEEKDAY_KR[d.weekday()], MAX_PERIOD)
    tt = get_effective_timetable_for_date(norm, version, use_test=use_test)
    occupied = set()
    if isinstance(tt, pd.DataFrame) and not tt.empty:
        for r in tt.itertuples(index=False):
            if str(getattr(r, '교사명', '')).strip() == teacher:
                p = safe_int(getattr(r, '교시', 0))
                if 1 <= p <= max_period:
                    occupied.add(p)
    return [p for p in range(1, max_period + 1) if p not in occupied]
def _daily_schedule_matrix(ref_date: date, *, teacher_filter=None, use_test=False, version=0):
    norm = ref_date.strftime("%Y-%m-%d") if isinstance(ref_date, date) else normalize_date_str(ref_date)
    e = get_effective_timetable_for_date(norm, version, use_test=use_test)
    names = []
    if teacher_filter:
        names = [str(teacher_filter).strip()]
    else:
        base = st.session_state.get("teachers", pd.DataFrame())
        if not base.empty and "교사명" in base.columns:
            names = base["교사명"].dropna().astype(str).str.strip().tolist()
        if not e.empty:
            names += e["교사명"].dropna().astype(str).str.strip().tolist()
        names = sorted(set(x for x in names if x))
    idx = {}
    if not e.empty:
        for r in e.itertuples(index=False):
            idx[(str(r.교사명).strip(), safe_int(r.교시))] = r
    rows=[]
    for t in names:
        row={"교사명":t}
        for pno in range(1, MAX_PERIOD+1):
            r=idx.get((t,pno))
            if r is None:
                row[f"{pno}교시"]=""
                continue
            cell=f"{str(r.학급).strip()} {str(r.과목).strip()}".strip()
            typ=str(getattr(r,"변경유형","원본")).strip()
            if typ=="교환": cell += " 🔄"
            elif typ=="테스트교환": cell += " 🧪"
            elif typ=="보강": cell += " 🟢"
            elif typ=="시간강사": cell += f" 🟡 {str(getattr(r,'원본교사','')).strip()}→시간강사"
            row[f"{pno}교시"]=cell
        rows.append(row)
    return pd.DataFrame(rows, columns=["교사명"]+[f"{p}교시" for p in range(1,MAX_PERIOD+1)])
def daily_schedule_picker(ref_date=None, key="daily_schedule", *, teacher_filter=None, use_test=False,
                          multi=False, height=430, help_text=None):
    ref_date = ref_date or _today_kst()
    picked_date = calendar_picker("날짜", ref_date, key=f"{key}_date")
    ver=st.session_state.get("_data_version",0)
    matrix=_daily_schedule_matrix(picked_date, teacher_filter=teacher_filter, use_test=use_test, version=ver)
    st.caption(f"{picked_date:%Y-%m-%d} ({WEEKDAY_KR[picked_date.weekday()]}) · 수업 셀을 클릭하면 작업 대상이 선택됩니다.")
    if matrix.empty:
        st.info("선택한 날짜의 시간표가 없습니다.")
        return picked_date, matrix, []
    teacher_key = str(teacher_filter or "all").strip().replace(" ", "_")
    matrix_key = f"{key}_matrix_{picked_date:%Y%m%d}_{teacher_key}_{'test' if use_test else 'live'}"
    state_key = f"_{key}_selected_cells_{picked_date:%Y%m%d}_{teacher_key}_{'test' if use_test else 'live'}"
    event=st.dataframe(_daily_styled_matrix(matrix), hide_index=True, width="stretch", height=height,
                       key=matrix_key, on_select="rerun",
                       selection_mode="multi-cell" if multi else "single-cell")
    cells=getattr(getattr(event,"selection",None),"cells",[]) or []
    if cells:
        st.session_state[state_key]=list(cells)
    elif event is not None and state_key in st.session_state:
        st.session_state[state_key]=[]
    saved_cells=st.session_state.get(state_key,[])
    selections=[]
    for row_idx,col_name in saved_cells:
        if row_idx<0 or row_idx>=len(matrix) or col_name=="교사명": continue
        teacher=str(matrix.iloc[row_idx]["교사명"]).strip()
        period=safe_int(str(col_name).replace("교시",""))
        value=str(matrix.iloc[row_idx][col_name]).strip()
        if teacher and period>0 and value:
            selections.append({"교사명":teacher,"일자":picked_date.strftime("%Y-%m-%d"),
                               "요일":WEEKDAY_KR[picked_date.weekday()],"교시":period,"표시":value})
    if selections:
        st.success("선택: " + " · ".join(f"{x['교사명']} {x['교시']}교시" for x in selections[:8]))
    if help_text: st.caption(help_text)
    return picked_date, matrix, selections
def range_calendar_matrix_picker(start_value=None, end_value=None, key="range_matrix"):
    start_value=start_value or _today_kst(); end_value=end_value or start_value
    st.markdown("#### 📅 기간 선택")
    start_date,end_date=calendar_range_picker(start_value,end_value,key=f"{key}_dates")
    st.markdown("#### 🕐 교시 범위 선택")
    c1,c2=st.columns(2)
    with c1:
        start_p=period_matrix_picker("시작일 적용 교시",f"{key}_start_period",st.session_state.get(f"{key}_start_periods",[]))
    with c2:
        end_p=period_matrix_picker("종료일 적용 교시",f"{key}_end_period",st.session_state.get(f"{key}_end_periods",[]))
    st.session_state[f"{key}_start_periods"]=start_p
    st.session_state[f"{key}_end_periods"]=end_p
    return start_date,end_date,start_p,end_p
def get_all_teacher_names():
    ts = []
    if "teachers" in st.session_state and not st.session_state.teachers.empty and "교사명" in st.session_state.teachers.columns:
        ts = st.session_state.teachers["교사명"].dropna().astype(str).str.strip().tolist()
    pts = []
    pt = st.session_state.get("part_time", pd.DataFrame())
    if not pt.empty and "시간강사명" in pt.columns:
        pts = pt["시간강사명"].dropna().astype(str).str.strip().tolist()
    return sorted(set([t for t in ts + pts if t]))
def current_user():
    return st.session_state.get("user_id", "")
def current_name():
    return st.session_state.get("user_name", "")
def current_role():
    return st.session_state.get("user_role", ROLE_GUEST)
def is_master():
    return current_role() == ROLE_MASTER
def is_edu_or_master():
    return current_role() in (ROLE_MASTER, ROLE_EDU)
def is_teacher():
    return current_role() == ROLE_TEACHER
def can_manage_ids():
    return is_edu_or_master()
def can_full_data():
    return is_edu_or_master()
def get_user_allowed_tabs():
    role = current_role()
    allowed_str = st.session_state.get("user_allowed_tabs", "")
    if allowed_str and isinstance(allowed_str, str) and allowed_str.strip():
        tabs = [t.strip() for t in allowed_str.split(",") if t.strip()]
        if not can_manage_ids():
            tabs = [t for t in tabs if t not in ("🔑 아이디·권한 관리", "📑 회원별 탭 권한 관리", "🛠️ 다중 출장·전체 조정 추천")]
        return tabs if tabs else DEFAULT_TABS.get(role, [])
    return DEFAULT_TABS.get(role, [])
def get_teacher_subject(teacher_name: str) -> str:
    teachers = st.session_state.get("teachers", pd.DataFrame())
    if not teachers.empty and "교사명" in teachers.columns:
        row = teachers[teachers["교사명"] == teacher_name]
        if not row.empty:
            for col in ["담당과목", "과목", "교과", "전공"]:
                if col in row.columns:
                    val = str(row.iloc[0][col]).strip()
                    if val and val not in ("nan", "None", ""):
                        return val
    tt = st.session_state.get("timetable", pd.DataFrame())
    if not tt.empty:
        sub = tt[tt["교사명"] == teacher_name]
        if not sub.empty and "과목군" in sub.columns:
            most = sub["과목군"].value_counts()
            if not most.empty:
                return most.index[0]
        if not sub.empty and "과목" in sub.columns:
            most = sub["과목"].value_counts()
            if not most.empty:
                return subject_group(most.index[0])
    return ""
@st.cache_resource(show_spinner=False)
def get_gspread_client():
    scopes = ["https://www.googleapis.com/auth/spreadsheets"]
    if "gcp_service_account" not in st.secrets:
        raise RuntimeError("GCP Secrets에 gcp_service_account가 없습니다.")
    creds = Credentials.from_service_account_info(st.secrets["gcp_service_account"], scopes=scopes)
    return gspread.authorize(creds)
def _clear_gsheet_runtime_cache():
    st.session_state.pop("_gsheet_spreadsheets", None)
    st.session_state.pop("_gsheet_worksheets", None)
    st.session_state.pop("_gsheet_last_error", None)
    try:
        get_gspread_client.clear()
    except Exception:
        pass
def _run_network_with_timeout(fn, timeout=3.0):
    result = []
    error = []
    def worker():
        try:
            result.append(fn())
        except BaseException as exc:
            error.append(exc)
    timeout = max(1.0, float(timeout))
    if not _GSHEET_NETWORK_SEMAPHORE.acquire(timeout=0.25):
        raise TimeoutError("Google Sheets 요청이 이전 네트워크 작업으로 잠겨 있습니다. 잠시 후 다시 시도해 주세요.")
    def guarded_worker():
        try:
            worker()
        finally:
            _GSHEET_NETWORK_SEMAPHORE.release()
    thread = threading.Thread(target=guarded_worker, name="gsheet-timeout-worker", daemon=True)
    thread.start()
    thread.join(timeout=timeout)
    if thread.is_alive():
        raise TimeoutError(f"Google Sheets 응답 시간 초과 ({timeout:.0f}초)")
    if error:
        raise error[0]
    return result[0] if result else None
def get_worksheet(spreadsheet_id: str, sheet_name: str):
    try:
        cache = st.session_state.setdefault("_gsheet_spreadsheets", {})
        spreadsheet = cache.get(spreadsheet_id)
        if spreadsheet is None:
            client = get_gspread_client()
            spreadsheet = _run_network_with_timeout(
                lambda: client.open_by_key(spreadsheet_id), timeout=3
            )
            cache[spreadsheet_id] = spreadsheet
        ws_key = f"{spreadsheet_id}:{sheet_name}"
        ws_cache = st.session_state.setdefault("_gsheet_worksheets", {})
        if ws_key in ws_cache:
            return ws_cache[ws_key]
        ws = _run_network_with_timeout(
            lambda: spreadsheet.worksheet(sheet_name), timeout=3
        )
        ws_cache[ws_key] = ws
        return ws
    except WorksheetNotFound:
        try:
            cache = st.session_state.setdefault("_gsheet_spreadsheets", {})
            spreadsheet = cache.get(spreadsheet_id)
            if spreadsheet is None:
                client = get_gspread_client()
                spreadsheet = _run_network_with_timeout(
                    lambda: client.open_by_key(spreadsheet_id), timeout=3
                )
                cache[spreadsheet_id] = spreadsheet
            ws = _run_network_with_timeout(
                lambda: spreadsheet.add_worksheet(title=sheet_name, rows=2000, cols=40), timeout=3
            )
            st.session_state.setdefault("_gsheet_worksheets", {})[f"{spreadsheet_id}:{sheet_name}"] = ws
            return ws
        except Exception as e:
            st.session_state["_gsheet_last_error"] = str(e)
            return None
    except Exception as e:
        st.session_state["_gsheet_last_error"] = str(e)
        return None
def df_from_worksheet(ws):
    if ws is None:
        return pd.DataFrame()
    try:
        data = _run_network_with_timeout(lambda: ws.get_all_values(), timeout=3)
        if not data or len(data) < 2:
            return pd.DataFrame()
        headers = [str(h).strip() for h in data[0]]
        rows = []
        for row in data[1:]:
            row = list(row) + [""] * max(0, len(headers) - len(row))
            rows.append(["" if c is None else str(c).strip() for c in row[:len(headers)]])
        return pd.DataFrame(rows, columns=headers).replace({"nan": "", "None": "", "NaN": ""})
    except Exception as e:
        st.session_state["_gsheet_last_error"] = str(e)
        return pd.DataFrame()
def df_to_worksheet(ws, df):
    if ws is None:
        raise RuntimeError("Google Sheets 워크시트를 열 수 없습니다.")
    try:
        if df is None or df.empty:
            _run_network_with_timeout(lambda: ws.clear(), timeout=5)
            return True
        clean = df.fillna("").astype(str)
        clean = clean.apply(lambda col: col.map(_safe_sheet_text))
        values = [clean.columns.tolist()] + clean.values.tolist()
        _run_network_with_timeout(
            lambda: ws.update("A1", values, value_input_option="USER_ENTERED"), timeout=8
        )
        new_row_count = len(values)
        total_rows = int(getattr(ws, "row_count", new_row_count))
        total_cols = int(getattr(ws, "col_count", max(1, len(values[0]))))
        if new_row_count < total_rows:
            clear_range = f"A{new_row_count + 1}:{get_column_letter(total_cols)}{total_rows}"
            _run_network_with_timeout(lambda: ws.batch_clear([clear_range]), timeout=5)
        return True
    except Exception as exc:
        st.session_state["_gsheet_last_error"] = str(exc)
        raise RuntimeError(f"Google Sheets 저장 실패: {exc}") from exc
def _empty_id_df():
    return pd.DataFrame(columns=["아이디", "이름", "권한", "허용탭"])
def _seed_master_id_df():
    return pd.DataFrame([{
        "아이디": MASTER_ID,
        "이름": "관리자",
        "권한": ROLE_MASTER,
        "허용탭": ",".join(ALL_APP_TABS),
    }])
def load_id_sheet():
    ws = get_worksheet(WORK_SHEET_ID, "아이디저장함")
    if ws is None:
        detail = st.session_state.get("_gsheet_last_error", "알 수 없는 Google Sheets 오류")
        st.session_state["_id_sheet_error"] = (
            "아이디저장함을 열 수 없습니다. Google Sheets 서비스 계정의 공유 권한과 "
            f"시트 ID/네트워크를 확인하세요. ({detail})"
        )
        return _empty_id_df()
    df = df_from_worksheet(ws)
    if df.empty or "아이디" not in df.columns:
        seed = _seed_master_id_df()
        try:
            df_to_worksheet(ws, seed)
            df = seed
            st.session_state.pop("_id_sheet_error", None)
        except Exception as exc:
            st.session_state["_id_sheet_error"] = (
                "아이디저장함은 열렸지만 초기 관리자 계정을 저장하지 못했습니다. "
                f"Google Sheets 편집 권한을 확인하세요. ({exc})"
            )
            return _empty_id_df()
    for c in ["아이디", "이름", "권한", "허용탭"]:
        if c not in df.columns:
            df[c] = ""
    df["권한"] = df["권한"].replace("", ROLE_TEACHER)
    df["이름"] = df["이름"].fillna("").astype(str)
    df["허용탭"] = df["허용탭"].fillna("").astype(str)
    st.session_state.pop("_id_sheet_error", None)
    return df
def _validate_id_admin_edit(existing: pd.DataFrame, edited: pd.DataFrame):
    if not can_manage_ids():
        return False, "아이디 권한을 변경할 권한이 없습니다."
    if not isinstance(edited, pd.DataFrame):
        return False, "잘못된 아이디 데이터입니다."
    work = edited.copy(deep=True)
    for c in ["아이디", "이름", "권한", "허용탭"]:
        if c not in work.columns:
            work[c] = ""
    work["아이디"] = work["아이디"].map(lambda x: _clean_user_text(x, MAX_ID_LENGTH))
    work["이름"] = work["이름"].map(lambda x: _clean_user_text(x, MAX_NAME_LENGTH))
    work["권한"] = work["권한"].map(_validated_role)
    work["허용탭"] = work["허용탭"].map(lambda x: _clean_user_text(x, MAX_TAB_TEXT_LENGTH))
    if work["아이디"].eq("").any() or work["이름"].eq("").any():
        return False, "아이디와 이름은 비워둘 수 없습니다."
    if work["아이디"].duplicated().any():
        return False, "중복된 아이디가 있습니다."
    old_master = set(existing.loc[existing["권한"] == ROLE_MASTER, "아이디"].astype(str).str.strip()) if isinstance(existing, pd.DataFrame) and "권한" in existing.columns else set()
    new_master = set(work.loc[work["권한"] == ROLE_MASTER, "아이디"].astype(str).str.strip())
    if not new_master:
        return False, "마스터 권한 아이디가 최소 1명은 있어야 합니다."
    if current_role() == ROLE_EDU and old_master != new_master:
        return False, "교육과정부는 마스터 권한을 생성·양도·해제할 수 없습니다."
    return True, work
def save_id_sheet(df):
    if not can_manage_ids():
        raise PermissionError("아이디 목록을 변경할 권한이 없습니다.")
    ws = get_worksheet(WORK_SHEET_ID, "아이디저장함")
    existing = df_from_worksheet(ws)
    ok, result = _validate_id_admin_edit(existing, df)
    if not ok:
        raise PermissionError(result)
    df_to_worksheet(ws, result)
    st.session_state.pop("_auth_last_check", None)
    return True
def save_id_request(name, email, desired_id, memo):
    name = _clean_user_text(name, MAX_NAME_LENGTH)
    email = _clean_user_text(email, MAX_EMAIL_LENGTH)
    desired_id = _clean_user_text(desired_id, MAX_ID_LENGTH)
    memo = _clean_user_text(memo, MAX_MEMO_LENGTH)
    if not name or not email or not desired_id:
        raise ValueError("이름, 이메일, 아이디는 필수입니다.")
    if "@" not in email or len(email) > MAX_EMAIL_LENGTH:
        raise ValueError("이메일 형식이 올바르지 않습니다.")
    ws = get_worksheet(WORK_SHEET_ID, "아이디추가요청")
    df = df_from_worksheet(ws)
    new = pd.DataFrame([{
        "이름": name, "이메일": email, "추가아이디": desired_id, "메모": memo,
        "요청시각": _now_text(), "처리상태": "대기"
    }])
    if df.empty:
        df = new
    else:
        for c in new.columns:
            if c not in df.columns:
                df[c] = ""
        df = pd.concat([df, new], ignore_index=True)
    df_to_worksheet(ws, df)
def load_budget_df():
    ws = get_worksheet(WORK_SHEET_ID, "예산")
    df = df_from_worksheet(ws)
    expected_cols = ["일시", "내용", "변동금액", "잔액"]
    if df.empty or not any(c in df.columns for c in expected_cols + ["보강예산 현황"]):
        init_df = pd.DataFrame([{
            "일시": _now_text(),
            "내용": "초기 예산 설정",
            "변동금액": 2200000,
            "잔액": 2200000
        }])
        df_to_worksheet(ws, init_df)
        st.session_state["_budget_df_snapshot"] = init_df.copy(deep=True)
        return init_df
    if "보강예산 현황" in df.columns and "잔액" not in df.columns:
        try:
            val = safe_int(df.iloc[0, 0]) if not df.empty else 2200000
        except Exception:
            val = 2200000
        df = pd.DataFrame([{
            "일시": _now_text(),
            "내용": "기존 예산 불러오기",
            "변동금액": 0,
            "잔액": val
        }])
    for c in expected_cols:
        if c not in df.columns:
            df[c] = 0 if c in ("변동금액", "잔액") else ""
    st.session_state["_budget_df_snapshot"] = df.copy(deep=True)
    return df
def get_current_budget():
    try:
        df = st.session_state.get("_budget_df_snapshot")
        if not isinstance(df, pd.DataFrame):
            df = load_budget_df()
        if df.empty:
            return 2200000
        last_val = safe_int(df.iloc[-1].get("잔액", 2200000))
        return max(0, last_val)
    except Exception:
        return 2200000
def update_budget(change_amount: int, reason: str = "보강"):
    try:
        df = st.session_state.get("_budget_df_snapshot")
        if not isinstance(df, pd.DataFrame):
            df = load_budget_df()
        current = safe_int(df.iloc[-1].get("잔액", 2200000)) if not df.empty else 2200000
        new_balance = max(0, current + change_amount)
        new_row = pd.DataFrame([{
            "일시": _now_text(),
            "내용": reason,
            "변동금액": change_amount,
            "잔액": new_balance
        }])
        df = pd.concat([df, new_row], ignore_index=True)
        ws = get_worksheet(WORK_SHEET_ID, "예산")
        df_to_worksheet(ws, df)
        st.session_state["_budget_df_snapshot"] = df.copy(deep=True)
        return new_balance
    except Exception as e:
        st.warning(f"예산 업데이트 실패: {e}")
        return None
SWAP_REQUEST_COLS = [
    "신청ID", "신청자", "신청자이름", "원본일자", "교사A", "요일A", "교시A", "학급A", "과목A",
    "목표일자", "교사B", "요일B", "교시B", "학급B", "과목B", "유형", "신청시각", "상태"
]
def load_swap_requests():
    ws = get_worksheet(WORK_SHEET_ID, "수업교체신청")
    df = df_from_worksheet(ws)
    if df.empty:
        return pd.DataFrame(columns=SWAP_REQUEST_COLS)
    for c in SWAP_REQUEST_COLS:
        if c not in df.columns:
            df[c] = ""
    return df
def save_swap_request(rec: dict):
    df = load_swap_requests()
    new = pd.DataFrame([rec])
    df = pd.concat([df, new], ignore_index=True)
    ws = get_worksheet(WORK_SHEET_ID, "수업교체신청")
    df_to_worksheet(ws, df)
def push_history(action_name="작업"):
    if "history" not in st.session_state:
        st.session_state.history=[]; st.session_state.history_index=-1
    state = {
        "action": action_name,
        "time": _now_text(),
        "absences": st.session_state.get("absences", pd.DataFrame()).copy(deep=True),
        "subs": st.session_state.get("subs", pd.DataFrame()).copy(deep=True),
        "swaps": st.session_state.get("swaps", pd.DataFrame()).copy(deep=True),
        "part_time": st.session_state.get("part_time", pd.DataFrame()).copy(deep=True),
        "duties": st.session_state.get("duties", pd.DataFrame()).copy(deep=True),
        "budget_df": load_budget_df().copy(deep=True) if st.session_state.get("logged_in", False) else pd.DataFrame(),
    }
    def fingerprint(x):
        return tuple((k, tuple(v.fillna("").astype(str).tolist())) for k,v in sorted(x.items()) if isinstance(v,pd.DataFrame))
    if st.session_state.history:
        prev=st.session_state.history[st.session_state.history_index]
        if fingerprint(prev)==fingerprint(state):
            prev["action"]=action_name; prev["time"]=state["time"]; return
    st.session_state.history=st.session_state.history[:st.session_state.history_index+1]
    st.session_state.history.append(state)
    st.session_state.history_index=len(st.session_state.history)-1
    if len(st.session_state.history)>MAX_HISTORY:
        st.session_state.history.pop(0); st.session_state.history_index-=1
def _restore_history_snapshot(snap):
    for k in ["absences", "subs", "swaps", "part_time", "duties"]:
        st.session_state[k] = snap.get(k, pd.DataFrame()).copy(deep=True)
    budget_df = snap.get("budget_df")
    if isinstance(budget_df, pd.DataFrame) and not budget_df.empty:
        budget_df = budget_df.copy(deep=True)
        df_to_worksheet(get_worksheet(WORK_SHEET_ID, "예산"), budget_df)
        st.session_state["_budget_df_snapshot"] = budget_df
        _invalidate_all_caches()
def undo():
    if st.session_state.get("history_index", 0) <= 0:
        return False
    st.session_state.history_index -= 1
    _restore_history_snapshot(st.session_state.history[st.session_state.history_index])
    return True
def redo():
    if st.session_state.get("history_index", -1) >= len(st.session_state.get("history", [])) - 1:
        return False
    st.session_state.history_index += 1
    _restore_history_snapshot(st.session_state.history[st.session_state.history_index])
    return True
def _invalidate_all_caches():
    st.session_state._data_version = st.session_state.get("_data_version", 0) + 1
    get_effective_timetable_for_date.clear()
    get_single_lesson_1to1_candidates.clear()
    get_single_lesson_linked_cycles.clear()
    effective_teacher_matrix.clear()
    teacher_matrix.clear()
    class_matrix.clear()
    cumulative_sub_count.clear()
    weekly_load.clear()
    _duty_slot_index.clear()
    get_teacher_availability_index.clear()
    get_test_affected_slots.clear()
    get_actual_direct_swap_affected_slots.clear()
    find_cycle_linked_swaps.clear()
    get_target_time_recommendations.clear()
DUTY_COLS = ["교사명", "일자", "교시", "사유", "상세사유", "등록시각", "입력자"]
PART_TIME_EXTRA_COLS = ["시작일", "종료일", "대체교사"]
def ensure_duty_columns(df):
    if df is None or not isinstance(df, pd.DataFrame):
        return pd.DataFrame(columns=DUTY_COLS)
    df = df.copy()
    for c in DUTY_COLS:
        if c not in df.columns:
            df[c] = 0 if c == "교시" else ""
    df["교시"] = df["교시"].apply(safe_int)
    return df
def ensure_part_time_columns(df):
    base_cols = ["번호", "시간강사명", "담당과목", "과목군", "비고"] + PART_TIME_EXTRA_COLS + [f"{d}{p}" for d in DAYS for p in range(1, 8)]
    if df is None or not isinstance(df, pd.DataFrame):
        return pd.DataFrame(columns=base_cols)
    df = df.copy()
    for c in base_cols:
        if c not in df.columns:
            df[c] = ""
    for c in ["시작일", "종료일"]:
        if c in df.columns:
            df[c] = df[c].apply(normalize_date_str)
    return df
def ensure_input_user(df, default=""):
    if df is None or not isinstance(df, pd.DataFrame):
        return df
    if "입력자" not in df.columns:
        df = df.copy()
        df["입력자"] = default
    return df
def load_timetable_from_gsheet():
    try:
        ti = df_from_worksheet(get_worksheet(TIMETABLE_SHEET_ID, "교사정보"))
        tt = df_from_worksheet(get_worksheet(TIMETABLE_SHEET_ID, "시간표"))
        if tt.empty:
            raise RuntimeError("'시간표' 시트에서 데이터를 읽지 못했습니다.")
        required = ["교사명", "요일", "교시", "과목", "학급"]
        missing = [c for c in required if c not in tt.columns]
        if missing:
            raise RuntimeError("'시간표' 시트에 필요한 열이 없습니다: " + ", ".join(missing))
        tt["교시"] = tt["교시"].apply(safe_int)
        for c in ["교사명", "요일", "과목", "학급"]:
            tt[c] = tt[c].astype(str).str.strip()
        if "과목군" not in tt.columns or tt["과목군"].eq("").all():
            tt["과목군"] = tt["과목"].map(subject_group)
        tt = tt[tt["요일"].isin(DAYS) & (tt["교시"] >= 1) & (tt["교시"] <= MAX_PERIOD)]
        tt = tt.drop_duplicates(subset=["교사명", "요일", "교시"]).reset_index(drop=True)
        st.session_state.pop("_gsheet_last_error", None)
        return ti, tt
    except Exception as e:
        st.session_state["_gsheet_last_error"] = str(e)
        return pd.DataFrame(), pd.DataFrame()
def load_work_data_from_gsheet():
    try:
        absences = ensure_input_user(df_from_worksheet(get_worksheet(WORK_SHEET_ID, "결강")))
        subs = ensure_input_user(df_from_worksheet(get_worksheet(WORK_SHEET_ID, "보강")))
        swaps = ensure_input_user(df_from_worksheet(get_worksheet(WORK_SHEET_ID, "맞교환")))
        part_time = ensure_part_time_columns(df_from_worksheet(get_worksheet(WORK_SHEET_ID, "시간강사")))
        cumulative = df_from_worksheet(get_worksheet(WORK_SHEET_ID, "누적보강"))
        duties = ensure_duty_columns(df_from_worksheet(get_worksheet(WORK_SHEET_ID, "복무")))
        for df in [absences, subs]:
            if not df.empty:
                if "교시" in df.columns:
                    df["교시"] = df["교시"].apply(safe_int)
                if "일자" in df.columns:
                    df["일자"] = df["일자"].apply(normalize_date_str)
        if not swaps.empty:
            for c in ["원본일자", "목표일자"]:
                if c in swaps.columns:
                    swaps[c] = swaps[c].apply(normalize_date_str)
            for c in ["교시A", "교시B"]:
                if c in swaps.columns:
                    swaps[c] = swaps[c].apply(safe_int)
        if not duties.empty and "일자" in duties.columns:
            duties["일자"] = duties["일자"].apply(normalize_date_str)
        return absences, subs, swaps, part_time, cumulative, duties
    except Exception as e:
        st.session_state["_gsheet_last_error"] = str(e)
        return (pd.DataFrame(),) * 5 + (pd.DataFrame(columns=DUTY_COLS),)
def save_work_data_to_gsheet(changed_sheets=None):
    if not can_full_data() and not is_teacher():
        st.warning("저장 권한이 없습니다.")
        return False
    try:
        if changed_sheets is None:
            changed_sheets = ["결강", "보강", "맞교환", "시간강사", "복무"]
        changed_sheets = set(changed_sheets)
        if "결강" in changed_sheets:
            df_to_worksheet(get_worksheet(WORK_SHEET_ID, "결강"), st.session_state.absences)
        if "보강" in changed_sheets:
            df_to_worksheet(get_worksheet(WORK_SHEET_ID, "보강"), st.session_state.subs)
        if "맞교환" in changed_sheets:
            df_to_worksheet(get_worksheet(WORK_SHEET_ID, "맞교환"), st.session_state.swaps)
        if "시간강사" in changed_sheets:
            st.session_state.part_time = ensure_part_time_columns(st.session_state.part_time)
            df_to_worksheet(get_worksheet(WORK_SHEET_ID, "시간강사"), st.session_state.part_time)
        if "복무" in changed_sheets:
            st.session_state.duties = ensure_duty_columns(st.session_state.duties)
            df_to_worksheet(get_worksheet(WORK_SHEET_ID, "복무"), st.session_state.duties)
        _invalidate_all_caches()
        return True
    except Exception as e:
        st.error(f"저장 실패: {e}")
        return False
def init_state():
    existing_tt = st.session_state.get("timetable")
    existing_ti = st.session_state.get("teachers")
    if isinstance(existing_tt, pd.DataFrame) and not existing_tt.empty:
        return True
    ti, tt = load_timetable_from_gsheet()
    if tt is None or tt.empty:
        st.session_state.pop("teachers", None)
        st.session_state.pop("timetable", None)
        detail = st.session_state.get("_gsheet_last_error", "")
        st.session_state["_load_error"] = (
            "시간표 시트에서 유효한 데이터를 받지 못했습니다. "
            "Google Sheets 인증/공유 권한, 시트 이름(시간표), 네트워크 상태를 확인하세요."
            + (f" 최근 오류: {detail}" if detail else "")
        )
        return False
    st.session_state.teachers = ti if isinstance(ti, pd.DataFrame) else pd.DataFrame()
    st.session_state.timetable = tt
    st.session_state.pop("_load_error", None)
    absences, subs, swaps, part_time, cumulative, duties = load_work_data_from_gsheet()
    if absences.empty:
        absences = pd.DataFrame(columns=["결강ID", "일자", "요일", "교사명", "사유", "상세사유", "교시", "학급", "과목", "등록시각", "입력자"])
    if subs.empty:
        subs = pd.DataFrame(columns=["결강ID", "일자", "요일", "교시", "학급", "과목", "결강교사", "보강교사", "배정방식", "우선순위", "비고", "등록시각", "입력자"])
    if swaps.empty:
        swaps = pd.DataFrame(columns=["원본일자", "교사A", "요일A", "교시A", "학급A", "과목A",
                                      "목표일자", "교사B", "요일B", "교시B", "학급B", "과목B", "유형", "시간강사구인", "등록시각", "입력자"])
    if part_time.empty:
        part_time = pd.DataFrame(columns=["번호", "시간강사명", "담당과목", "과목군", "비고", "시작일", "종료일", "대체교사"] + [f"{d}{p}" for d in DAYS for p in range(1, 8)])
    st.session_state.absences = ensure_input_user(absences)
    st.session_state.subs = ensure_input_user(subs)
    st.session_state.swaps = ensure_input_user(swaps)
    st.session_state.part_time = ensure_part_time_columns(part_time)
    st.session_state.cumulative = cumulative
    st.session_state.duties = ensure_duty_columns(duties)
    st.session_state._data_version = 0
    st.session_state.history = []
    st.session_state.history_index = -1
    st.session_state.test_swaps = pd.DataFrame()
    push_history("초기 상태")
    return True
CHANGE_TYPES = {"원본", "교환", "테스트교환", "보강", "시간강사"}
def _new_change_id(prefix="CHG"):
    return f"{prefix}-{datetime.now().strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:8]}"
def _lesson_record(teacher, day, period, subject, class_name, *, orig_teacher=None,
                   orig_date="", orig_period=0, change_type="원본", change_source="",
                   change_id="", change_detail=""):
    return {
        "교사명": str(teacher).strip(), "요일": day, "교시": safe_int(period),
        "과목": str(subject).strip(), "학급": str(class_name).strip(),
        "과목군": subject_group(str(subject).strip()),
        "원본교사": str(orig_teacher if orig_teacher is not None else teacher).strip(),
        "원본일자": normalize_date_str(orig_date), "원본교시": safe_int(orig_period or period),
        "변경유형": change_type, "변경출처": str(change_source).strip(),
        "변경ID": str(change_id).strip(), "변경상세": str(change_detail).strip(),
    }
def _truthy_availability(value):
    if value is None or pd.isna(value):
        return False
    s = str(value).strip().lower()
    return s not in {"", "nan", "none", "0", "x", "n", "no", "불가", "아니오", "×"}
def _part_time_available(prow, day, period):
    col = f"{day}{safe_int(period)}"
    return col in prow.index and _truthy_availability(prow.get(col, ""))

_TEACHER_SLOT_COL_RE = re.compile(r"^(월|화|수|목|금)\s*(?:요일)?\s*[-_/]?\s*(\d{1,2})\s*(?:교시)?$")
_TEACHER_DAY_COL_RE = re.compile(r"^(월|화|수|목|금)(?:요일)?(?:\s|_|-|/)*(?:가능|가능시간|가능시간대|교시|시간)?$")
_TEACHER_AVAIL_TEXT_KEYS = (
    "가능시간", "가능 시간", "가능시간대", "가능 시간대", "가능교시", "가능 교시",
    "교사 가능시간", "교사 가능 시간", "가용시간", "가용 시간", "가용교시", "가용 교시"
)

def _normalize_teacher_col(value):
    return re.sub(r"\s+", "", str(value or "").strip())

def _parse_period_tokens(text):
    result=set()
    for m in re.finditer(r"(\d{1,2})\s*(?:~|\-|–|—)\s*(\d{1,2})", str(text)):
        a,b=safe_int(m.group(1)),safe_int(m.group(2))
        if 1 <= a <= 7 and 1 <= b <= 7:
            lo,hi=sorted((a,b)); result.update(range(lo,hi+1))
    for m in re.finditer(r"(?<!\d)([1-7])(?!\d)", str(text)):
        result.add(safe_int(m.group(1)))
    return result

def _parse_teacher_availability_text(text):
    text=str(text or "").strip()
    if not text:
        return set()
    if text.lower() in {"전체", "모두", "전부", "매일", "전일", "all"}:
        return {(d,p) for d in DAYS for p in range(1, PERIODS_PER_DAY.get(d, MAX_PERIOD)+1)}
    result=set()
    day_pattern=r"(월|화|수|목|금)(?:요일)?"
    matches=list(re.finditer(day_pattern, text))
    if matches:
        for i,m in enumerate(matches):
            day=m.group(1)
            body=text[m.end():matches[i+1].start() if i+1<len(matches) else len(text)]
            periods=_parse_period_tokens(body)
            for p in periods:
                if p <= PERIODS_PER_DAY.get(day, MAX_PERIOD):
                    result.add((day,p))
        return result
    # 요일 정보가 없는 단독 숫자 표기는 호출하는 쪽의 요일 문맥에서 해석한다.
    return set()

@st.cache_data(show_spinner=False, ttl=300)
def get_teacher_availability_index(version=0):
    ti=st.session_state.get("teachers", pd.DataFrame())
    if not isinstance(ti, pd.DataFrame) or ti.empty or "교사명" not in ti.columns:
        return {}
    columns={_normalize_teacher_col(c): c for c in ti.columns}
    slot_cols={}
    day_cols={}
    for norm_col, original_col in columns.items():
        m=_TEACHER_SLOT_COL_RE.fullmatch(norm_col)
        if m:
            day,p=m.group(1), safe_int(m.group(2))
            if 1 <= p <= MAX_PERIOD:
                slot_cols[(day,p)]=original_col
            continue
        m=_TEACHER_DAY_COL_RE.fullmatch(norm_col)
        if m and norm_col[:1] in DAYS:
            day_cols[m.group(1)]=original_col
    text_cols=[]
    for key in _TEACHER_AVAIL_TEXT_KEYS:
        nk=_normalize_teacher_col(key)
        if nk in columns:
            text_cols.append(columns[nk])
    # '가능시간'을 포함한 비표준 열 이름도 지원
    for original_col in ti.columns:
        norm_col=_normalize_teacher_col(original_col)
        if any(k.replace(" ","") in norm_col for k in ("가능시간","가용시간")) and original_col not in text_cols:
            text_cols.append(original_col)
    index={}
    for row in ti.to_dict("records"):
        teacher=str(row.get("교사명", "")).strip()
        if not teacher:
            continue
        configured=False
        allowed=set()
        if slot_cols:
            configured=True
            for slot,col in slot_cols.items():
                if _truthy_availability(row.get(col, "")):
                    allowed.add(slot)
        else:
            # 요일별 열: '월=1,2,4' / '화=3~5' 같은 입력을 지원
            for day,col in day_cols.items():
                raw=row.get(col, "")
                if str(raw).strip():
                    configured=True
                    for p in _parse_period_tokens(raw):
                        if p <= PERIODS_PER_DAY.get(day, MAX_PERIOD):
                            allowed.add((day,p))
            # 통합 가능시간 열: '월 1,2 / 화 4~6' 지원
            for col in text_cols:
                raw=row.get(col, "")
                if str(raw).strip():
                    configured=True
                    allowed.update(_parse_teacher_availability_text(raw))
        index[teacher]=(configured, frozenset(allowed))
    return index

def teacher_slot_is_available(teacher: str, day: str, period: int, version=0):
    teacher=str(teacher or "").strip(); day=str(day or "").strip(); p=safe_int(period)
    if not teacher or day not in DAYS or p <= 0:
        return False
    item=get_teacher_availability_index(version).get(teacher)
    if item is None:
        return True
    configured, allowed=item
    return (day,p) in allowed if configured else True
def validate_part_time_table(df):
    if df is None or df.empty: return True, ""
    seen = {}
    for _, r in df.iterrows():
        name=str(r.get("시간강사명","")).strip(); orig=str(r.get("대체교사","")).strip()
        start,end=normalize_date_str(r.get("시작일","")),normalize_date_str(r.get("종료일",""))
        if not name or not orig or not start or not end: continue
        if start>end: return False, f"시간강사 {name}: 시작일이 종료일보다 늦습니다."
        for d in DAYS:
            for p in range(1,8):
                if _part_time_available(r,d,p):
                    key=(orig,d,p)
                    if key in seen: return False, f"{orig}의 {d}{p}교시 시간강사 대체가 중복됩니다."
                    seen[key]=name
    return True, ""
def _effective_index(e_tt):
    idx = {}
    if e_tt is None or e_tt.empty:
        return idx
    for r in e_tt.itertuples(index=False):
        idx[(str(r.교사명).strip(), safe_int(r.교시))] = r
    return idx
def _class_slot_teachers(e_tt, class_name, period):
    if e_tt is None or e_tt.empty:
        return []
    m = e_tt[(e_tt["학급"].astype(str).str.strip() == str(class_name).strip()) & (e_tt["교시"].apply(safe_int) == safe_int(period))]
    return [str(x).strip() for x in m["교사명"].tolist()]
def validate_swap(a, b, date_a, date_b, *, is_test=False):
    da, db = normalize_date_str(date_a), normalize_date_str(date_b)
    ta, tb = str(a.get("교사명", "")).strip(), str(b.get("교사명", "")).strip()
    pa, pb = safe_int(a.get("교시", 0)), safe_int(b.get("교시", 0))
    if not da or not db or not ta or not tb or pa <= 0 or pb <= 0:
        return False, "교사·일자·교시 정보가 올바르지 않습니다."
    if ta == tb and da == db and pa == pb:
        return False, "동일한 교사·일자·교시는 교환할 수 없습니다."
    ver = st.session_state.get("_data_version", 0)
    e_a = get_effective_timetable_for_date(da, ver, use_test=is_test)
    e_b = get_effective_timetable_for_date(db, ver, use_test=is_test)
    ma = e_a[(e_a["교사명"] == ta) & (e_a["교시"].apply(safe_int) == pa)] if not e_a.empty else pd.DataFrame()
    mb = e_b[(e_b["교사명"] == tb) & (e_b["교시"].apply(safe_int) == pb)] if not e_b.empty else pd.DataFrame()
    if ma.empty:
        return False, f"현재 적용 시간표에서 {ta}의 {da} {pa}교시 수업을 찾을 수 없습니다."
    if mb.empty:
        return False, f"현재 적용 시간표에서 {tb}의 {db} {pb}교시 수업을 찾을 수 없습니다."
    for label, info, m in [("A", a, ma), ("B", b, mb)]:
        cls, subj = str(info.get("학급", "")).strip(), str(info.get("과목", "")).strip()
        if cls and str(m.iloc[0]["학급"]).strip() != cls:
            return False, f"{label} 수업의 학급이 현재 시간표와 달라졌습니다. 다시 검색하세요."
        if subj and str(m.iloc[0]["과목"]).strip() != subj:
            return False, f"{label} 수업의 과목이 현재 시간표와 달라졌습니다. 다시 검색하세요."
    if db != da or pb != pa:
        a_at_target = e_b[(e_b["교사명"] == ta) & (e_b["교시"].apply(safe_int) == pb)] if not e_b.empty else pd.DataFrame()
        b_at_target = e_a[(e_a["교사명"] == tb) & (e_a["교시"].apply(safe_int) == pa)] if not e_a.empty else pd.DataFrame()
        if not a_at_target.empty:
            return False, f"A 교사의 목표 슬롯 {db} {pb}교시에 이미 수업이 있습니다."
        if not b_at_target.empty:
            return False, f"B 교사의 목표 슬롯 {da} {pa}교시에 이미 수업이 있습니다."
    if not is_test:
        used = get_actual_direct_swap_affected_slots(st.session_state.get("_data_version", 0))
        if (da, ta, pa) in used or (db, tb, pb) in used:
            return False, "이미 다른 맞교환에 사용된 슬롯입니다."
    else:
        if (da, ta, pa) in get_test_affected_slots(st.session_state.get("_data_version", 0)) or (db, tb, pb) in get_test_affected_slots(st.session_state.get("_data_version", 0)):
            return False, "테스트에서 이미 사용된 슬롯입니다."
    day_a = WEEKDAY_KR[datetime.strptime(da, "%Y-%m-%d").weekday()]
    day_b = WEEKDAY_KR[datetime.strptime(db, "%Y-%m-%d").weekday()]
    if not teacher_slot_is_available(ta, day_b, pb, ver):
        return False, f"{ta} 교사는 {day_b}{pb}교시를 가능 시간으로 등록하지 않았습니다."
    if not teacher_slot_is_available(tb, day_a, pa, ver):
        return False, f"{tb} 교사는 {day_a}{pa}교시를 가능 시간으로 등록하지 않았습니다."
    if has_duty(ta, da, pa) or has_duty(tb, db, pb):
        return False, "복무가 등록된 슬롯은 맞교환할 수 없습니다."
    return True, ""
def validate_substitute(cid, on_date, period, class_name, subject, absent_teacher, sub_teacher, *, e_tt=None):
    norm = normalize_date_str(on_date); p = safe_int(period)
    absent_teacher, sub_teacher = str(absent_teacher).strip(), str(sub_teacher).strip()
    if not norm or p <= 0 or not absent_teacher or not sub_teacher:
        return False, "보강에 필요한 정보가 부족합니다."
    if absent_teacher == sub_teacher:
        return False, "결강교사와 보강교사가 같을 수 없습니다."
    ver = st.session_state.get("_data_version", 0)
    if e_tt is None:
        e_tt = get_effective_timetable_for_date(norm, ver)
    source = e_tt[(e_tt["교사명"] == absent_teacher) & (e_tt["교시"].apply(safe_int) == p)] if not e_tt.empty else pd.DataFrame()
    if source.empty:
        base = st.session_state.timetable
        source = base[(base["교사명"] == absent_teacher) & (base["요일"] == WEEKDAY_KR[datetime.strptime(norm, "%Y-%m-%d").weekday()]) & (base["교시"].apply(safe_int) == p)] if not base.empty else pd.DataFrame()
    if source.empty:
        return False, f"{absent_teacher}의 해당 교시 원본 수업을 찾을 수 없습니다."
    if str(source.iloc[0].get("학급", "")).strip() != str(class_name).strip() or str(source.iloc[0].get("과목", "")).strip() != str(subject).strip():
        return False, "결강 수업 정보가 현재 시간표와 일치하지 않습니다."
    day = WEEKDAY_KR[datetime.strptime(norm, "%Y-%m-%d").weekday()]
    if not is_free(sub_teacher, day, p, norm, e_tt):
        return False, f"{sub_teacher} 교사는 {day}{p}교시에 공강이 아닙니다."
    class_teachers = _class_slot_teachers(e_tt, class_name, p)
    others = [t for t in class_teachers if t != absent_teacher]
    if others:
        return False, f"{class_name}은 {p}교시에 이미 다른 교사({', '.join(others)})가 담당하고 있습니다."
    subs = st.session_state.get("subs", pd.DataFrame())
    if not subs.empty:
        same = subs[(subs["일자"].map(normalize_date_str) == norm) & (subs["교시"].apply(safe_int) == p)]
        same = same[(same["결강교사"].astype(str).str.strip() == absent_teacher) | (same["보강교사"].astype(str).str.strip() == sub_teacher)]
        if not same.empty:
            return False, "같은 시간대에 중복 보강 배정이 있습니다."
    return True, ""
def _apply_sub_to_effective(e_tt, class_name, subject, absent_teacher, sub_teacher, period, norm, change_id=""):
    if e_tt is None:
        return pd.DataFrame()
    current = [dict(x) for x in e_tt.to_dict("records")]
    current = [x for x in current if not (str(x.get("교사명", "")).strip() == absent_teacher and safe_int(x.get("교시")) == safe_int(period))]
    day = WEEKDAY_KR[datetime.strptime(norm, "%Y-%m-%d").weekday()]
    current.append(_lesson_record(sub_teacher, day, period, subject, class_name,
                                  orig_teacher=absent_teacher, orig_date=norm, orig_period=period,
                                  change_type="보강", change_source=f"{absent_teacher} 결강 → {sub_teacher} 보강",
                                  change_id=change_id or _new_change_id("SUB"),
                                  change_detail=f"{class_name} {subject}"))
    return pd.DataFrame(current)
def _is_direct_swap_type(typ: str) -> bool:
    return str(typ).strip() in ["1:1 맞교환", "1:1맞교환", "직접1:1"]
@st.cache_data(show_spinner=False, ttl=180)
def get_test_affected_slots(version=0) -> set:
    affected = set()
    test_swaps = st.session_state.get("test_swaps", pd.DataFrame())
    if test_swaps is None or test_swaps.empty:
        return affected
    for sw in test_swaps.itertuples(index=False):
        date_a = normalize_date_str(getattr(sw, "원본일자", ""))
        date_b = normalize_date_str(getattr(sw, "목표일자", ""))
        t_a = str(getattr(sw, "교사A", "")).strip()
        t_b = str(getattr(sw, "교사B", "")).strip()
        p_a = safe_int(getattr(sw, "교시A", 0))
        p_b = safe_int(getattr(sw, "교시B", 0))
        typ = str(getattr(sw, "유형", "")).strip()
        if date_a and t_a and p_a > 0:
            affected.add((date_a, t_a, p_a))
        if date_b and t_b and p_b > 0:
            affected.add((date_b, t_b, p_b))
        if _is_direct_swap_type(typ):
            if date_a and t_b and p_a > 0:
                affected.add((date_a, t_b, p_a))
            if date_b and t_a and p_b > 0:
                affected.add((date_b, t_a, p_b))
        elif "연계" in typ:
            if date_b and t_a and p_b > 0:
                affected.add((date_b, t_a, p_b))
    return affected
def _test_slot_is_affected(on_date: str, teacher: str, period: int) -> bool:
    norm = normalize_date_str(on_date)
    return (norm, str(teacher).strip(), safe_int(period)) in get_test_affected_slots(st.session_state.get("_data_version", 0))
@st.cache_data(show_spinner=False, ttl=180)
def get_actual_direct_swap_affected_slots(version=0) -> set:
    affected = set()
    swaps = st.session_state.get("swaps", pd.DataFrame())
    if swaps is None or swaps.empty:
        return affected
    for sw in swaps.itertuples(index=False):
        if not _is_direct_swap_type(str(getattr(sw, "유형", ""))):
            continue
        date_a = normalize_date_str(getattr(sw, "원본일자", ""))
        date_b = normalize_date_str(getattr(sw, "목표일자", ""))
        t_a = str(getattr(sw, "교사A", "")).strip()
        t_b = str(getattr(sw, "교사B", "")).strip()
        p_a = safe_int(getattr(sw, "교시A", 0))
        p_b = safe_int(getattr(sw, "교시B", 0))
        for item in [
            (date_a, t_a, p_a), (date_a, t_b, p_a),
            (date_b, t_b, p_b), (date_b, t_a, p_b),
        ]:
            if item[0] and item[1] and item[2] > 0:
                affected.add(item)
    return affected
def _actual_direct_slot_is_affected(on_date: str, teacher: str, period: int) -> bool:
    norm = normalize_date_str(on_date)
    return (norm, str(teacher).strip(), safe_int(period)) in get_actual_direct_swap_affected_slots(st.session_state.get("_data_version", 0))
def _effective_swap_origin_info(teacher: str, on_date: str, period: int, use_test: bool = False) -> str:
    norm = normalize_date_str(on_date)
    teacher = str(teacher).strip()
    period = safe_int(period)
    if not norm or not teacher or period <= 0:
        return ""
    tables = []
    actual = st.session_state.get("swaps", pd.DataFrame())
    if actual is not None and not actual.empty:
        tables.append(actual)
    if use_test:
        test = st.session_state.get("test_swaps", pd.DataFrame())
        if test is not None and not test.empty:
            tables.append(test)
    for swaps in tables:
        for sw in swaps.itertuples(index=False):
            typ = str(getattr(sw, "유형", "")).strip()
            date_a = normalize_date_str(getattr(sw, "원본일자", ""))
            date_b = normalize_date_str(getattr(sw, "목표일자", ""))
            teacher_a = str(getattr(sw, "교사A", "")).strip()
            teacher_b = str(getattr(sw, "교사B", "")).strip()
            period_a = safe_int(getattr(sw, "교시A", 0))
            period_b = safe_int(getattr(sw, "교시B", 0))
            if _is_direct_swap_type(typ):
                if date_a == norm and teacher_b == teacher and period_a == period:
                    return f"{date_b} {period_b}교시의 {teacher_a} 수업과 맞교환"
                if date_b == norm and teacher_a == teacher and period_b == period:
                    return f"{date_a} {period_a}교시의 {teacher_b} 수업과 맞교환"
            elif "연계" in typ:
                if date_b == norm and teacher_a == teacher and period_b == period:
                    return f"{date_a} {period_a}교시의 {teacher_b}와 연계교환"
    return ""
def _effective_sub_origin_info(teacher: str, on_date: str, period: int) -> str:
    norm = normalize_date_str(on_date)
    teacher = str(teacher).strip()
    period = safe_int(period)
    if not norm or not teacher or period <= 0:
        return ""
    subs = st.session_state.get("subs", pd.DataFrame())
    if subs is None or subs.empty:
        return ""
    mask = (
        (subs["일자"].astype(str).map(normalize_date_str) == norm)
        & (subs["교시"].apply(safe_int) == period)
        & (subs["보강교사"].astype(str).str.strip() == teacher)
    )
    matches = subs[mask]
    if matches.empty:
        return ""
    row = matches.iloc[-1]
    absent_teacher = str(row.get("결강교사", "")).strip()
    method = str(row.get("배정방식", "")).strip()
    priority = str(row.get("우선순위", "")).strip()
    memo = str(row.get("비고", "")).strip()
    parts = []
    if absent_teacher:
        parts.append(f"{absent_teacher} 결강 → {teacher} 보강")
    else:
        parts.append(f"{teacher} 보강 배정")
    if method:
        parts.append(f"배정방식: {method}")
    if priority:
        parts.append(f"우선순위: {priority}")
    if memo:
        parts.append(f"비고: {memo}")
    return " / ".join(parts)
@st.cache_data(show_spinner=False, ttl=180)
def get_effective_timetable_for_date(on_date: str, version: int = 0, use_test: bool = False) -> pd.DataFrame:
    norm = normalize_date_str(on_date)
    columns = ["교사명","요일","교시","과목","학급","과목군","원본교사","원본일자","원본교시","변경유형","변경출처","변경ID","변경상세"]
    if not norm:
        base = st.session_state.timetable.copy()
        return base.assign(**{c: "" for c in columns if c not in base.columns})
    try:
        day = WEEKDAY_KR[datetime.strptime(norm, "%Y-%m-%d").weekday()]
    except Exception:
        return st.session_state.timetable.copy()
    tt = st.session_state.timetable
    if tt.empty:
        return pd.DataFrame(columns=columns)
    current = {}
    base = tt[tt["요일"] == day]
    for r in base.itertuples(index=False):
        p = safe_int(r.교시); t = str(r.교사명).strip()
        current[(t,p)] = _lesson_record(t, day, p, getattr(r,"과목", ""), getattr(r,"학급", ""),
                                         orig_teacher=t, orig_date=norm, orig_period=p, change_type="원본")
    def apply_swap_table(swaps, is_test_table=False):
        nonlocal current
        if swaps is None or swaps.empty:
            return
        mask = (swaps["원본일자"] == norm) | (swaps["목표일자"] == norm)
        for sw in swaps[mask].itertuples(index=False):
            ta, tb = str(getattr(sw,"교사A","")).strip(), str(getattr(sw,"교사B","")).strip()
            da, db = normalize_date_str(getattr(sw,"원본일자","")), normalize_date_str(getattr(sw,"목표일자",""))
            pa, pb = safe_int(getattr(sw,"교시A",0)), safe_int(getattr(sw,"교시B",0))
            typ = str(getattr(sw,"유형","")).strip(); cid = str(getattr(sw,"변경ID","")).strip() or ""
            sa, ca = str(getattr(sw,"과목A","")).strip(), str(getattr(sw,"학급A","")).strip()
            sb, cb = str(getattr(sw,"과목B","")).strip(), str(getattr(sw,"학급B","")).strip()
            if _is_direct_swap_type(typ):
                if da == db and pa == pb and da == norm:
                    current.pop((ta,pa), None); current.pop((tb,pb), None)
                    current[(tb,pa)] = _lesson_record(tb, day, pa, sa, ca,
                        orig_teacher=str(getattr(sw,"원본교사A",ta)).strip() or ta, orig_date=normalize_date_str(getattr(sw,"실제원본일자A",da)) or da, orig_period=safe_int(getattr(sw,"실제원본교시A",pa)) or pa,
                        change_type="테스트교환" if is_test_table else "교환", change_source=f"{ta} ↔ {tb}", change_id=cid,
                        change_detail=f"{da} {pa}교시의 {ta} 수업을 {tb}가 담당")
                    current[(ta,pb)] = _lesson_record(ta, day, pb, sb, cb,
                        orig_teacher=str(getattr(sw,"원본교사B",tb)).strip() or tb, orig_date=normalize_date_str(getattr(sw,"실제원본일자B",db)) or db, orig_period=safe_int(getattr(sw,"실제원본교시B",pb)) or pb,
                        change_type="테스트교환" if is_test_table else "교환", change_source=f"{ta} ↔ {tb}", change_id=cid,
                        change_detail=f"{db} {pb}교시의 {tb} 수업을 {ta}가 담당")
                else:
                    if da == norm:
                        current.pop((ta,pa), None)
                        if tb:
                            current[(tb,pa)] = _lesson_record(tb, day, pa, sa, ca,
                                orig_teacher=str(getattr(sw,"원본교사A",ta)).strip() or ta, orig_date=normalize_date_str(getattr(sw,"실제원본일자A",da)) or da, orig_period=safe_int(getattr(sw,"실제원본교시A",pa)) or pa,
                                change_type="테스트교환" if is_test_table else "교환", change_source=f"{ta} ↔ {tb}", change_id=cid,
                                change_detail=f"{da} {pa}교시의 {ta} 수업을 {tb}가 담당")
                    if db == norm:
                        current.pop((tb,pb), None)
                        if ta:
                            current[(ta,pb)] = _lesson_record(ta, day, pb, sb, cb,
                                orig_teacher=str(getattr(sw,"원본교사B",tb)).strip() or tb, orig_date=normalize_date_str(getattr(sw,"실제원본일자B",db)) or db, orig_period=safe_int(getattr(sw,"실제원본교시B",pb)) or pb,
                                change_type="테스트교환" if is_test_table else "교환", change_source=f"{ta} ↔ {tb}", change_id=cid,
                                change_detail=f"{db} {pb}교시의 {tb} 수업을 {ta}가 담당")
            elif "연계" in typ and db == norm and ta:
                current.pop((tb,pb), None)
                current[(ta,pb)] = _lesson_record(ta, day, pb, sb, cb,
                    orig_teacher=ta, orig_date=da, orig_period=pa,
                    change_type="테스트교환" if is_test_table else "교환",
                    change_source=f"{ta} → {tb}", change_id=cid,
                    change_detail=f"{da} {pa}교시 수업의 연계 이동")
                if da == norm:
                    current.pop((ta,pa), None)
    apply_swap_table(st.session_state.get("swaps", pd.DataFrame()), False)
    if use_test:
        apply_swap_table(st.session_state.get("test_swaps", pd.DataFrame()), True)
    subs = st.session_state.get("subs", pd.DataFrame())
    if subs is not None and not subs.empty:
        for r in subs[subs["일자"] == norm].itertuples(index=False):
            p = safe_int(getattr(r,"교시",0)); abs_t = str(getattr(r,"결강교사","")).strip(); sub_t = str(getattr(r,"보강교사","")).strip()
            if p <= 0 or not sub_t: continue
            current.pop((abs_t,p), None)
            current.pop((sub_t,p), None)
            current[(sub_t,p)] = _lesson_record(sub_t, day, p, getattr(r,"과목",""), getattr(r,"학급",""),
                orig_teacher=abs_t, orig_date=norm, orig_period=p, change_type="보강",
                change_source=f"{abs_t} 결강 → {sub_t} 보강",
                change_id=str(getattr(r,"결강ID","")).strip() + f"-{p}",
                change_detail=str(getattr(r,"배정방식","")).strip())
    pt_df = st.session_state.get("part_time", pd.DataFrame())
    if pt_df is not None and not pt_df.empty:
        for prow in pt_df.itertuples(index=False):
            start, end = normalize_date_str(getattr(prow, "시작일", "")), normalize_date_str(getattr(prow, "종료일", ""))
            if not (start and end and start <= norm <= end): continue
            orig, pt_name = str(getattr(prow, "대체교사", "")).strip(), str(getattr(prow, "시간강사명", "")).strip()
            if not orig or not pt_name or orig == pt_name: continue
            for (teacher,p), lesson in list(current.items()):
                if teacher != orig or not _part_time_available(pd.Series(prow._asdict()), day, p): continue
                current.pop((teacher,p), None)
                lesson = dict(lesson); lesson["교사명"] = pt_name; lesson["원본교사"] = orig
                lesson["변경유형"] = "시간강사"; lesson["변경출처"] = f"{orig} → {pt_name}"
                lesson["변경ID"] = lesson.get("변경ID") or _new_change_id("PT")
                lesson["변경상세"] = f"{day}{p} 가능시간에 따른 대체"
                current[(pt_name,p)] = lesson
    df = pd.DataFrame(list(current.values()))
    for c in columns:
        if c not in df.columns: df[c] = ""
    return df[columns].reset_index(drop=True)
def get_swap_origin_info(teacher: str, on_date: str, period: int) -> str:
    norm_date = normalize_date_str(on_date)
    if not norm_date:
        return ""
    swaps = st.session_state.get("swaps", pd.DataFrame())
    if swaps.empty:
        return ""
    p = safe_int(period)
    mask1 = (swaps["목표일자"] == norm_date) & (swaps["교사A"] == teacher) & (swaps["교시B"] == p)
    if mask1.any():
        row = swaps[mask1].iloc[0]
        return f"{row.get('요일A','')}{safe_int(row.get('교시A',0))}({row.get('교사B','')})"
    mask2 = (swaps["원본일자"] == norm_date) & (swaps["교사B"] == teacher) & (swaps["교시A"] == p)
    if mask2.any():
        row = swaps[mask2].iloc[0]
        return f"{row.get('요일B','')}{safe_int(row.get('교시B',0))}({row.get('교사A','')})"
    return ""
@st.cache_data(show_spinner=False, ttl=180)
def _duty_slot_index(version=0):
    duties = st.session_state.get("duties", pd.DataFrame())
    index = defaultdict(set)
    if isinstance(duties, pd.DataFrame) and not duties.empty:
        cols = {"교사명", "일자", "교시"}
        if cols.issubset(duties.columns):
            for r in duties.itertuples(index=False):
                teacher = str(getattr(r, "교사명", "")).strip()
                if not teacher:
                    continue
                day = normalize_date_str(getattr(r, "일자", ""))
                p = safe_int(getattr(r, "교시", 0))
                if day:
                    index[teacher].add((day, 0 if p <= 0 else p))
    return dict(index)
def has_duty(teacher: str, on_date: str, period: int = None) -> bool:
    norm = normalize_date_str(on_date)
    slots = _duty_slot_index(st.session_state.get("_data_version", 0)).get(str(teacher).strip(), set())
    if not slots or not norm:
        return False
    if period is None:
        return any(day == norm for day, _ in slots)
    p = safe_int(period)
    return (norm, 0) in slots or (norm, p) in slots
def is_free(teacher: str, day: str, period: int, on_date: str = None, e_tt=None) -> bool:
    p = safe_int(period)
    norm = normalize_date_str(on_date)
    if has_duty(teacher, norm, p):
        return False
    if e_tt is None:
        e_tt = get_effective_timetable_for_date(norm, st.session_state.get("_data_version", 0))
    if not e_tt.empty and ((e_tt["교사명"] == teacher) & (e_tt["교시"] == p)).any():
        return False
    return True
@st.cache_data(show_spinner=False)
def cumulative_sub_count(start_date=None, end_date=None, version=0):
    s = st.session_state.subs
    base = {t: 0 for t in st.session_state.teachers["교사명"].tolist()} if not st.session_state.teachers.empty else {}
    if s.empty or "보강교사" not in s.columns:
        return base
    if start_date and end_date:
        s = s[(s["일자"] >= normalize_date_str(start_date)) & (s["일자"] <= normalize_date_str(end_date))]
    counts = s["보강교사"].value_counts()
    for k, v in counts.items():
        if k in base:
            base[k] = int(v)
    return base
@st.cache_data(show_spinner=False)
def weekly_load(version=0):
    tt = st.session_state.timetable
    return tt["교사명"].value_counts().to_dict() if not tt.empty else {}
def recommend_substitutes(day, period, subject, class_name, absent_teacher, on_date, top_n=20, include_part_time=False, e_tt=None):
    teachers = st.session_state.teachers
    norm = normalize_date_str(on_date)
    if e_tt is None:
        e_tt = get_effective_timetable_for_date(norm, st.session_state.get("_data_version", 0))
    grp, grade = subject_group(subject), grade_of(class_name)
    ver = st.session_state.get("_data_version", 0)
    cum, load = cumulative_sub_count(version=ver), weekly_load(version=ver)
    max_cum = max(cum.values()) if cum else 0
    rows=[]
    occupied = {(str(r.교사명).strip(), safe_int(r.교시)) for r in e_tt.itertuples(index=False)} if e_tt is not None and not e_tt.empty else set()
    duty_set = set()
    duties_df = st.session_state.get("duties", pd.DataFrame())
    if duties_df is not None and not duties_df.empty:
        dmask = (duties_df["교사명"].astype(str).str.strip() == str(absent_teacher).strip()) & (duties_df["일자"].astype(str).map(normalize_date_str) == norm)
        duty_set = {(str(r.교사명).strip(), safe_int(r.교시)) for r in duties_df[duties_df["일자"].astype(str).map(normalize_date_str) == norm].itertuples(index=False)}
    teacher_groups = defaultdict(set); teacher_grades = defaultdict(set)
    if e_tt is not None and not e_tt.empty:
        for rr in e_tt.itertuples(index=False):
            tn = str(rr.교사명).strip()
            if not tn: continue
            teacher_groups[tn].add(str(getattr(rr, "과목군", "")).strip() or subject_group(getattr(rr, "과목", "")))
            teacher_grades[tn].add(grade_of(getattr(rr, "학급", "")))
    teacher_subject = {}
    if not teachers.empty and "교사명" in teachers.columns:
        for rr in teachers.itertuples(index=False):
            tn = str(getattr(rr, "교사명", "")).strip()
            if tn: teacher_subject[tn] = str(getattr(rr, "담당과목", "")).strip()
    for t in (teachers["교사명"].astype(str).str.strip().tolist() if not teachers.empty and "교사명" in teachers.columns else []):
        if t == absent_teacher or (t, safe_int(period)) in occupied or (t, safe_int(period)) in duty_set: continue
        groups=teacher_groups.get(t,set()); grades=teacher_grades.get(t,set())
        if grp in groups and grade in grades: prio,label,score=1,"1순위 · 동일 과목 & 동일 학년",100
        elif grp in groups: prio,label,score=2,"2순위 · 동일 과목",70
        elif grade in grades: prio,label,score=3,"3순위 · 동일 학년",45
        else: prio,label,score=4,"4순위 · 전체 공강",20
        score += (max_cum-cum.get(t,0))*2 + max(0,22-load.get(t,0))*0.3
        rows.append({"보강교사":t,"유형":"정규교사","우선순위":label,"_prio":prio,"담당과목":teacher_subject.get(t,""),"주당시수":load.get(t,0),"누적보강":cum.get(t,0),"추천점수":round(score,1)})
    if include_part_time:
        pt=st.session_state.get("part_time",pd.DataFrame())
        if not pt.empty:
            for _,prow in pt.iterrows():
                t=str(prow.get("시간강사명","")).strip()
                if not t or t==absent_teacher or not _part_time_available(prow,day,period): continue
                if not is_free(t,day,period,norm,e_tt): continue
                pgrp=subject_group(str(prow.get("담당과목","")).strip() or str(prow.get("과목군","")).strip())
                prio=1 if pgrp and pgrp==grp else 2 if pgrp else 4
                label="1순위 · 시간강사 동일 과목" if prio==1 else "2순위 · 시간강사" if prio==2 else "4순위 · 시간강사 공강"
                score=90 if prio==1 else 60 if prio==2 else 15
                rows.append({"보강교사":t,"유형":"시간강사","우선순위":label,"_prio":prio,"담당과목":prow.get("담당과목","") ,"주당시수":0,"누적보강":0,"추천점수":round(score,1)})
    if not rows: return pd.DataFrame()
    return pd.DataFrame(rows).sort_values(["_prio","추천점수"],ascending=[True,False]).drop(columns=["_prio"]).head(top_n).reset_index(drop=True)
@st.cache_data(show_spinner=False, ttl=120)
def get_cached_substitute_recommendations(day, period, subject, class_name, absent_teacher, on_date, top_n=10, include_part_time=True, version=0):
    return recommend_substitutes(
        day, period, subject, class_name, absent_teacher, on_date,
        top_n=top_n, include_part_time=include_part_time
    )
def add_substitute(cid, on_date, day, period, class_name, subject, absent_teacher, sub_teacher, method, priority, memo, *, save=True, history=True):
    ver = st.session_state.get("_data_version", 0)
    e_tt = get_effective_timetable_for_date(normalize_date_str(on_date), ver)
    ok, msg = validate_substitute(cid, on_date, period, class_name, subject, absent_teacher, sub_teacher, e_tt=e_tt)
    if not ok:
        st.warning(msg)
        return False
    p = safe_int(period); norm = normalize_date_str(on_date)
    s = st.session_state.subs.copy(deep=True)
    old = s[(s["결강ID"].astype(str) == str(cid)) & (s["교시"].apply(safe_int) == p)] if not s.empty else pd.DataFrame()
    if not old.empty:
        s = s.drop(old.index)
    new = pd.DataFrame([{
        "결강ID": cid, "일자": norm, "요일": day, "교시": p, "학급": class_name, "과목": subject,
        "결강교사": absent_teacher, "보강교사": sub_teacher, "배정방식": method, "우선순위": priority,
        "비고": memo, "등록시각": _now_text("%Y-%m-%d %H:%M"), "입력자": current_user()
    }])
    st.session_state.subs = pd.concat([s, new], ignore_index=True)
    if save:
        if not save_work_data_to_gsheet(["보강"]):
            st.session_state.subs = s
            _invalidate_all_caches()
            return False
        if old.empty and update_budget(-SUB_COST, f"보강 1건 ({sub_teacher} ← {absent_teacher})") is None:
            st.session_state.subs = s
            try:
                save_work_data_to_gsheet(["보강"])
            except Exception:
                pass
            _invalidate_all_caches()
            return False
    _invalidate_all_caches()
    if history:
        push_history(f"보강 배정 ({sub_teacher})")
    return True
def add_substitutes_batch(assignments, action_name="자동 보강"):
    if not assignments: return 0, []
    original_subs = st.session_state.subs.copy(deep=True)
    working = original_subs.copy(deep=True)
    added_new = 0; accepted = 0; errors = []
    for rec in assignments:
        cid, norm, p = rec["cid"], normalize_date_str(rec["on_date"]), safe_int(rec["period"])
        st.session_state.subs = working
        get_effective_timetable_for_date.clear()
        e_tt = get_effective_timetable_for_date(norm, st.session_state.get("_data_version",0))
        ok, msg = validate_substitute(cid, norm, p, rec["class_name"], rec["subject"], rec["absent_teacher"], rec["sub_teacher"], e_tt=e_tt)
        if not ok:
            errors.append(f"{p}교시: {msg}"); continue
        old = working[(working["결강ID"].astype(str) == str(cid)) & (working["교시"].apply(safe_int) == p)] if not working.empty else pd.DataFrame()
        if not old.empty: working = working.drop(old.index)
        else: added_new += 1
        working = pd.concat([working, pd.DataFrame([{
            "결강ID": cid, "일자": norm, "요일": rec["day"], "교시": p, "학급": rec["class_name"], "과목": rec["subject"],
            "결강교사": rec["absent_teacher"], "보강교사": rec["sub_teacher"], "배정방식": rec.get("method","자동"),
            "우선순위": rec.get("priority",""), "비고": rec.get("memo",""), "등록시각": _now_text("%Y-%m-%d %H:%M"), "입력자": current_user()
        }])], ignore_index=True)
        accepted += 1
    st.session_state.subs = working.reset_index(drop=True)
    if accepted:
        if save_work_data_to_gsheet(["보강"]):
            if added_new and update_budget(-(added_new * SUB_COST), f"{action_name} {added_new}건") is None:
                st.session_state.subs = original_subs
                try:
                    save_work_data_to_gsheet(["보강"])
                except Exception:
                    pass
                accepted = 0
                errors.append("예산 저장 실패로 보강 작업을 원복했습니다.")
            else:
                push_history(action_name)
        else:
            st.session_state.subs = original_subs
            accepted = 0
            errors.append("보강 저장 실패로 작업을 원복했습니다.")
    else:
        if st.session_state.history and st.session_state.history_index == len(st.session_state.history)-1:
            st.session_state.history.pop(); st.session_state.history_index -= 1
    _invalidate_all_caches()
    return accepted, errors
def cancel_substitute(cid, period):
    if not can_full_data():
        s = st.session_state.subs; p = safe_int(period)
        m = s[(s["결강ID"] == cid) & (s["교시"] == p)]
        if not m.empty and m.iloc[0].get("입력자") != current_user():
            st.warning("본인이 입력한 데이터만 취소할 수 있습니다.")
            return False
    p = safe_int(period); s = st.session_state.subs.copy(deep=True)
    m = s[(s["결강ID"] == cid) & (s["교시"] == p)]
    if m.empty:
        return False
    before = s
    st.session_state.subs = s.drop(m.index).reset_index(drop=True)
    if not save_work_data_to_gsheet(["보강"]):
        st.session_state.subs = before
        _invalidate_all_caches()
        return False
    if update_budget(+SUB_COST, f"보강 취소 복구 ({period}교시)") is None:
        st.session_state.subs = before
        try:
            save_work_data_to_gsheet(["보강"])
        except Exception:
            pass
        _invalidate_all_caches()
        return False
    _invalidate_all_caches()
    push_history(f"보강 취소 ({period}교시)")
    return True
def do_swap(a, b, date_a, date_b, is_part_time_purpose=False, is_test=False):
    class_a = str(a.get("학급", "")).strip()
    class_b = str(b.get("학급", "")).strip()
    if class_a and class_b and class_a != class_b:
        if is_test:
            st.error("1:1 맞교환 테스트는 동일 학급 수업끼리만 가능합니다.")
        else:
            st.warning("1:1 맞교환은 동일 학급 수업끼리만 가능합니다.")
        return False
    ok, msg = validate_swap(a, b, date_a, date_b, is_test=is_test)
    if not ok:
        if not is_test: st.warning(msg)
        return False
    ver = st.session_state.get("_data_version", 0)
    e_a_now = get_effective_timetable_for_date(normalize_date_str(date_a), ver, use_test=is_test)
    e_b_now = get_effective_timetable_for_date(normalize_date_str(date_b), ver, use_test=is_test)
    ma_now = e_a_now[(e_a_now["교사명"] == a["교사명"]) & (e_a_now["교시"].apply(safe_int) == safe_int(a["교시"]))] if not e_a_now.empty else pd.DataFrame()
    mb_now = e_b_now[(e_b_now["교사명"] == b["교사명"]) & (e_b_now["교시"].apply(safe_int) == safe_int(b["교시"]))] if not e_b_now.empty else pd.DataFrame()
    oa = ma_now.iloc[0] if not ma_now.empty else a
    ob = mb_now.iloc[0] if not mb_now.empty else b
    rec = {
        "변경ID": _new_change_id("TESTSW" if is_test else "SWAP"),
        "원본교사A": str(oa.get("원본교사", a["교사명"])), "실제원본일자A": str(oa.get("원본일자", normalize_date_str(date_a))), "실제원본교시A": safe_int(oa.get("원본교시", a["교시"])),
        "원본교사B": str(ob.get("원본교사", b["교사명"])), "실제원본일자B": str(ob.get("원본일자", normalize_date_str(date_b))), "실제원본교시B": safe_int(ob.get("원본교시", b["교시"])),
        "원본일자": normalize_date_str(date_a), "교사A": a["교사명"], "요일A": a["요일"], "교시A": safe_int(a["교시"]),
        "학급A": a.get("학급", ""), "과목A": a.get("과목", ""), "목표일자": normalize_date_str(date_b),
        "교사B": b["교사명"], "요일B": b["요일"], "교시B": safe_int(b["교시"]), "학급B": b.get("학급", ""), "과목B": b.get("과목", ""),
        "유형": "1:1 맞교환", "시간강사구인": "Y" if is_part_time_purpose else "N",
        "등록시각": _now_text("%Y-%m-%d %H:%M"), "입력자": current_user()
    }
    if is_test:
        st.session_state.test_swaps = pd.concat([st.session_state.get("test_swaps", pd.DataFrame()), pd.DataFrame([rec])], ignore_index=True)
        get_effective_timetable_for_date.clear(); effective_teacher_matrix.clear(); get_single_lesson_1to1_candidates.clear(); get_single_lesson_linked_cycles.clear(); get_test_affected_slots.clear(); find_cycle_linked_swaps.clear(); get_target_time_recommendations.clear()
        return True
    before = st.session_state.swaps.copy(deep=True)
    st.session_state.swaps = pd.concat([before, pd.DataFrame([rec])], ignore_index=True)
    if not save_work_data_to_gsheet(["맞교환"]):
        st.session_state.swaps = before
        _invalidate_all_caches()
        return False
    push_history(f"맞교환 ({a['교사명']} ↔ {b['교사명']})")
    return True
def do_linked_swap(a, teacher_b, date_a, date_b, day_b, period_b, is_part_time_purpose=False, is_test=False, subject_b=None, *, save=True, history=True):
    b_probe = {"교사명": teacher_b, "일자": date_b, "요일": day_b, "교시": period_b, "학급": "", "과목": ""}
    ver = st.session_state.get("_data_version", 0)
    e_b = get_effective_timetable_for_date(normalize_date_str(date_b), ver, use_test=is_test)
    norm_a=normalize_date_str(date_a)
    norm_b=normalize_date_str(date_b)
    day_a=WEEKDAY_KR[datetime.strptime(norm_a, "%Y-%m-%d").weekday()] if norm_a else ""
    if not is_free(teacher_b, day_b, period_b, norm_b, e_b):
        return False
    if not is_test and _actual_direct_slot_is_affected(norm_b, teacher_b, period_b):
        return False
    if not teacher_slot_is_available(str(a.get("교사명", "")), day_b, period_b, ver):
        return False
    if not teacher_slot_is_available(teacher_b, day_a, safe_int(a.get("교시", 0)), ver):
        return False
    if is_test and _test_slot_is_affected(date_b, teacher_b, period_b):
        return False
    rec = {
        "변경ID": _new_change_id("TESTLINK" if is_test else "LINK"),
        "원본일자": normalize_date_str(date_a), "교사A": a["교사명"], "요일A": a["요일"], "교시A": safe_int(a["교시"]),
        "학급A": a.get("학급", ""), "과목A": a.get("과목", ""), "목표일자": normalize_date_str(date_b),
        "교사B": teacher_b, "요일B": day_b, "교시B": safe_int(period_b), "학급B": a.get("학급", ""),
        "과목B": subject_b if subject_b is not None else a.get("과목", ""), "유형": "연계 공강 교환",
        "시간강사구인": "Y" if is_part_time_purpose else "N", "등록시각": _now_text("%Y-%m-%d %H:%M"), "입력자": current_user()
    }
    if is_test:
        st.session_state.test_swaps = pd.concat([st.session_state.get("test_swaps", pd.DataFrame()), pd.DataFrame([rec])], ignore_index=True)
        get_effective_timetable_for_date.clear(); effective_teacher_matrix.clear(); get_single_lesson_1to1_candidates.clear(); get_single_lesson_linked_cycles.clear(); get_test_affected_slots.clear(); find_cycle_linked_swaps.clear(); get_target_time_recommendations.clear()
        return True
    before = st.session_state.swaps.copy(deep=True)
    st.session_state.swaps = pd.concat([before, pd.DataFrame([rec])], ignore_index=True)
    if save and not save_work_data_to_gsheet(["맞교환"]):
        st.session_state.swaps = before
        _invalidate_all_caches()
        return False
    if history: push_history(f"연계교환 ({a['교사명']} → {teacher_b})")
    return True
def apply_cycle_swaps(moves, is_test=False):
    if not moves:
        return False
    if is_test:
        before = st.session_state.get("test_swaps", pd.DataFrame()).copy(deep=True)
        before_has_cycle = bool(st.session_state.get("test_has_cycle", False))
        try:
            for m in moves:
                a_info = {
                    "교사명": m["teacher"],
                    "요일": m.get("day_from", WEEKDAY_KR[datetime.strptime(m["from_date"], "%Y-%m-%d").weekday()]),
                    "교시": m["from_period"],
                    "학급": m["class"],
                    "과목": m["subject"],
                }
                to_day = WEEKDAY_KR[datetime.strptime(m["to_date"], "%Y-%m-%d").weekday()]
                if not do_linked_swap(
                    a_info, m.get("next_teacher", m["teacher"]),
                    m["from_date"], m["to_date"], to_day, m["to_period"],
                    is_test=True, subject_b=m.get("target_subject", m["subject"])
                ):
                    raise ValueError("순환 테스트의 일부 이동을 적용할 수 없습니다.")
            return True
        except Exception:
            st.session_state.test_swaps = before
            st.session_state["test_has_cycle"] = before_has_cycle
            get_effective_timetable_for_date.clear()
            effective_teacher_matrix.clear()
            get_single_lesson_1to1_candidates.clear()
            get_single_lesson_linked_cycles.clear()
            get_test_affected_slots.clear()
            find_cycle_linked_swaps.clear()
            get_target_time_recommendations.clear()
            return False
    before = st.session_state.swaps.copy(deep=True)
    for m in moves:
        a_info = {"교사명": m["teacher"], "요일": m.get("day_from", WEEKDAY_KR[datetime.strptime(m["from_date"], "%Y-%m-%d").weekday()]), "교시": m["from_period"], "학급": m["class"], "과목": m["subject"]}
        to_day = WEEKDAY_KR[datetime.strptime(m["to_date"], "%Y-%m-%d").weekday()]
        if not do_linked_swap(a_info, m.get("next_teacher", m["teacher"]), m["from_date"], m["to_date"], to_day, m["to_period"], is_test=False, subject_b=m.get("target_subject", m["subject"]), save=False, history=False):
            st.session_state.swaps = before
            return False
    if not save_work_data_to_gsheet(["맞교환"]):
        st.session_state.swaps = before
        _invalidate_all_caches()
        return False
    push_history(f"{len(moves)}인 연계교환")
    return True
@st.cache_data(show_spinner=False, ttl=180)
def find_cycle_linked_swaps(teacher_a, date_a_str, period_a, class_a, subject_a,
                           date_b_str, period_b, min_cycle=2, max_cycle=3, future_days=7, version=0, use_test=False):
    original_slot = (normalize_date_str(date_a_str), safe_int(period_a))
    target_slot = (normalize_date_str(date_b_str), safe_int(period_b))
    if not original_slot[0] or not target_slot[0]:
        return [], "날짜 오류"
    try:
        da, db = date.fromisoformat(original_slot[0]), date.fromisoformat(target_slot[0])
    except ValueError:
        return [], "날짜 오류"
    if da.weekday() >= 5 or db.weekday() >= 5:
        return [], "날짜 오류"
    ver = version or st.session_state.get("_data_version", 0)
    candidates=[]
    cur=min(da,db)-timedelta(days=2); end=max(da,db)+timedelta(days=future_days)
    while cur<=end:
        if cur.weekday()<5:
            candidates.append(cur.isoformat())
        cur+=timedelta(days=1)
    weekday_cache={d:WEEKDAY_KR[date.fromisoformat(d).weekday()] for d in candidates}
    e_cache={d:get_effective_timetable_for_date(d,ver,use_test=use_test) for d in candidates}
    affected_test=get_test_affected_slots(ver) if use_test else set()
    affected_actual=get_actual_direct_swap_affected_slots(ver) if not use_test else set()
    excluded=affected_test|affected_actual
    class_slots={}; teacher_occupied=defaultdict(set)
    for d,e in e_cache.items():
        if e is None or e.empty: continue
        day=weekday_cache[d]
        for r in e.itertuples(index=False):
            t=str(getattr(r,"교사명","")).strip(); p=safe_int(getattr(r,"교시",0))
            if not t or p<=0: continue
            teacher_occupied[t].add((d,p))
            if str(getattr(r,"학급","")).strip()==class_a and (d,t,p) not in excluded:
                class_slots[(d,p)]={"teacher":t,"subject":str(getattr(r,"과목","")).strip(),"day":day}
    if original_slot not in class_slots or class_slots[original_slot]["teacher"]!=teacher_a:
        return [], "원본 슬롯/교사 불일치"
    if target_slot not in class_slots:
        return [], "목표 슬롯에 학급 수업 없음 (공강 생성 금지)"
    avail=get_teacher_availability_index(ver); duties=_duty_slot_index(ver)
    def free_at(t,d,p):
        occ=teacher_occupied.get(t,set())
        if (d,p) in occ or (d,t,p) in excluded: return False
        ds=duties.get(t,set())
        if (d,0) in ds or (d,p) in ds: return False
        item=avail.get(t)
        if item is not None and item[0] and (weekday_cache.get(d,WEEKDAY_KR[date.fromisoformat(d).weekday()]),p) not in item[1]: return False
        return True
    if not free_at(teacher_a,target_slot[0],target_slot[1]):
        return [], "교사A 목표시간 수업 있음"
    slot_items=list(class_slots.items())
    free_moves=defaultdict(list)
    for from_s,info in slot_items:
        t=info["teacher"]; occupied=teacher_occupied.get(t,set()); ds=duties.get(t,set()); item=avail.get(t)
        for to_s,_ in slot_items:
            if to_s==from_s or to_s in occupied or (to_s[0],t,to_s[1]) in excluded: continue
            d,p=to_s; day=weekday_cache[d]
            if (d,0) in ds or (d,p) in ds: continue
            if item is not None and item[0] and (day,p) not in item[1]: continue
            free_moves[from_s].append(to_s)
    cycles=[]; max_found=6
    def dfs(current,path,visited):
        if len(cycles)>=max_found or len(path)>max_cycle: return
        if current!=original_slot and (current[0],class_slots.get(current,{}).get("teacher",""),current[1]) in excluded: return
        if current==original_slot:
            if len(path)<min_cycle: return
            cycle_slots=[original_slot]+path[:-1]; moves=[]; n=len(cycle_slots)
            for i,from_s in enumerate(cycle_slots):
                to_s=cycle_slots[(i+1)%n]; info=class_slots[from_s]; tgt=class_slots[to_s]
                moves.append({"teacher":info["teacher"],"from_date":from_s[0],"from_period":from_s[1],"to_date":to_s[0],"to_period":to_s[1],"class":class_a,"subject":info["subject"],"target_subject":tgt["subject"],"day_from":info["day"],"next_teacher":tgt["teacher"]})
            path_desc=" → ".join(f"{m['teacher']}({m['class']} {m['from_date'][5:]} {m['from_period']}→{m['to_date'][5:]} {m['to_period']})" for m in moves)
            cycles.append({"length":n,"moves":moves,"path_desc":path_desc,"score":110-n*12}); return
        for nxt in free_moves.get(current,()):
            if nxt in visited: continue
            visited.add(nxt); path.append(nxt); dfs(nxt,path,visited); path.pop(); visited.remove(nxt)
    dfs(target_slot,[target_slot],{target_slot})
    cycles.sort(key=lambda x:(x["length"],-x["score"]))
    return cycles[:max_found], f"{len(cycles)}개 순환 경로 발견" if cycles else "순환 경로 없음"
@st.cache_data(show_spinner=False, ttl=180)
def get_target_time_recommendations(teacher_a, date_a_str, period_a, class_a, subject_a, date_b_str, period_b, budget_factor=1.0, version=0, use_test=False):
    norm_a,norm_b=normalize_date_str(date_a_str),normalize_date_str(date_b_str); ver=version or st.session_state.get("_data_version",0)
    ti=st.session_state.get("teachers",pd.DataFrame()); e_a=get_effective_timetable_for_date(norm_a,ver,use_test=use_test); e_b=get_effective_timetable_for_date(norm_b,ver,use_test=use_test)
    if ti.empty or e_a.empty or e_b.empty: return pd.DataFrame(),[],""
    p_a,p_b=safe_int(period_a),safe_int(period_b); day_a=WEEKDAY_KR[date.fromisoformat(norm_a).weekday()]; day_b=WEEKDAY_KR[date.fromisoformat(norm_b).weekday()]
    my_class=str(class_a).strip(); my_grade=grade_of(my_class); my_group=subject_group(subject_a); cum=cumulative_sub_count(version=ver)
    avail=get_teacher_availability_index(ver); duties=_duty_slot_index(ver)
    def avail_ok(t,d,p):
        item=avail.get(t); return True if item is None or not item[0] else (d,p) in item[1]
    occ_a={(str(r.교사명).strip(),safe_int(r.교시)) for r in e_a.itertuples(index=False)}; occ_b={(str(r.교사명).strip(),safe_int(r.교시)) for r in e_b.itertuples(index=False)}
    def free(t,d,p,occ):
        if (t,p) in occ: return False
        ds=duties.get(t,set()); return (d,0) not in ds and (d,p) not in ds
    if not free(teacher_a,day_b,p_b,occ_b) or not avail_ok(teacher_a,day_b,p_b): return pd.DataFrame(),[],""
    teacher_names={str(x).strip() for x in ti["교사명"].dropna().tolist()}
    b_slots=defaultdict(list)
    for r in e_b.itertuples(index=False):
        if safe_int(getattr(r,"교시",0))==p_b:
            t=str(getattr(r,"교사명","")).strip()
            if t in teacher_names and t: b_slots[t].append(r)
    swap_recs=[]
    for t_b,b_lessons in b_slots.items():
        if t_b==teacher_a: continue
        ds=duties.get(t_b,set())
        if (norm_b,0) in ds or (norm_b,p_b) in ds or not avail_ok(t_b,day_a,p_a) or not free(t_b,day_a,p_a,occ_a): continue
        for r in b_lessons:
            other_class=str(getattr(r,"학급","")).strip()
            if other_class!=my_class: continue
            other_subject=str(getattr(r,"과목","")).strip(); other_grade=grade_of(other_class); same_group=subject_group(other_subject)==my_group
            score=(200+(40 if same_group else 0)+(15 if norm_a==norm_b else 0)-cum.get(t_b,0)*3)*budget_factor
            swap_recs.append({"유형":"1:1","교사B":t_b,"현재 수업":f"{day_b}{p_b}교시 · {other_class} · {other_subject}","학급":other_class,"학년":other_grade,"same_class":True,"same_grade":other_grade==my_grade,"점수":score,"b_info":{"교사명":t_b,"일자":norm_b,"요일":day_b,"교시":p_b,"학급":other_class,"과목":other_subject}})
    df_swap=(pd.DataFrame(swap_recs).sort_values(["same_class","same_grade","점수"],ascending=[False,False,False]).reset_index(drop=True) if swap_recs else pd.DataFrame())
    cycles,msg=find_cycle_linked_swaps(teacher_a,norm_a,p_a,my_class,subject_a,norm_b,p_b,max_cycle=3,future_days=7,version=ver,use_test=use_test)
    return df_swap,cycles,msg
@st.cache_data(show_spinner=False, ttl=180)
def get_weekly_1to1_swap_table(teacher: str, ref_date: date, future_days: int = 0, version: int = 0) -> pd.DataFrame:
    weekday = ref_date.weekday()
    monday = ref_date - timedelta(days=weekday)
    week_dates = [(monday + timedelta(days=i)).strftime("%Y-%m-%d") for i in range(5)]
    ver = version or st.session_state.get("_data_version", 0)
    results = []
    seen = set()
    search_dates = week_dates[:]
    if future_days > 0:
        last = datetime.strptime(week_dates[-1], "%Y-%m-%d").date()
        for i in range(1, future_days + 1):
            nd = last + timedelta(days=i)
            if nd.weekday() < 5:
                search_dates.append(nd.strftime("%Y-%m-%d"))
    search_dates = sorted(set(search_dates))
    for d_str in week_dates:
        day_kr = WEEKDAY_KR[datetime.strptime(d_str, "%Y-%m-%d").weekday()]
        e_tt = get_effective_timetable_for_date(d_str, ver)
        if e_tt.empty:
            continue
        my_lessons = e_tt[(e_tt["교사명"] == teacher) & (e_tt["요일"] == day_kr)].sort_values("교시")
        for _, lesson in my_lessons.iterrows():
            p = safe_int(lesson["교시"])
            my_class = str(lesson["학급"]).strip()
            my_subj  = str(lesson["과목"]).strip()
            my_grade = grade_of(my_class)
            my_group = subject_group(my_subj)
            for td_str in search_dates:
                if td_str == d_str:
                    continue
                tday = WEEKDAY_KR[datetime.strptime(td_str, "%Y-%m-%d").weekday()]
                e_b = get_effective_timetable_for_date(td_str, ver)
                if e_b.empty:
                    continue
                others = e_b[(e_b["교시"] == p) & (e_b["교사명"] != teacher)]
                others = others.drop_duplicates(subset=["교사명", "교시"])
                for _, o in others.iterrows():
                    other_teacher = str(o["교사명"]).strip()
                    dup_key = (d_str, p, td_str, other_teacher)
                    if dup_key in seen:
                        continue
                    if not is_free(teacher, tday, p, td_str, e_b):
                        continue
                    if not teacher_slot_is_available(teacher, tday, p, ver):
                        continue
                    if not is_free(other_teacher, day_kr, p, d_str, e_tt):
                        continue
                    if not teacher_slot_is_available(other_teacher, day_kr, p, ver):
                        continue
                    other_class = str(o["학급"]).strip()
                    same_class = (other_class == my_class)
                    if not same_class:
                        continue
                    other_grade = grade_of(other_class)
                    same_grade  = (other_grade == my_grade)
                    score = 200
                    if subject_group(str(o["과목"])) == my_group: score += 40
                    if td_str[:7] == d_str[:7]: score += 10
                    seen.add(dup_key)
                    results.append({
                        "원본일자": d_str,
                        "원본요일": day_kr,
                        "원본교시": p,
                        "원본학급": my_class,
                        "원본과목": my_subj,
                        "이동희망일": td_str,
                        "이동요일": tday,
                        "상대교사": other_teacher,
                        "상대학급": str(other_class), "상대과목": str(o["과목"]),
                        "상대수업": f"{other_class} {o['과목']}",
                        "동일학급": "🏆" if same_class else "",
                        "동학년": "⚠" if same_grade and not same_class else "",
                        "점수": score,
                        "_sort": (0 if same_class else 1, 0 if same_grade else 1, -score)
                    })
    if not results:
        return pd.DataFrame(columns=[
            "원본일자", "원본요일", "원본교시", "원본학급", "원본과목",
            "이동희망일", "이동요일", "상대교시", "상대교사", "상대수업",
            "동일학급", "동학년", "점수"
        ])
    df = pd.DataFrame(results)
    df = df.sort_values("_sort").drop(columns=["_sort"]).reset_index(drop=True)
    return df
@st.cache_data(show_spinner=False, ttl=180)
def get_single_lesson_1to1_candidates(
    teacher: str, orig_date_str: str, orig_period: int,
    orig_class: str, orig_subject: str, future_days: int = 0, version: int = 0, use_test: bool = False,
    target_date_str: str | None = None, target_period: int | None = None
) -> pd.DataFrame:
    cols=["원본일자","원본요일","원본교시","원본학급","원본과목","이동희망일","이동요일","이동희망교시","상대교사","상대학급","상대과목","상대수업","동일학급","동학년","교환가능사유","점수"]
    source_str=normalize_date_str(orig_date_str); teacher=str(teacher).strip(); source_period=safe_int(orig_period); source_class=str(orig_class).strip(); source_subject=str(orig_subject).strip()
    if not teacher or not source_str or source_period<=0 or not source_class: return pd.DataFrame(columns=cols)
    try: source_date=date.fromisoformat(source_str)
    except ValueError: return pd.DataFrame(columns=cols)
    target_date=None
    if target_date_str:
        try: target_date=date.fromisoformat(normalize_date_str(target_date_str))
        except ValueError: target_date=None
    if target_date is not None:
        search_dates=[target_date]
    else:
        monday=source_date-timedelta(days=source_date.weekday()); friday=monday+timedelta(days=4); end=friday+timedelta(days=max(0,int(future_days)))
        search_dates=[source_date+timedelta(days=i) for i in range((end-source_date).days+1) if (source_date+timedelta(days=i)).weekday()<5]
    ver=version or st.session_state.get("_data_version",0); source_day=WEEKDAY_KR[source_date.weekday()]; source_tt=get_effective_timetable_for_date(source_str,ver,use_test=use_test)
    if source_tt is None or source_tt.empty: return pd.DataFrame(columns=cols)
    source_rows=[r for r in source_tt.itertuples(index=False) if str(getattr(r,"교사명","")).strip()==teacher and safe_int(getattr(r,"교시",0))==source_period and str(getattr(r,"학급","")).strip()==source_class and str(getattr(r,"과목","")).strip()==source_subject]
    if not source_rows: return pd.DataFrame(columns=cols)
    avail=get_teacher_availability_index(ver); duties=_duty_slot_index(ver); test_affected=get_test_affected_slots(ver) if use_test else set(); actual_affected=get_actual_direct_swap_affected_slots(ver)
    source_busy={(str(r.교사명).strip(),safe_int(r.교시)) for r in source_tt.itertuples(index=False)}; source_group=subject_group(source_subject)
    def available(t,d,p):
        item=avail.get(t); return True if item is None or not item[0] else (d,p) in item[1]
    def duty_blocked(t,d,p):
        ds=duties.get(t,set()); return (d,0) in ds or (d,p) in ds
    results=[]; seen=set()
    for td in search_dates:
        if td.weekday()>=5: continue
        target_str=td.isoformat(); target_day=WEEKDAY_KR[td.weekday()]; target_tt=get_effective_timetable_for_date(target_str,ver,use_test=use_test)
        if target_tt is None or target_tt.empty: continue
        tp=safe_int(target_period) if target_date is not None and target_period is not None else None
        rows=[r for r in target_tt.itertuples(index=False) if tp is None or safe_int(getattr(r,"교시",0))==tp]
        periods=[tp] if tp is not None and 1<=tp<=PERIODS_PER_DAY.get(target_day,MAX_PERIOD) else sorted({safe_int(getattr(r,"교시",0)) for r in rows if 1<=safe_int(getattr(r,"교시",0))<=PERIODS_PER_DAY.get(target_day,MAX_PERIOD)}) if tp is None else []
        for target_p in periods:
            if target_str==source_str and target_p==source_period: continue
            if not available(teacher,target_day,target_p) or duty_blocked(teacher,target_str,target_p): continue
            for r in rows:
                if safe_int(getattr(r,"교시",0))!=target_p: continue
                other_teacher=str(getattr(r,"교사명","")).strip(); other_class=str(getattr(r,"학급","")).strip(); other_subject=str(getattr(r,"과목","")).strip()
                if not other_teacher or other_teacher==teacher or other_class!=source_class or not other_subject: continue
                if (other_teacher,source_period) in source_busy or duty_blocked(other_teacher,source_str,source_period) or not available(other_teacher,source_day,source_period): continue
                slots={(target_str,other_teacher,target_p),(source_str,other_teacher,source_period),(target_str,teacher,target_p)}
                if slots & test_affected or slots & actual_affected: continue
                key=(target_str,target_p,other_teacher,other_class,other_subject)
                if key in seen: continue
                seen.add(key); same_group=subject_group(other_subject)==source_group; score=300+(40 if same_group else 0)+(20+max(0,8-abs(target_p-source_period)) if target_str==source_str else 0)
                results.append({"원본일자":source_str,"원본요일":source_day,"원본교시":source_period,"원본학급":source_class,"원본과목":source_subject,"이동희망일":target_str,"이동요일":target_day,"이동희망교시":target_p,"상대교사":other_teacher,"상대학급":other_class,"상대과목":other_subject,"상대수업":f"{other_class} {other_subject}","동일학급":"🏆","동학년":"","교환가능사유":"동일 학급 · 내 공강 · 상대 교사 공강"+(" · 같은 과목군" if same_group else ""),"점수":score})
    if not results: return pd.DataFrame(columns=cols)
    return pd.DataFrame(results).sort_values(["동일학급","점수","이동희망일","이동희망교시","상대교사","상대학급"],ascending=[False,False,True,True,True,True],kind="stable").reset_index(drop=True)
@st.cache_data(show_spinner=False, ttl=180)
def get_single_lesson_linked_cycles(
    teacher: str, orig_date_str: str, orig_period: int,
    orig_class: str, orig_subject: str, future_days: int = 0, version: int = 0,
    min_cycle: int = 2, max_cycle: int = 3, use_test: bool = False
):
    source_str=normalize_date_str(orig_date_str); source_period=safe_int(orig_period); orig_class=str(orig_class).strip(); teacher=str(teacher).strip()
    try: source_date=date.fromisoformat(source_str)
    except ValueError: return [], "날짜 오류"
    monday=source_date-timedelta(days=source_date.weekday()); search_dates=[monday+timedelta(days=i) for i in range(5)]
    if future_days>0:
        friday=search_dates[-1]; search_dates.extend(friday+timedelta(days=i) for i in range(1,future_days+1) if (friday+timedelta(days=i)).weekday()<5)
    ver=version or st.session_state.get("_data_version",0); test_affected=get_test_affected_slots(ver) if use_test else set(); avail=get_teacher_availability_index(ver); duty_idx=_duty_slot_index(ver)
    target_slots=[]
    for d in search_dates:
        target_str=d.isoformat()
        if target_str<=source_str: continue
        day=WEEKDAY_KR[d.weekday()]; e=get_effective_timetable_for_date(target_str,ver,use_test=use_test)
        if e is None or e.empty: continue
        occupied={(str(r.교사명).strip(),safe_int(r.교시)) for r in e.itertuples(index=False)}
        ds=duty_idx.get(teacher,set()); item=avail.get(teacher)
        for r in e.itertuples(index=False):
            if str(getattr(r,"학급","")).strip()!=orig_class: continue
            p=safe_int(getattr(r,"교시",0)); current_t=str(getattr(r,"교사명","")).strip()
            if (teacher,p) in occupied or (target_str,0) in ds or (target_str,p) in ds: continue
            if item is not None and item[0] and (day,p) not in item[1]: continue
            if use_test and (target_str,current_t,p) in test_affected: continue
            priority=(0 if p==source_period else 1,abs((d-source_date).days),0 if subject_group(str(getattr(r,"과목","")))==subject_group(orig_subject) else 1)
            target_slots.append((priority,target_str,p))
    if not target_slots: return [], "연계 순환을 시작할 수 있는 빈 시간대가 없습니다."
    target_slots.sort(); all_cycles=[]; seen_paths=set()
    for _,target_str,target_period in target_slots[:10]:
        cycles,_=find_cycle_linked_swaps(teacher,source_str,source_period,orig_class,orig_subject,target_str,target_period,min_cycle=min_cycle,max_cycle=max_cycle,future_days=future_days,version=ver,use_test=use_test)
        for cycle in cycles:
            key=tuple((m["teacher"],m["from_date"],m["from_period"],m["to_date"],m["to_period"]) for m in cycle["moves"])
            if key not in seen_paths: seen_paths.add(key); all_cycles.append(cycle)
        if len(all_cycles)>=6: break
    all_cycles.sort(key=lambda c:(c["length"],-c["score"]))
    return (all_cycles[:6],f"{len(all_cycles)}개 연계 순환 경로 발견") if all_cycles else ([],"조건을 만족하는 연계 순환 경로가 없습니다.")
@st.cache_data(show_spinner=False)
def teacher_matrix(version=0):
    tt = st.session_state.timetable
    if tt.empty: return pd.DataFrame()
    idx = {(str(r.교사명).strip(), str(r.요일).strip(), safe_int(r.교시)): f"{r.학급} {r.과목}" for r in tt.itertuples(index=False)}
    teachers = sorted(tt["교사명"].astype(str).str.strip().unique())
    rows=[]
    for t in teachers:
        row={"교사명":t}
        for d in DAYS:
            for p in range(1, PERIODS_PER_DAY.get(d,7)+1): row[f"{d}{p}"]=idx.get((t,d,p),"")
        rows.append(row)
    return pd.DataFrame(rows)
@st.cache_data(show_spinner=False, ttl=180)
def effective_teacher_matrix(ref_date: date, version: int = 0, use_test: bool = False) -> pd.DataFrame:
    monday = ref_date - timedelta(days=ref_date.weekday())
    daily_timetables={}; teacher_names=set()
    base_tt=st.session_state.get("timetable",pd.DataFrame())
    if not base_tt.empty: teacher_names.update(base_tt["교사명"].dropna().astype(str).str.strip())
    for i,d in enumerate(DAYS):
        day_date=monday+timedelta(days=i); ds=day_date.strftime("%Y-%m-%d")
        day_tt=get_effective_timetable_for_date(ds,version,use_test=use_test); daily_timetables[d]=(ds,day_tt)
        if not day_tt.empty: teacher_names.update(day_tt["교사명"].dropna().astype(str).str.strip())
    daily_indexes = {}
    for d, (_, day_tt) in daily_timetables.items():
        daily_indexes[d] = (
            {(str(r.교사명).strip(), safe_int(r.교시)): r for r in day_tt.itertuples(index=False)}
            if not day_tt.empty else {}
        )
    rows=[]
    for t in sorted(x for x in teacher_names if x):
        row={"교사명":t}
        for d in DAYS:
            idx = daily_indexes[d]
            for p in range(1,PERIODS_PER_DAY.get(d,7)+1):
                r=idx.get((t,p))
                if r is None: row[f"{d}{p}"]=""; continue
                cell=f"{r.학급} {r.과목}".strip(); typ=str(getattr(r,"변경유형", "원본")).strip()
                if typ=="교환": cell += " 🔄 교환"
                elif typ=="테스트교환": cell += " 🧪 테스트교환"
                elif typ=="보강": cell += " 🟢 보강"
                elif typ=="시간강사": cell += f" 🟡 {r.원본교사}→시간강사"
                row[f"{d}{p}"]=cell
        rows.append(row)
    return pd.DataFrame(rows)
@st.cache_data(show_spinner=False)
def class_matrix(version=0, ref_date=None, use_test=False):
    ref=ref_date or _today_kst(); monday=ref-timedelta(days=ref.weekday())
    daily={}
    classes=set()
    for i,d in enumerate(DAYS):
        ds=(monday+timedelta(days=i)).strftime("%Y-%m-%d")
        e=get_effective_timetable_for_date(ds,version,use_test=use_test); daily[d]=e
        if not e.empty: classes.update(e["학급"].dropna().astype(str).str.strip())
    daily_indexes = {}
    for d, e in daily.items():
        idx = {}
        if not e.empty:
            for r in e.itertuples(index=False):
                cls = str(getattr(r, "학급", "")).strip()
                if not cls:
                    continue
                idx[(cls, safe_int(getattr(r, "교시", 0)))] = r
        daily_indexes[d] = idx
    rows=[]
    for c in sorted(x for x in classes if x):
        row={"학급":c}
        for d in DAYS:
            idx = daily_indexes[d]
            for p in range(1,PERIODS_PER_DAY.get(d,7)+1):
                r=idx.get((c,p))
                if r is not None:
                    cell = str(getattr(r, "과목", "")).strip()
                    typ = str(getattr(r, "변경유형", "원본")).strip()
                    marker = {
                        "교환": "🔄",
                        "테스트교환": "🧪",
                        "보강": "🟢",
                        "시간강사": "🟡",
                    }.get(typ, "")
                    row[f"{d}{p}"] = f"{cell} {marker}".strip()
                else:
                    row[f"{d}{p}"]=""
        rows.append(row)
    return pd.DataFrame(rows)
def _weekly_cell_parts(value):
    text = "" if value is None else str(value).strip()
    if not text:
        return "", "", "원본", ""
    marker = "원본"
    if "🧪" in text or "테스트교환" in text:
        marker = "테스트교환"
    elif "🔄" in text or "교환" in text:
        marker = "교환"
    elif "🟢" in text or "보강" in text:
        marker = "보강"
    elif "🟡" in text or "시간강사" in text:
        marker = "시간강사"
    clean = text.replace("🔄 교환", "").replace("🧪 테스트교환", "")
    clean = clean.replace("🟢 보강", "").replace("🟡 시간강사", "")
    clean = clean.replace(" 🔄", "").replace(" 🧪", "").replace(" 🟢", "")
    if "🟡" in clean:
        clean = clean.split("🟡", 1)[0].strip()
    if "[결강]" in clean:
        clean = clean.replace("[결강]", "").strip()
    parts = clean.split(None, 1)
    cls = parts[0] if parts else ""
    subject = parts[1] if len(parts) > 1 else ""
    icon = {"교환":"🔄", "테스트교환":"🧪", "보강":"🟢", "시간강사":"🟡"}.get(marker, "")
    return cls, subject, marker, icon
def _resolve_weekly_selection(selection, ref_date, use_test=False):
    if not selection:
        return None
    monday = ref_date - timedelta(days=ref_date.weekday())
    picked_date = monday + timedelta(days=selection["day_index"])
    ds = picked_date.strftime("%Y-%m-%d")
    ver = st.session_state.get("_data_version", 0)
    e = get_effective_timetable_for_date(ds, ver, use_test=use_test)
    if e.empty:
        return None
    p = selection["period"]
    if selection["row_label"] == "교사명":
        m = e[(e["교사명"].astype(str).str.strip() == selection["row_name"]) & (e["교시"].apply(safe_int) == p)]
    else:
        m = e[(e["학급"].astype(str).str.strip() == selection["row_name"]) & (e["교시"].apply(safe_int) == p)]
    if m.empty:
        return None
    r = m.iloc[0]
    return {
        "교사명": str(r.get("교사명", "")).strip(), "일자": ds, "요일": DAYS[selection["day_index"]],
        "교시": p, "학급": str(r.get("학급", "")).strip(), "과목": str(r.get("과목", "")).strip(),
        "변경유형": str(r.get("변경유형", "원본")).strip(), "변경출처": str(r.get("변경출처", "")).strip(),
        "변경ID": str(r.get("변경ID", "")).strip(), "원본교사": str(r.get("원본교사", r.get("교사명", ""))).strip(),
        "원본일자": str(r.get("원본일자", ds)), "원본교시": safe_int(r.get("원본교시", p)),
    }
def _register_absence_from_weekly(lesson, reason, detail=""):
    if not lesson:
        return False
    on_date, teacher, p = lesson["일자"], lesson["교사명"], safe_int(lesson["교시"])
    existing = st.session_state.absences
    if not existing.empty:
        dup = existing[(existing["일자"].astype(str) == on_date) & (existing["교사명"].astype(str).str.strip() == teacher) & (existing["교시"].apply(safe_int) == p)]
        if not dup.empty:
            st.warning("이미 등록된 결강입니다.")
            return False
    cid = f"{on_date}-{teacher}"
    new = pd.DataFrame([{
        "결강ID": cid, "일자": on_date, "요일": lesson["요일"], "교사명": teacher,
        "사유": reason, "상세사유": detail, "교시": p, "학급": lesson["학급"], "과목": lesson["과목"],
        "등록시각": _now_text("%Y-%m-%d %H:%M"), "입력자": current_user()
    }])
    before_absences = existing.copy(deep=True)
    st.session_state.absences = pd.concat([before_absences, new], ignore_index=True)
    if not save_work_data_to_gsheet(["결강"]):
        st.session_state.absences = before_absences
        _invalidate_all_caches()
        return False
    _invalidate_all_caches()
    push_history(f"주간표에서 결강 등록 ({teacher} {p}교시)")
    return True
def _weekly_display_matrix(matrix: pd.DataFrame) -> pd.DataFrame:
    if matrix is None or matrix.empty:
        return matrix.copy(deep=True) if isinstance(matrix, pd.DataFrame) else pd.DataFrame()
    out = matrix.copy(deep=True)
    for col in out.columns:
        if str(col) in ("교사명", "학급", "교시"):
            continue
        vals = []
        for raw in out[col].tolist():
            if "학급" in out.columns and "교사명" not in out.columns:
                vals.append("" if raw is None else str(raw).strip())
                continue
            cls, subject, marker, icon = _weekly_cell_parts(raw)
            if not cls and not subject:
                vals.append("")
            else:
                text = " ".join([x for x in (cls, subject, icon) if x])
                vals.append(text)
        out[col] = vals
    return out
_CHANGED_LESSON_MARKERS = ("🔄", "🧪", "🟢", "🟡")

def _is_changed_lesson_cell(value) -> bool:
    text = "" if value is None else str(value)
    return any(marker in text for marker in _CHANGED_LESSON_MARKERS)

def _daily_styled_matrix(matrix: pd.DataFrame):
    if matrix is None or matrix.empty:
        return matrix
    styler = matrix.style
    data_cols = [c for c in matrix.columns if str(c) != "교사명"]
    if data_cols:
        def _changed_text_style(v):
            return "color:#c62828;font-weight:700;" if _is_changed_lesson_cell(v) else ""
        styler = styler.map(_changed_text_style, subset=data_cols) if hasattr(styler, "map") else styler.applymap(_changed_text_style, subset=data_cols)
    return styler

def _weekly_styled_matrix(matrix: pd.DataFrame, *, drop_teacher_name: bool = False):
    display = _weekly_display_matrix(matrix)
    if display is None or display.empty:
        return display
    if drop_teacher_name and "교사명" in display.columns:
        display = display.drop(columns=["교사명"])
    styler = display.style
    styler = styler.set_properties(**{
        "text-align": "center",
        "vertical-align": "middle",
        "white-space": "nowrap",
        "line-height": "1.0",
        "font-size": "10.5px",
        "padding": "2px 1px",
    })
    if "교사명" in display.columns:
        styler = styler.set_properties(subset=["교사명"], **{
            "font-weight": "700", "text-align": "left", "white-space": "nowrap"
        })
    if "학급" in display.columns:
        styler = styler.set_properties(subset=["학급"], **{
            "font-weight": "700", "text-align": "left", "white-space": "nowrap"
        })
    day_rgba = [
        "rgba(59,130,246,.035)", "rgba(16,185,129,.035)", "rgba(245,158,11,.040)",
        "rgba(139,92,246,.035)", "rgba(236,72,153,.035)"
    ]
    for di, day in enumerate(DAYS):
        cols = [f"{day}{p}" for p in range(1, MAX_PERIOD + 1) if f"{day}{p}" in display.columns]
        if not cols:
            continue
        styler = styler.set_properties(subset=cols, **{"background-color": day_rgba[di]})
        first = cols[0]
        styler = styler.set_properties(subset=[first], **{"border-left": "3px solid rgba(71,85,105,.42)"})
    def _status_style(v):
        s = "" if v is None else str(v)
        if "🔄" in s:
            return "color:#c62828; box-shadow: inset 0 0 0 2px rgba(239,68,68,.82); font-weight:700;"
        if "🧪" in s:
            return "color:#c62828; box-shadow: inset 0 0 0 2px rgba(124,58,237,.82); font-weight:700;"
        if "🟢" in s:
            return "color:#c62828; box-shadow: inset 0 0 0 2px rgba(22,163,74,.82); font-weight:700;"
        if "🟡" in s:
            return "color:#c62828; box-shadow: inset 0 0 0 2px rgba(217,119,6,.82); font-weight:700;"
        return ""
    data_cols = [c for c in display.columns if str(c) not in ("교사명", "학급")]
    if data_cols:
        styler = styler.map(_status_style, subset=data_cols) if hasattr(styler, "map") else styler.applymap(_status_style, subset=data_cols)
    return styler
def _clear_weekly_selection():
    st.session_state["weekly_matrix_epoch"] = int(st.session_state.get("weekly_matrix_epoch", 0) or 0) + 1
    st.session_state.pop("weekly_selected_lesson", None)
    st.session_state.pop("weekly_swap_source", None)
    st.session_state.pop("weekly_dialog_action_mode", None)
    st.session_state.pop("weekly_dialog_open", None)
    st.session_state.pop("weekly_swap_candidates", None)
    st.session_state.pop("weekly_swap_candidates_key", None)
    st.session_state.pop("weekly_dialog_swap_selected_row", None)
    st.session_state.pop("weekly_cycle_candidates", None)
    st.session_state.pop("weekly_cycle_candidates_msg", None)
    st.session_state.pop("weekly_cycle_candidates_key", None)
def _validate_current_weekly_selection(lesson, *, use_test=False):
    if not isinstance(lesson, dict):
        return False
    try:
        ds = normalize_date_str(lesson.get("일자", ""))
        period = safe_int(lesson.get("교시", 0))
        teacher = str(lesson.get("교사명", "")).strip()
        klass = str(lesson.get("학급", "")).strip()
        subject = str(lesson.get("과목", "")).strip()
        if not ds or not teacher or period <= 0 or not klass or not subject:
            return False
        ver = st.session_state.get("_data_version", 0)
        e = get_effective_timetable_for_date(ds, ver, use_test=bool(use_test))
        if e is None or e.empty:
            return False
        m = e[(e["교사명"].astype(str).str.strip() == teacher)
               & (e["교시"].apply(safe_int) == period)
               & (e["학급"].astype(str).str.strip() == klass)
               & (e["과목"].astype(str).str.strip() == subject)]
        return not m.empty
    except Exception:
        return False
def _filter_current_swap_candidates(df, lesson, *, use_test=False):
    if df is None or df.empty or not isinstance(lesson, dict):
        return pd.DataFrame(columns=list(df.columns) if isinstance(df, pd.DataFrame) else [])
    ver = st.session_state.get("_data_version", 0)
    source_date = normalize_date_str(lesson.get("일자", ""))
    source_period = safe_int(lesson.get("교시", 0))
    source_teacher = str(lesson.get("교사명", "")).strip()
    source_class = str(lesson.get("학급", "")).strip()
    source_subject = str(lesson.get("과목", "")).strip()
    source_e = get_effective_timetable_for_date(source_date, ver, use_test=bool(use_test))
    if source_e is None or source_e.empty:
        return df.iloc[0:0].copy()
    source_ok = not source_e[(source_e["교사명"].astype(str).str.strip() == source_teacher)
                              & (source_e["교시"].apply(safe_int) == source_period)
                              & (source_e["학급"].astype(str).str.strip() == source_class)
                              & (source_e["과목"].astype(str).str.strip() == source_subject)].empty
    if not source_ok:
        return df.iloc[0:0].copy()
    valid_rows = []
    target_cache = {}
    for idx, r in df.iterrows():
        td = normalize_date_str(r.get("이동희망일", ""))
        tp = safe_int(r.get("이동희망교시", r.get("원본교시", source_period)))
        tt = str(r.get("상대교사", "")).strip()
        tc = str(r.get("상대학급", "")).strip()
        ts = str(r.get("상대과목", "")).strip()
        if (not td or not tt or tp <= 0 or not tc or not ts
                or td < source_date):
            continue
        if td not in target_cache:
            target_cache[td] = get_effective_timetable_for_date(td, ver, use_test=bool(use_test))
        te = target_cache[td]
        if te is None or te.empty:
            continue
        m = te[(te["교사명"].astype(str).str.strip() == tt)
               & (te["교시"].apply(safe_int) == tp)
               & (te["학급"].astype(str).str.strip() == tc)
               & (te["과목"].astype(str).str.strip() == ts)]
        if m.empty:
            continue
        target_day = WEEKDAY_KR[datetime.strptime(td, "%Y-%m-%d").weekday()]
        if not is_free(source_teacher, target_day, tp, td, te):
            continue
        if not teacher_slot_is_available(source_teacher, target_day, tp, ver):
            continue
        source_day_check=WEEKDAY_KR[datetime.strptime(source_date, "%Y-%m-%d").weekday()]
        if not is_free(tt, source_day_check, source_period, source_date, source_e):
            continue
        if not teacher_slot_is_available(tt, source_day_check, source_period, ver):
            continue
        valid_rows.append(idx)
    return df.loc[valid_rows].reset_index(drop=True)
def _weekly_fragment_rerun():
    try:
        st.rerun(scope="fragment")
    except TypeError:
        st.rerun()
@contextmanager
def _weekly_dialog_loading(label: str):
    with st.spinner(f"🔄 로딩 중... {label}"):
        yield
@st.dialog("🎯 수업 작업", width="large")
def _weekly_action_dialog():
    lesson = st.session_state.get("weekly_selected_lesson")
    if not lesson:
        st.info("선택한 수업이 없습니다.")
        return
    use_test = bool(st.session_state.get("weekly_dialog_use_test", False))
    if not _validate_current_weekly_selection(lesson, use_test=use_test):
        _clear_weekly_selection()
        _weekly_fragment_rerun()
    title = st.session_state.get("weekly_dialog_title", "주간표 작업")
    status = lesson.get("변경유형") or "원본"
    ver = st.session_state.get("_data_version", 0)
    st.markdown(f"### {title}")
    st.info(
        f"**{lesson['일자']} ({lesson['요일']}) · {lesson['교사명']} · "
        f"{lesson['교시']}교시 · {lesson['학급']} · {lesson['과목']}**\n\n"
        f"현재 상태: **{status}**"
    )
    if st.session_state.get("weekly_dialog_result"):
        st.success(st.session_state.weekly_dialog_result)
        st.caption("시간표가 갱신되었습니다. 팝업을 닫으면 최신 주간표를 확인할 수 있습니다.")
        if st.button("✖ 닫기", key="dlg_result_close", width="stretch"):
            st.session_state.pop("weekly_dialog_result", None)
            _clear_weekly_selection()
            _weekly_fragment_rerun()
        return
    action_mode = st.session_state.get("weekly_dialog_action_mode", "swap")
    if action_mode not in {"swap", "target", "absence", "substitute", "detail"}:
        action_mode = "swap"
        st.session_state.weekly_dialog_action_mode = "swap"
    if action_mode == "swap":
        st.markdown("#### 🔄 1:1 기본 맞교환")
        st.caption("가장 자주 사용하는 1:1 맞교환을 기본 화면으로 표시합니다. 특정 날짜·교시를 지정하면 그 시간에 1:1이 없을 때 자동으로 연계 순환을 탐색합니다.")
    else:
        nav_cols = st.columns(5)
        nav_items = [
            ("swap", "🔄 1:1 맞교환"),
            ("target", "📅 날짜·교시 지정"),
            ("absence", "📌 결강"),
            ("substitute", "🟢 보강"),
            ("detail", "ℹ️ 상세"),
        ]
        for col, (mode, label) in zip(nav_cols, nav_items):
            with col:
                if st.button(label, type="primary" if mode == action_mode else "secondary",
                             key=f"dlg_action_{mode}", width="stretch"):
                    st.session_state.weekly_dialog_action_mode = mode
                    _weekly_fragment_rerun()
        st.markdown(f"#### {dict(nav_items)[action_mode]}")
    if action_mode == "swap":
        alt_cols = st.columns(4)
        alt_items = [("target", "📅 날짜·교시 지정"), ("absence", "📌 결강"),
                     ("substitute", "🟢 보강"), ("detail", "ℹ️ 상세")]
        for col, (mode, label) in zip(alt_cols, alt_items):
            with col:
                if st.button(label, type="secondary", key=f"dlg_alt_{mode}", width="stretch"):
                    st.session_state.weekly_dialog_action_mode = mode
                    _weekly_fragment_rerun()
    if status != "원본" and action_mode == "detail":
        st.caption(
            f"변경출처: {lesson.get('변경출처') or '-'} · 변경ID: {lesson.get('변경ID') or '-'} · "
            f"원본: {lesson.get('원본교사') or '-'} / {lesson.get('원본일자') or '-'} / "
            f"{lesson.get('원본교시') or '-'}교시"
        )
    extra_days = int(st.session_state.get("weekly_dialog_extra_days", 7))
    if action_mode in ("swap", "cycle"):
        st.markdown('<div class="weekly-search-range">🔎 <strong>검색 범위</strong><span> · 기본 미래 7일 추가</span></div>', unsafe_allow_html=True)
        extra_days = st.slider(
            "미래 추가 검색 일수", 0, 21, extra_days,
            key="weekly_dialog_extra_days_input",
            help="미래 날짜를 추가로 검색할 범위입니다.",
            label_visibility="collapsed",
        )
        st.caption(f"미래 {extra_days}일 추가 검색")
        if extra_days != st.session_state.get("weekly_dialog_extra_days"):
            st.session_state.weekly_dialog_extra_days = extra_days
            st.session_state.pop("weekly_swap_candidates_key", None)
            st.session_state.pop("weekly_cycle_candidates_key", None)
            _weekly_fragment_rerun()
    if action_mode == "swap":
        cache_key = (
            "same-class-only-v2",
            str(lesson.get("교사명", "")), str(lesson.get("일자", "")), safe_int(lesson.get("교시", 0)),
            str(lesson.get("학급", "")), str(lesson.get("과목", "")), int(extra_days), int(ver), bool(use_test)
        )
        stored_key = st.session_state.get("weekly_swap_candidates_key")
        if stored_key != cache_key:
            st.session_state.pop("weekly_dialog_swap_selected_row", None)
            with _weekly_dialog_loading("1:1 교환 후보 검색 중"):
                df_swap = get_single_lesson_1to1_candidates(
                    lesson["교사명"], lesson["일자"], safe_int(lesson["교시"]),
                    str(lesson["학급"]), str(lesson["과목"]),
                    future_days=extra_days, version=ver, use_test=use_test,
                )
            st.session_state.weekly_swap_candidates = df_swap
            st.session_state.weekly_swap_candidates_key = cache_key
        else:
            df_swap = st.session_state.get("weekly_swap_candidates", pd.DataFrame())
        df_swap = _filter_current_swap_candidates(df_swap, lesson, use_test=use_test)
        if not df_swap.empty and "동일학급" in df_swap.columns:
            df_swap = df_swap[df_swap["동일학급"].astype(str).str.strip() == "🏆"].copy()
        if df_swap.empty:
            st.info("현재 같은 학급 조건에서 가능한 1:1 맞교환 위치가 없습니다.")
        else:
            same_df = df_swap.reset_index(drop=True).copy()
            st.markdown(
                f'<div class="swap-result-summary"><strong>{len(same_df)}개</strong> 동일 학급 1:1 교환 가능</div>',
                unsafe_allow_html=True,
            )
            def _swap_result_view(src):
                view = src.copy()
                view["날짜"] = view.apply(lambda r: f"{r['이동희망일']} ({r['이동요일']})", axis=1)
                view["교시"] = view["이동희망교시"].apply(lambda x: f"{safe_int(x)}교시")
                view = view.rename(columns={
                    "상대교사": "상대교사", "상대학급": "상대학급", "상대과목": "상대과목",
                    "교환가능사유": "교환 가능 사유",
                })
                return view[["날짜", "교시", "상대교사", "상대학급", "상대과목", "교환 가능 사유"]]
            st.markdown("#### 교환할 수업 선택")
            st.caption("아래 표에서 교환할 수업의 행을 클릭하면 바로 선택됩니다. 별도의 선택 목록은 없습니다.")
            table_view = _swap_result_view(same_df)
            dialog_instance = int(st.session_state.get("weekly_dialog_instance", 0) or 0)
            table_key = f"weekly_dialog_swap_table_{dialog_instance}"
            selected_row_key = "weekly_dialog_swap_selected_row"
            table_event = st.dataframe(
                table_view, width="stretch", hide_index=True,
                height=min(520, 44 + max(1, len(table_view)) * 35), key=table_key,
                on_select="rerun", selection_mode="single-row",
                column_config={
                    "날짜": st.column_config.TextColumn("날짜", width="small"),
                    "교시": st.column_config.TextColumn("교시", width="small"),
                    "상대교사": st.column_config.TextColumn("상대 교사", width="small"),
                    "상대학급": st.column_config.TextColumn("상대 학급", width="small"),
                    "상대과목": st.column_config.TextColumn("상대 과목", width="small"),
                    "교환 가능 사유": st.column_config.TextColumn("교환 가능 사유", width="large"),
                },
            )
            selected_rows = list(getattr(getattr(table_event, "selection", None), "rows", []) or [])
            if selected_rows:
                row_idx = int(selected_rows[0])
                if 0 <= row_idx < len(same_df):
                    st.session_state[selected_row_key] = row_idx
            selected_row = st.session_state.get(selected_row_key)
            if selected_row is not None and not (0 <= int(selected_row) < len(same_df)):
                selected_row = None
                st.session_state.pop(selected_row_key, None)
            picked = same_df.iloc[int(selected_row)] if selected_row is not None else None
            if picked is None:
                st.info("교환하려는 수업을 위 표에서 클릭해 주세요.")
            else:
                st.success(
                    f"선택됨: **{picked['상대교사']} · {picked['이동희망일']} · "
                    f"{safe_int(picked.get('이동희망교시', picked['원본교시']))}교시 · "
                    f"{picked['상대학급']} · {picked['상대과목']}**"
                )
                b_info = {
                    "교사명": str(picked["상대교사"]), "일자": str(picked["이동희망일"]),
                    "요일": str(picked["이동요일"]), "교시": safe_int(picked.get("이동희망교시", picked["원본교시"])),
                    "학급": str(picked["상대학급"]), "과목": str(picked["상대과목"]),
                }
            button_label = "🧪 1:1 맞교환 테스트" if use_test else "✅ 1:1 맞교환 실행"
            if st.button(button_label, type="primary", key="dlg_direct_swap", width="stretch", disabled=picked is None):
                try:
                    with _weekly_dialog_loading("1:1 맞교환 처리 중"):
                        ok = do_swap(lesson, b_info, lesson["일자"], b_info["일자"], is_test=use_test)
                except Exception as exc:
                    st.error(f"맞교환 처리 중 오류가 발생했습니다: {exc}")
                    ok = False
                if ok:
                    st.success("테스트 맞교환이 적용되었습니다." if use_test else "1:1 맞교환이 반영되었습니다.")
                    st.session_state.weekly_dialog_result = (
                        "테스트 맞교환이 적용되었습니다." if use_test else "1:1 맞교환이 반영되었습니다."
                    )
                    if use_test:
                        st.rerun(scope="fragment")
                    else:
                        _weekly_fragment_rerun()
                else:
                    st.error("현재 상태에서는 이 1:1 맞교환을 적용할 수 없습니다. 최신 시간표 상태를 다시 확인해 주세요.")
    elif action_mode == "target":
        st.session_state["weekly_dialog_open"] = True
        target_date_key = "weekly_dialog_target_date"
        target_period_key = "weekly_dialog_target_period"
        default_target = st.session_state.get(target_date_key)
        if not default_target:
            try:
                default_target = normalize_date_str(lesson.get("일자", ""))
                default_target = datetime.strptime(default_target, "%Y-%m-%d").date()
            except Exception:
                default_target = _today_kst()
        elif isinstance(default_target, str):
            try:
                default_target = datetime.strptime(normalize_date_str(default_target), "%Y-%m-%d").date()
            except Exception:
                default_target = _today_kst()
        if default_target.weekday() >= 5:
            default_target -= timedelta(days=default_target.weekday() - 4)
        previous_target_date = st.session_state.get(target_date_key)
        target_date = calendar_picker("교환 희망일", default_target, key="weekly_dialog_target_calendar", rerun_scope="fragment")
        st.session_state[target_date_key] = target_date
        if previous_target_date and normalize_date_str(previous_target_date) != target_date.strftime("%Y-%m-%d"):
            st.session_state.weekly_target_search_requested = False
            st.session_state.weekly_target_search_dirty_key = None
            st.session_state.weekly_target_candidates_key = None
            st.session_state.weekly_target_swap_candidates = pd.DataFrame()
            st.session_state.weekly_target_cycle_candidates = None
            st.session_state.weekly_target_cycle_msg = ""
            st.session_state.pop("weekly_target_swap_selected_row", None)
        saved_period = safe_int(st.session_state.get(target_period_key, 0))
        max_target_period = PERIODS_PER_DAY.get(WEEKDAY_KR[target_date.weekday()], MAX_PERIOD)
        if saved_period > max_target_period:
            saved_period = 0
            st.session_state[target_period_key] = 0
        target_free_periods = _teacher_empty_periods_on_date(
            lesson.get("교사명", ""), target_date, version=ver, use_test=use_test
        )
        if saved_period and saved_period not in target_free_periods:
            saved_period = 0
            st.session_state[target_period_key] = 0
        selected_periods = period_matrix_picker(
            "교환 희망 교시", "weekly_dialog_target_period_picker",
            selected=[saved_period] if saved_period else [], allow_all=False, rerun_scope="fragment",
            single=True, state_key=target_period_key, allowed_periods=target_free_periods
        )
        target_period = safe_int(selected_periods[0]) if selected_periods else 0
        st.session_state[target_period_key] = target_period
        if not target_free_periods:
            st.warning("선택한 교사가 이 날짜에는 모든 교시에 수업이 있어 교환 가능한 공강 교시가 없습니다.")
        elif target_period <= 0:
            st.info("교환 희망 교시를 선택해 주세요.")
        else:
            target_date_str = target_date.strftime("%Y-%m-%d")
            source_date_str = normalize_date_str(lesson.get("일자", ""))
            source_period = safe_int(lesson.get("교시", 0))
            if target_date_str == source_date_str and target_period == source_period:
                st.warning("원본 수업과 같은 날짜·교시는 선택할 수 없습니다.")
            else:
                st.markdown(
                    f'<div class="swap-result-summary"><strong>지정 시간</strong> · {target_date_str} ({WEEKDAY_KR[target_date.weekday()]}) · {target_period}교시</div>',
                    unsafe_allow_html=True,
                )
                target_cache_key = (
                    "target-first-fallback-v1", str(lesson.get("교사명", "")), source_date_str, source_period,
                    str(lesson.get("학급", "")), str(lesson.get("과목", "")), target_date_str, target_period, int(ver), bool(use_test)
                )
                stored_target_key = st.session_state.get("weekly_target_candidates_key")
                search_requested = bool(st.session_state.get("weekly_target_search_requested", False))
                search_dirty_key = st.session_state.get("weekly_target_search_dirty_key")
                if search_dirty_key != target_cache_key:
                    st.session_state.weekly_target_search_dirty_key = target_cache_key
                    st.session_state.weekly_target_search_requested = False
                    st.session_state.pop("weekly_target_swap_selected_row", None)
                    st.session_state.pop("weekly_target_candidates_key", None)
                    st.session_state.weekly_target_swap_candidates = pd.DataFrame()
                    st.session_state.weekly_target_cycle_candidates = None
                    st.session_state.weekly_target_cycle_msg = ""
                    search_requested = False
                search_button_label = "🔎 이 날짜·교시로 교환 찾기"
                if st.button(search_button_label, type="primary", key="dlg_target_search", width="stretch"):
                    st.session_state.weekly_target_search_requested = True
                    search_requested = True
                    st.session_state.pop("weekly_target_swap_selected_row", None)
                    st.session_state.weekly_target_candidates_key = None
                    st.session_state.weekly_target_cycle_candidates = None
                    st.session_state.weekly_target_cycle_msg = ""
                    st.rerun(scope="fragment")
                if not search_requested:
                    st.caption("날짜와 교시를 선택한 뒤 위 버튼을 누르면 해당 시간의 1:1 교환을 먼저 확인합니다. 없으면 연계 순환을 자동으로 탐색합니다.")
                    return
                stored_target_key = st.session_state.get("weekly_target_candidates_key")
                if stored_target_key != target_cache_key:
                    st.session_state.pop("weekly_target_swap_selected_row", None)
                    with _weekly_dialog_loading("지정 시간의 1:1 교환 가능 여부 확인 중"):
                        df_target = get_single_lesson_1to1_candidates(
                            lesson["교사명"], lesson["일자"], source_period,
                            str(lesson["학급"]), str(lesson["과목"]),
                            future_days=0, version=ver, use_test=use_test,
                            target_date_str=target_date_str, target_period=target_period,
                        )
                    if not df_target.empty:
                        df_target = df_target[
                            (df_target["이동희망일"].astype(str) == target_date_str) &
                            (df_target["이동희망교시"].apply(safe_int) == target_period)
                        ].reset_index(drop=True)
                    st.session_state.weekly_target_swap_candidates = df_target
                    st.session_state.weekly_target_candidates_key = target_cache_key
                    st.session_state.weekly_target_cycle_candidates = None
                    st.session_state.weekly_target_cycle_msg = ""
                else:
                    df_target = st.session_state.get("weekly_target_swap_candidates", pd.DataFrame())

                if not df_target.empty:
                    st.success(f"지정한 시간에 1:1 교환 가능 수업이 {len(df_target)}개 있습니다.")
                    st.markdown("#### 교환할 수업 선택")
                    st.caption("표에서 교환할 수업을 클릭하면 바로 선택됩니다.")
                    table_view = _swap_result_view(df_target) if "_swap_result_view" in locals() else df_target[[
                        "이동희망일","이동희망교시","상대교사","상대학급","상대과목","교환가능사유"
                    ]].rename(columns={
                        "이동희망일":"날짜","이동희망교시":"교시","교환가능사유":"교환 가능 사유"
                    })
                    table_view["날짜"] = df_target["이동희망일"].map(lambda x: f"{x} ({WEEKDAY_KR[datetime.strptime(str(x), '%Y-%m-%d').weekday()]})")
                    table_view["교시"] = df_target["이동희망교시"].apply(lambda x: f"{safe_int(x)}교시")
                    table_view = table_view[["날짜","교시","상대교사","상대학급","상대과목","교환 가능 사유"]]
                    selected_rows = []
                    target_table_key = f"weekly_dialog_target_swap_table_{int(st.session_state.get('weekly_dialog_instance',0) or 0)}"
                    event = st.dataframe(
                        table_view, width="stretch", hide_index=True, key=target_table_key,
                        on_select="rerun", selection_mode="single-row",
                        column_config={
                            "날짜": st.column_config.TextColumn("날짜", width="small"),
                            "교시": st.column_config.TextColumn("교시", width="small"),
                            "상대교사": st.column_config.TextColumn("상대 교사", width="small"),
                            "상대학급": st.column_config.TextColumn("상대 학급", width="small"),
                            "상대과목": st.column_config.TextColumn("상대 과목", width="small"),
                            "교환 가능 사유": st.column_config.TextColumn("교환 가능 사유", width="large"),
                        },
                    )
                    selected_rows = list(getattr(getattr(event,"selection",None),"rows",[]) or [])
                    if selected_rows:
                        st.session_state.weekly_target_swap_selected_row = int(selected_rows[0])
                    picked_idx = st.session_state.get("weekly_target_swap_selected_row")
                    picked = df_target.iloc[int(picked_idx)] if picked_idx is not None and 0 <= int(picked_idx) < len(df_target) else None
                    if picked is not None and (str(picked.get("이동희망일", "")) != target_date_str or safe_int(picked.get("이동희망교시", 0)) != target_period):
                        picked = None
                        st.session_state.pop("weekly_target_swap_selected_row", None)
                    if picked is not None:
                        st.success(
                            f"선택됨: **{picked['상대교사']} · {picked['이동희망일']} · {safe_int(picked['이동희망교시'])}교시 · {picked['상대학급']} · {picked['상대과목']}**"
                        )
                        b_info = {
                            "교사명": str(picked["상대교사"]), "일자": str(picked["이동희망일"]),
                            "요일": str(picked["이동요일"]), "교시": safe_int(picked["이동희망교시"]),
                            "학급": str(picked["상대학급"]), "과목": str(picked["상대과목"]),
                        }
                        target_button_label = "🧪 지정 시간 1:1 맞교환 테스트" if use_test else "✅ 지정 시간 1:1 맞교환 실행"
                        if st.button(target_button_label, type="primary", key="dlg_target_direct_swap", width="stretch"):
                            try:
                                with _weekly_dialog_loading("지정 시간 1:1 맞교환 처리 중"):
                                    ok = do_swap(lesson, b_info, lesson["일자"], b_info["일자"], is_test=use_test)
                            except Exception as exc:
                                st.error(f"맞교환 처리 중 오류가 발생했습니다: {exc}")
                                ok = False
                            if ok:
                                st.session_state.weekly_dialog_result = "테스트 맞교환이 적용되었습니다." if use_test else "1:1 맞교환이 반영되었습니다."
                                st.rerun(scope="fragment")
                    else:
                        st.info("교환하려는 수업을 위 표에서 클릭해 주세요.")
                else:
                    st.warning("지정한 날짜·교시에 가능한 1:1 교환이 없습니다. 같은 시간대를 대상으로 연계 순환을 탐색합니다.")
                    target_cycle_key = target_cache_key + ("cycle",)
                    stored_cycle_key = st.session_state.get("weekly_target_cycle_key")
                    if stored_cycle_key != target_cycle_key:
                        with _weekly_dialog_loading("지정 시간의 연계 순환 탐색 중"):
                            cycles, cycle_msg = find_cycle_linked_swaps(
                                lesson["교사명"], lesson["일자"], source_period,
                                str(lesson["학급"]), str(lesson["과목"]),
                                target_date_str, target_period,
                                min_cycle=2, max_cycle=3, future_days=max(0, extra_days), version=ver, use_test=use_test
                            )
                        st.session_state.weekly_target_cycle_candidates = cycles
                        st.session_state.weekly_target_cycle_msg = cycle_msg
                        st.session_state.weekly_target_cycle_key = target_cycle_key
                    else:
                        cycles = st.session_state.get("weekly_target_cycle_candidates") or []
                        cycle_msg = st.session_state.get("weekly_target_cycle_msg", "")
                    st.caption(cycle_msg or "지정한 날짜·교시를 기준으로 연계 순환 가능성을 검사합니다.")
                    if not cycles:
                        st.info("지정한 날짜·교시에서 가능한 연계 순환 경로도 없습니다.")
                    else:
                        st.markdown("#### 🔗 지정 시간 연계 순환")
                        for idx, cyc in enumerate(cycles[:6]):
                            with st.container(border=True):
                                st.markdown(
                                    f"**{'🔗' if cyc.get('length', 0) > 2 else '↔️'} {cyc.get('length','')}인 순환**"
                                )
                                st.caption(cyc.get("path_desc", ""))
                                st.caption("지정한 날짜·교시에 1:1 교환이 없어서 연계 순환으로 찾은 경로입니다.")
                                cycle_button = "🧪 이 연계 순환 테스트" if use_test else "✅ 이 연계 순환 실행"
                                if st.button(cycle_button, key=f"dlg_target_cycle_{idx}", type="primary", width="stretch"):
                                    try:
                                        with _weekly_dialog_loading("지정 시간 연계 순환 처리 중"):
                                            ok = apply_cycle_swaps(cyc["moves"], is_test=True if use_test else False)
                                    except Exception as exc:
                                        st.error(f"연계 순환 처리 중 오류가 발생했습니다: {exc}")
                                        ok = False
                                    if ok is not False:
                                        st.session_state.weekly_dialog_result = (
                                            f"테스트 {cyc.get('length','')}인 연계 순환이 적용되었습니다. 실제 저장되지는 않습니다."
                                            if use_test else
                                            f"{cyc.get('length','')}인 연계 순환이 적용되었습니다."
                                        )
                                        if use_test:
                                            st.rerun(scope="fragment")
                                        else:
                                            _weekly_fragment_rerun()
    elif action_mode == "absence":
        r1, r2 = st.columns([1, 2])
        with r1:
            reason = st.selectbox("사유", ABSENCE_REASONS, key="dlg_abs_reason")
        with r2:
            detail = st.text_input("상세사유", key="dlg_abs_detail")
        if st.button("📌 결강 등록", type="primary", key="dlg_abs_submit", width="stretch"):
            try:
                with _weekly_dialog_loading("결강 정보 저장 중"):
                    ok = _register_absence_from_weekly(lesson, reason, detail)
            except Exception as exc:
                st.error(f"결강 등록 중 오류가 발생했습니다: {exc}")
                ok = False
            if ok:
                st.success("결강이 등록되었습니다.")
                st.session_state.weekly_dialog_result = "결강이 등록되었습니다."
                _weekly_fragment_rerun()
    elif action_mode == "substitute":
        with _weekly_dialog_loading("보강 후보 확인 중"):
            cand = get_cached_substitute_recommendations(
                lesson["요일"], lesson["교시"], lesson["과목"], lesson["학급"], lesson["교사명"], lesson["일자"],
                top_n=10, include_part_time=True, version=ver
            )
        if cand.empty:
            st.warning("현재 조건에서 추천 가능한 보강 교사가 없습니다.")
        else:
            st.dataframe(cand, width="stretch", hide_index=True, height=240)
            abs_df = st.session_state.absences
            abs_match = (
                abs_df[(abs_df["일자"].astype(str) == lesson["일자"]) &
                       (abs_df["교사명"].astype(str).str.strip() == lesson["교사명"]) &
                       (abs_df["교시"].apply(safe_int) == lesson["교시"])]
                if not abs_df.empty else pd.DataFrame()
            )
            if not abs_match.empty:
                cid = str(abs_match.iloc[0]["결강ID"])
                labels = cand["보강교사"].astype(str).tolist()
                pick = st.selectbox("보강 교사", labels, key="dlg_sub_pick")
                picked = cand[cand["보강교사"].astype(str) == str(pick)].iloc[0]
                if st.button("🟢 선택 교사로 보강 배정", type="primary", key="dlg_sub_submit", width="stretch"):
                    try:
                        with _weekly_dialog_loading("보강 배정 처리 중"):
                            ok = add_substitute(
                                cid, lesson["일자"], lesson["요일"], lesson["교시"], lesson["학급"], lesson["과목"],
                                lesson["교사명"], str(picked["보강교사"]), "주간표", picked.get("우선순위", ""),
                                "주간 시간표 셀에서 배정"
                            )
                    except Exception as exc:
                        st.error(f"보강 배정 중 오류가 발생했습니다: {exc}")
                        ok = False
                    if ok:
                        st.success("보강이 배정되었습니다.")
                        st.session_state.weekly_dialog_result = "보강이 배정되었습니다."
                        _weekly_fragment_rerun()
            else:
                st.caption("이 수업에 등록된 결강이 없습니다. 결강 등록 후 바로 보강을 배정할 수 있습니다.")
    elif action_mode == "detail":
        detail_rows = [
            ("교사", lesson.get("교사명", "")), ("일자", lesson.get("일자", "")),
            ("요일", lesson.get("요일", "")), ("교시", lesson.get("교시", "")),
            ("학급", lesson.get("학급", "")), ("과목", lesson.get("과목", "")),
            ("변경유형", lesson.get("변경유형", "원본")), ("변경출처", lesson.get("변경출처", "")),
            ("변경ID", lesson.get("변경ID", "")), ("원본교사", lesson.get("원본교사", "")),
            ("원본일자", lesson.get("원본일자", "")), ("원본교시", lesson.get("원본교시", "")),
        ]
        st.dataframe(pd.DataFrame(detail_rows, columns=["항목", "내용"]), width="stretch", hide_index=True)
    if st.button("✖ 닫기", key="dlg_close", width="stretch"):
        st.session_state.pop("weekly_dialog_result", None)
        _clear_weekly_selection()
        _weekly_fragment_rerun()
def _resolve_matrix_cell_selection(matrix, ref_date, row_label, selected_cells, *, use_test=False):
    if not selected_cells:
        return None
    try:
        row_idx, column_name = selected_cells[0]
        row_idx = int(row_idx)
        column_name = str(column_name)
    except Exception:
        return None
    if row_idx < 0 or row_idx >= len(matrix) or column_name in ("교사명", "학급"):
        return None
    period_grid = row_label == "교시" and column_name in DAYS
    if period_grid:
        day = column_name
        period = safe_int(matrix.iloc[row_idx].get("교시", 0))
        row_name = str(matrix.iloc[row_idx].get("교사명", "")).strip()
    else:
        day = column_name[:1]
        period = safe_int(column_name[1:])
        row_name = str(matrix.iloc[row_idx].get(row_label, "")).strip()
    if day not in DAYS or not (1 <= period <= MAX_PERIOD):
        return None
    if not row_name or not str(matrix.iloc[row_idx].get(column_name, "")).strip():
        return None
    monday = ref_date - timedelta(days=ref_date.weekday())
    picked_date = monday + timedelta(days=DAYS.index(day))
    ds = picked_date.strftime("%Y-%m-%d")
    ver = st.session_state.get("_data_version", 0)
    e = get_effective_timetable_for_date(ds, ver, use_test=use_test)
    if e.empty:
        return None
    if row_label == "교사명" or period_grid:
        m = e[(e["교사명"].astype(str).str.strip() == row_name) & (e["교시"].apply(safe_int) == period)]
    else:
        m = e[(e["학급"].astype(str).str.strip() == row_name) & (e["교시"].apply(safe_int) == period)]
    if m.empty:
        return None
    r = m.iloc[0]
    return {
        "교사명": str(r.get("교사명", "")).strip(), "일자": ds, "요일": day, "교시": period,
        "학급": str(r.get("학급", "")).strip(), "과목": str(r.get("과목", "")).strip(),
        "변경유형": str(r.get("변경유형", "원본")).strip(), "변경출처": str(r.get("변경출처", "")).strip(),
        "변경ID": str(r.get("변경ID", "")).strip(), "원본교사": str(r.get("원본교사", r.get("교사명", ""))).strip(),
        "원본일자": str(r.get("원본일자", ds)), "원본교시": safe_int(r.get("원본교시", period)),
    }
def render_weekly_selection_panel(ref_date, *, use_test=False, title="선택 수업 작업"):
    lesson = st.session_state.get("weekly_selected_lesson")
    if lesson:
        st.session_state.weekly_dialog_use_test = use_test
        st.session_state.weekly_dialog_title = title
        st.session_state.weekly_dialog_open = True
    else:
        st.caption("주간표의 수업 셀을 클릭하면 작은 팝업에서 결강·맞교환·보강 작업을 시작할 수 있습니다.")
def render_standard_weekly_matrix(matrix: pd.DataFrame, ref_date: date, *, row_label="교사명", key="weekly_matrix", title=None, use_test=False, height=900, open_dialog=True):
    return render_weekly_matrix(
        matrix, ref_date, row_label=row_label, height=height, key=key,
        title=title, show_week_dates=True, use_test=use_test, open_dialog=open_dialog
    )
def _weekly_selection_signature(selected_cells):
    try:
        return tuple((int(r), str(c)) for r, c in selected_cells)
    except Exception:
        return tuple()
def render_weekly_matrix(matrix: pd.DataFrame, ref_date: date, *, row_label="교사명", height=900,
                         key="weekly_matrix", title=None, show_week_dates=True, use_test=False, open_dialog=True):
    if matrix is None or matrix.empty:
        st.info("표시할 주간 시간표가 없습니다.")
        return None
    monday = ref_date - timedelta(days=ref_date.weekday())
    if title:
        st.markdown(f"<div style='font-size:.86rem;font-weight:600;color:#6b7280;margin:0 0 .12rem .1rem'>{html_lib.escape(str(title))}</div>", unsafe_allow_html=True)
    if show_week_dates:
        dates = [monday + timedelta(days=i) for i in range(5)]
        st.markdown("<div class='compact-nav'>" + "　".join(f"{DAYS[i]} {dates[i]:%m.%d}" for i in range(5)) + "</div>", unsafe_allow_html=True)
    visible_matrix = _hide_past_week_slots(matrix, ref_date, hide_past=True)
    teacher_period_grid = row_label == "교시" and all(d in visible_matrix.columns for d in DAYS)
    display = _weekly_styled_matrix(visible_matrix, drop_teacher_name=teacher_period_grid)
    column_config = {}
    if teacher_period_grid:
        column_config["교시"] = st.column_config.NumberColumn("교시", width=54, format="%d")
        monday = ref_date - timedelta(days=ref_date.weekday())
        week_dates = [monday + timedelta(days=i) for i in range(5)]
        for day_idx, day in enumerate(DAYS):
            if day in display.columns:
                column_config[day] = st.column_config.TextColumn(f"{day}", width=230)
    else:
        compact_period_width = 38
        row_name_width = 90
        if row_label in display.columns:
            column_config[row_label] = st.column_config.TextColumn(row_label, width=row_name_width)
        monday = ref_date - timedelta(days=ref_date.weekday())
        week_dates = [monday + timedelta(days=i) for i in range(5)]
        for day_idx, day in enumerate(DAYS):
            day_date = week_dates[day_idx]
            for p in range(1, MAX_PERIOD + 1):
                col = f"{day}{p}"
                if col in display.columns:
                    label = f"{day} {day_date.day}" if p == 1 else f"{day}{p}"
                    column_config[col] = st.column_config.TextColumn(label, width=compact_period_width)
    event = None
    selected_cells = []
    matrix_epoch = int(st.session_state.get("weekly_matrix_epoch", 0) or 0)
    widget_key = f"{key}__{ref_date:%Y%m%d}__{row_label}__{'test' if use_test else 'live'}__sel{matrix_epoch}"
    try:
        event = st.dataframe(
            display,
            hide_index=True,
            width="stretch",
            height=height,
            row_height=27 if not teacher_period_grid else 36,
            key=widget_key,
            on_select="rerun",
            selection_mode="single-cell",
            column_config=column_config,
        )
        try:
            selected_cells = list(event.selection.cells)
        except Exception:
            selected_cells = []
        selection_seen_key = f"_weekly_selection_seen__{widget_key}"
        current_signature = _weekly_selection_signature(selected_cells)
        previous_signature = st.session_state.get(selection_seen_key, tuple())
        selection_changed = current_signature != previous_signature
        st.session_state[selection_seen_key] = current_signature
        if not selection_changed:
            selected_cells = []
    except TypeError:
        st.dataframe(
            display,
            hide_index=True,
            width="stretch",
            height=height,
            key=f"{key}_legacy__sel{matrix_epoch}",
            column_config=column_config,
        )
        st.warning("현재 Streamlit 버전에서는 주간표 셀 클릭 기능을 지원하지 않습니다. Streamlit을 최신 버전으로 업데이트하면 셀 클릭 팝업을 사용할 수 있습니다.")
        return None
    lesson = _resolve_matrix_cell_selection(matrix, ref_date, row_label, selected_cells, use_test=use_test)
    if lesson and not open_dialog:
        _clear_weekly_selection()
        return None
    if lesson:
        st.session_state.weekly_selected_lesson = lesson
        st.session_state.weekly_dialog_action_mode = "swap"
        st.session_state.weekly_dialog_extra_days = 7
        st.session_state.pop("weekly_dialog_extra_days_input", None)
        st.session_state.weekly_dialog_use_test = bool(use_test)
        st.session_state.weekly_dialog_title = title or "주간표 작업"
        st.session_state["weekly_dialog_instance"] = int(st.session_state.get("weekly_dialog_instance", 0) or 0) + 1
        st.session_state.weekly_dialog_open = bool(open_dialog)
        if open_dialog:
            _weekly_action_dialog()
    elif selection_changed:
        _clear_weekly_selection()
    return lesson
def get_teacher_week_view(teacher: str, ref_date: date, use_test=False):
    monday=ref_date-timedelta(days=ref_date.weekday()); week_dates=[monday+timedelta(days=i) for i in range(5)]
    ver=st.session_state.get("_data_version",0); grid=[]
    for p in range(1,MAX_PERIOD+1):
        row={"교사명": teacher, "교시":p}
        for i,d in enumerate(DAYS):
            ds=week_dates[i].strftime("%Y-%m-%d"); e=get_effective_timetable_for_date(ds,ver,use_test=use_test)
            m=e[(e["교사명"]==teacher)&(e["교시"].apply(safe_int)==p)] if not e.empty else pd.DataFrame()
            if m.empty: row[d]=""; continue
            r=m.iloc[0]; cell=f"{r['학급']} {r['과목']}".strip(); typ=str(r.get("변경유형","원본"))
            if typ=="교환": cell += f" 🔄 {r.get('변경출처','교환')}"
            elif typ=="테스트교환": cell += f" 🧪 {r.get('변경출처','테스트교환')}"
            elif typ=="보강": cell += f" 🟢 {r.get('변경출처','보강')}"
            elif typ=="시간강사": cell += f" 🟡 {r.get('원본교사','')}→시간강사"
            if not st.session_state.absences.empty and ((st.session_state.absences["일자"]==ds)&(st.session_state.absences["교사명"]==teacher)&(st.session_state.absences["교시"]==p)).any(): cell=f"[결강] {cell}"
            row[d]=cell
        grid.append(row)
    return pd.DataFrame(grid),week_dates
def get_changed_teachers_for_week(ref_date: date):
    weekday = ref_date.weekday()
    monday = ref_date - timedelta(days=weekday)
    week_dates = [(monday + timedelta(days=i)).strftime("%Y-%m-%d") for i in range(5)]
    changed = set()
    for df, cols in [(st.session_state.absences, ["교사명"]), (st.session_state.subs, ["결강교사", "보강교사"])]:
        if not df.empty and "일자" in df.columns:
            mask = df["일자"].isin(week_dates)
            for c in cols:
                if c in df.columns:
                    changed.update(df.loc[mask, c].dropna().tolist())
    if not st.session_state.swaps.empty:
        for col in ["원본일자", "목표일자"]:
            if col in st.session_state.swaps.columns:
                mask = st.session_state.swaps[col].isin(week_dates)
                changed.update(st.session_state.swaps.loc[mask, "교사A"].tolist())
                changed.update(st.session_state.swaps.loc[mask, "교사B"].tolist())
    pt = st.session_state.get("part_time", pd.DataFrame())
    if not pt.empty and "시작일" in pt.columns:
        for _, r in pt.iterrows():
            start = normalize_date_str(r.get("시작일", ""))
            end = normalize_date_str(r.get("종료일", ""))
            if start and end:
                for wd in week_dates:
                    if start <= wd <= end:
                        if r.get("대체교사"):
                            changed.add(str(r["대체교사"]).strip())
                        if r.get("시간강사명"):
                            changed.add(str(r["시간강사명"]).strip())
    return sorted(t for t in changed if t)
def filter_by_owner(df):
    if can_full_data() or df.empty or "입력자" not in df.columns:
        return df
    return df[df["입력자"] == current_user()].copy()
def build_personal_plan_html(teacher_name: str, on_date: str) -> str:
    try:
        dt = datetime.strptime(on_date, "%Y-%m-%d")
        day_kr = WEEKDAY_KR[dt.weekday()]
        date_display = f"{dt.year}년 {dt.month}월 {dt.day}일 ({day_kr})"
    except Exception:
        date_display = on_date
        day_kr = ""
    subject_dept = get_teacher_subject(teacher_name)
    dept_line = f"{subject_dept} 과" if subject_dept else "과"
    abs_df = st.session_state.get("absences", pd.DataFrame())
    my_abs = pd.DataFrame()
    if not abs_df.empty:
        my_abs = abs_df[(abs_df["교사명"] == teacher_name) & (abs_df["일자"] == on_date)].copy()
    subs_df = st.session_state.get("subs", pd.DataFrame())
    my_subs = pd.DataFrame()
    if not subs_df.empty:
        my_subs = subs_df[
            (subs_df["결강교사"] == teacher_name) & (subs_df["일자"] == on_date)
        ].copy()
    req_df = load_swap_requests()
    my_reqs = pd.DataFrame()
    if not req_df.empty:
        my_reqs = req_df[
            ((req_df["신청자이름"] == teacher_name) | (req_df["교사A"] == teacher_name)) &
            ((req_df["원본일자"] == on_date) | (req_df["목표일자"] == on_date))
        ].copy()
    swaps = st.session_state.get("swaps", pd.DataFrame())
    my_swaps = pd.DataFrame()
    if not swaps.empty:
        my_swaps = swaps[
            ((swaps["교사A"] == teacher_name) | (swaps["교사B"] == teacher_name)) &
            ((swaps["원본일자"] == on_date) | (swaps["목표일자"] == on_date))
        ].copy()
    reason = str(my_abs.iloc[0].get("사유", "")).strip() if not my_abs.empty else ""
    abs_list = []
    if not my_abs.empty:
        for _, r in my_abs.iterrows():
            p = safe_int(r.get("교시"))
            sub_teacher = ""
            if not my_subs.empty:
                match = my_subs[my_subs["교시"] == p]
                if not match.empty:
                    sub_teacher = str(match.iloc[0].get("보강교사", "")).strip()
            abs_list.append({
                "월일": on_date[5:].replace("-", "/"),
                "교시": p,
                "학년반": str(r.get("학급", "")),
                "과목": str(r.get("과목", "")),
                "보강교사": sub_teacher
            })
    swap_list = []
    for src in [my_reqs, my_swaps]:
        if not src.empty:
            for _, r in src.iterrows():
                target_date = str(r.get("목표일자", r.get("원본일자", "")))
                swap_list.append({
                    "월일": target_date[5:].replace("-", "/") if len(target_date) >= 10 else target_date,
                    "교시": safe_int(r.get("교시B", r.get("교시A", 0))),
                    "과목": str(r.get("과목B", r.get("과목A", ""))),
                    "교사": str(r.get("교사B", r.get("교사A", "")))
                })
    rows_html = ""
    max_rows = 6
    for i in range(max_rows):
        a = abs_list[i] if i < len(abs_list) else {"월일": "", "교시": "", "학년반": "", "과목": "", "보강교사": ""}
        s = swap_list[i] if i < len(swap_list) else {"월일": "", "교시": "", "과목": "", "교사": ""}
        rows_html += f"""
        <tr>
            <td style="height:29px;">{a['월일']}</td>
            <td>{a['교시']}</td>
            <td>{a['학년반']}</td>
            <td>{a['과목']}</td>
            <td>{a['보강교사']}</td>
            <td>{s['월일']}</td>
            <td>{s['교시']}</td>
            <td>{s['과목']}</td>
            <td>{s['교사']}</td>
        </tr>"""
    html = f"""<!DOCTYPE html>
<html lang="ko">
<head>
<meta charset="utf-8">
<title>결·보강 계획서 - {teacher_name}</title>
<style>
    @page {{ size: A4; margin: 12mm 14mm; }}
    * {{ box-sizing: border-box; }}
    body {{
        font-family: '맑은 고딕', 'Malgun Gothic', 'Apple SD Gothic Neo', sans-serif;
        font-size: 12.5px;
        line-height: 1.25;
        margin: 0;
        padding: 6px 10px;
        color: #000;
        background: #fff;
    }}
    table {{
        border-collapse: collapse;
        width: 100%;
        table-layout: fixed;
    }}
    th, td {{
        border: 1px solid #000;
        padding: 2px 2px;
        text-align: center;
        vertical-align: middle;
        word-break: keep-all;
    }}
    .title {{
        text-align: center;
        font-size: 22px;
        font-weight: bold;
        letter-spacing: 5px;
        margin: 2px 0 8px 0;
        text-decoration: underline;
        text-underline-offset: 3px;
    }}
    .top-right {{
        width: 150px;
        float: right;
        margin-top: -36px;
    }}
    .top-right td {{
        height: 26px;
        font-size: 12px;
        font-weight: bold;
    }}
    .dept-line {{
        font-size: 13.5px;
        margin: 4px 0 6px 2px;
    }}
    .date-box td {{
        height: 32px;
        font-size: 12.5px;
    }}
    .section-header {{
        background-color: #d6e3f0;
        font-weight: bold;
        font-size: 12px;
    }}
    .sub-header {{
        background-color: #eef3f9;
        font-size: 11.5px;
        font-weight: bold;
    }}
    .note-box {{
        border: 1px solid #000;
        min-height: 88px;
        margin-top: 0;
        padding: 6px;
    }}
</style>
</head>
<body>
<div class="title">결 · 보 강  계 획</div>
<table class="top-right">
    <tr>
        <td style="width:50%;">수업계</td>
        <td style="width:50%;">교육과정</td>
    </tr>
    <tr>
        <td style="height:34px;"></td>
        <td></td>
    </tr>
</table>
<div class="dept-line">
    <b>{dept_line}</b> &nbsp;&nbsp; 교 사 : <b>{teacher_name}</b> &nbsp;&nbsp;&nbsp; (인)
</div>
<table class="date-box" style="margin-bottom: 9px;">
    <tr>
        <td style="width: 68px; background:#f0f0f0; font-weight:bold;">결강<br>일자</td>
        <td style="text-align:left; padding-left:10px;">
            {date_display}<br>
            <span style="display:inline-block; margin-top:2px;">사유 : {reason}</span>
        </td>
    </tr>
</table>
<table>
    <thead>
        <tr>
            <th colspan="4" class="section-header">결강수업</th>
            <th class="section-header">보강수업</th>
            <th colspan="4" class="section-header">교체수업</th>
        </tr>
        <tr class="sub-header">
            <th style="width:9.5%;">월일</th>
            <th style="width:7%;">교시</th>
            <th style="width:10%;">학년반</th>
            <th style="width:11%;">과목</th>
            <th style="width:12%;">교사(인)</th>
            <th style="width:9.5%;">월일</th>
            <th style="width:7%;">교시</th>
            <th style="width:11%;">과목</th>
            <th style="width:13%;">교사(인)</th>
        </tr>
    </thead>
    <tbody>
        {rows_html}
    </tbody>
</table>
<div style="margin-top: 13px;">
    <div style="text-align:center; font-weight:bold; border:1px solid #000; border-bottom:none; padding:5px 0;">
        추가 기재 사항
    </div>
    <div class="note-box"></div>
</div>
<div style="margin-top: 10px; font-size: 10.5px; color:#555; text-align:right;">
    {SCHOOL_NAME} · {SCHOOL_YEAR}학년도 &nbsp;|&nbsp; 생성시각 {datetime.now().strftime('%Y-%m-%d %H:%M')}
</div>
</body>
</html>"""
    return html
def build_test_swaps_report_html(test_swaps: pd.DataFrame) -> str:
    if test_swaps.empty:
        return "<html><body><p>테스트 중인 맞교환 내역이 없습니다.</p></body></html>"
    my_name = current_name().strip()
    if not my_name:
        my_name = current_user()
    my_swaps = test_swaps[
        (test_swaps["교사A"] == my_name) | (test_swaps["교사B"] == my_name)
    ].copy()
    if my_swaps.empty:
        return f"<html><body><p>{my_name} 선생님의 테스트 맞교환 내역이 없습니다.</p></body></html>"
    origin_dates = []
    for _, r in my_swaps.iterrows():
        da = str(r.get("원본일자", "")).strip()
        if da and str(r.get("교사A", "")).strip() == my_name:
            origin_dates.append(da)
    if not origin_dates:
        for _, r in my_swaps.iterrows():
            da = str(r.get("원본일자", "")).strip()
            if da:
                origin_dates.append(da)
    if not origin_dates:
        return f"<html><body><p>{my_name} 선생님의 결강 대상 일자가 없습니다.</p></body></html>"
    on_date = sorted(set(origin_dates))[0]
    teacher_name = my_name
    try:
        dt = datetime.strptime(on_date, "%Y-%m-%d")
        day_kr = WEEKDAY_KR[dt.weekday()]
        date_display = f"{dt.year}년 {dt.month}월 {dt.day}일 ({day_kr})"
    except Exception:
        date_display = on_date
        day_kr = ""
    subject_dept = get_teacher_subject(teacher_name)
    dept_line = f"{subject_dept} 과" if subject_dept else "과"
    day_swaps = my_swaps[my_swaps["원본일자"] == on_date].copy()
    absence_list = []
    swap_list = []
    for _, r in day_swaps.iterrows():
        if str(r.get("교사A", "")) == teacher_name:
            absence_list.append({
                "월일": on_date[5:].replace("-", "/") if len(on_date) >= 10 else on_date,
                "교시": safe_int(r.get("교시A", 0)),
                "학년반": str(r.get("학급A", "")),
                "과목": str(r.get("과목A", "")),
                "보강교사": ""
            })
            target_date = str(r.get("목표일자", ""))
            swap_list.append({
                "월일": target_date[5:].replace("-", "/") if len(target_date) >= 10 else target_date,
                "교시": safe_int(r.get("교시B", 0)),
                "과목": str(r.get("과목B", "")),
                "교사": str(r.get("교사B", ""))
            })
    rows_html = ""
    max_rows = 6
    for i in range(max_rows):
        a = absence_list[i] if i < len(absence_list) else {
            "월일": "", "교시": "", "학년반": "", "과목": "", "보강교사": ""
        }
        s = swap_list[i] if i < len(swap_list) else {
            "월일": "", "교시": "", "과목": "", "교사": ""
        }
        rows_html += f"""
        <tr>
            <td style="height:29px;">{a['월일']}</td>
            <td>{a['교시']}</td>
            <td>{a['학년반']}</td>
            <td>{a['과목']}</td>
            <td>{a['보강교사']}</td>
            <td>{s['월일']}</td>
            <td>{s['교시']}</td>
            <td>{s['과목']}</td>
            <td>{s['교사']}</td>
        </tr>"""
    page_html = f"""
<div class="page">
    <div class="title">결 · 보 강  계 획 </div>
    <table class="top-right">
        <tr>
            <td style="width:50%;">수업계</td>
            <td style="width:50%;">교육과정</td>
        </tr>
        <tr>
            <td style="height:34px;"></td>
            <td></td>
        </tr>
    </table>
    <div class="dept-line">
        <b>{dept_line}</b> &nbsp;&nbsp; 교 사 : <b>{teacher_name}</b> &nbsp;&nbsp;&nbsp; (인)
    </div>
    <table class="date-box" style="margin-bottom: 9px;">
        <tr>
            <td style="width: 68px; background:#f0f0f0; font-weight:bold;">해당<br>일자</td>
            <td style="text-align:left; padding-left:10px;">
                {date_display}<br>
                <span style="display:inline-block; margin-top:2px;">사유 : </span>
            </td>
        </tr>
    </table>
    <table>
        <thead>
            <tr>
                <th colspan="4" class="section-header">결강수업</th>
                <th class="section-header">보강수업</th>
                <th colspan="4" class="section-header">교체수업</th>
            </tr>
            <tr class="sub-header">
                <th style="width:9.5%;">월일</th>
                <th style="width:7%;">교시</th>
                <th style="width:10%;">학년반</th>
                <th style="width:11%;">과목</th>
                <th style="width:12%;">교사(인)</th>
                <th style="width:9.5%;">월일</th>
                <th style="width:7%;">교시</th>
                <th style="width:11%;">과목</th>
                <th style="width:13%;">교사(인)</th>
            </tr>
        </thead>
        <tbody>
            {rows_html}
        </tbody>
    </table>
    <div style="margin-top: 13px;">
        <div style="text-align:center; font-weight:bold; border:1px solid #000; border-bottom:none; padding:5px 0;">
            추가 기재 사항
        </div>
        <div class="note-box"></div>
    </div>
    <div style="margin-top: 10px; font-size: 10.5px; color:#555; text-align:right;">
        {SCHOOL_NAME} · {SCHOOL_YEAR}학년도 &nbsp;|&nbsp; 생성시각 {datetime.now().strftime('%Y-%m-%d %H:%M')}
    </div>
</div>"""
    html = f"""<!DOCTYPE html>
<html lang="ko">
<head>
<meta charset="utf-8">
<title>결보강 계획서</title>
<style>
    @page {{ size: A4; margin: 12mm 14mm; }}
    * {{ box-sizing: border-box; }}
    body {{
        font-family: '맑은 고딕', 'Malgun Gothic', 'Apple SD Gothic Neo', sans-serif;
        font-size: 12.5px;
        line-height: 1.25;
        margin: 0;
        padding: 0;
        color: #000;
        background: #fff;
    }}
    .page {{
        padding: 6px 10px;
        page-break-after: always;
    }}
    .page:last-child {{
        page-break-after: auto;
    }}
    table {{
        border-collapse: collapse;
        width: 100%;
        table-layout: fixed;
    }}
    th, td {{
        border: 1px solid #000;
        padding: 2px 2px;
        text-align: center;
        vertical-align: middle;
        word-break: keep-all;
    }}
    .title {{
        text-align: center;
        font-size: 22px;
        font-weight: bold;
        letter-spacing: 5px;
        margin: 2px 0 8px 0;
        text-decoration: underline;
        text-underline-offset: 3px;
    }}
    .top-right {{
        width: 150px;
        float: right;
        margin-top: -36px;
    }}
    .top-right td {{
        height: 26px;
        font-size: 12px;
        font-weight: bold;
    }}
    .dept-line {{
        font-size: 13.5px;
        margin: 4px 0 6px 2px;
    }}
    .date-box td {{
        height: 32px;
        font-size: 12.5px;
    }}
    .section-header {{
        background-color: #d6e3f0;
        font-weight: bold;
        font-size: 12px;
    }}
    .sub-header {{
        background-color: #eef3f9;
        font-size: 11.5px;
        font-weight: bold;
    }}
    .note-box {{
        border: 1px solid #000;
        min-height: 88px;
        margin-top: 0;
        padding: 6px;
    }}
</style>
</head>
<body>
{page_html}
</body>
</html>"""
    return html
def build_report_html(norm_date: str) -> str:
    day = WEEKDAY_KR.get(datetime.strptime(norm_date, "%Y-%m-%d").weekday(), "")
    a = st.session_state.absences
    s = st.session_state.subs
    w = st.session_state.swaps
    def rows_abs():
        if a.empty or "일자" not in a.columns:
            return "<tr><td colspan='6'>없음</td></tr>"
        sub = a[a["일자"] == norm_date]
        if sub.empty:
            return "<tr><td colspan='6'>없음</td></tr>"
        return "".join(
            f"<tr><td>{r.get('교사명','')}</td><td>{r.get('사유','')}</td><td>{safe_int(r.get('교시'))}</td>"
            f"<td>{r.get('학급','')}</td><td>{r.get('과목','')}</td><td>{r.get('상세사유','')}</td></tr>"
            for _, r in sub.iterrows()
        )
    def rows_sub():
        if s.empty or "일자" not in s.columns:
            return "<tr><td colspan='7'>없음</td></tr>"
        sub = s[s["일자"] == norm_date]
        if sub.empty:
            return "<tr><td colspan='7'>없음</td></tr>"
        return "".join(
            f"<tr><td>{safe_int(r.get('교시'))}</td><td>{r.get('학급','')}</td><td>{r.get('과목','')}</td>"
            f"<td>{r.get('결강교사','')}</td><td><b>{r.get('보강교사','')}</b></td>"
            f"<td>{r.get('우선순위','')}</td><td>{r.get('비고','')}</td></tr>"
            for _, r in sub.iterrows()
        )
    def rows_swap():
        if w.empty:
            return "<tr><td colspan='4'>없음</td></tr>"
        mask = (w["원본일자"] == norm_date) | (w["목표일자"] == norm_date)
        sub = w[mask]
        if sub.empty:
            return "<tr><td colspan='4'>없음</td></tr>"
        return "".join(
            f"<tr><td>{r.get('교사A','')}</td>"
            f"<td>[{r.get('원본일자')}] {r.get('요일A')} {safe_int(r.get('교시A'))}교시 ↔ "
            f"[{r.get('목표일자')}] {r.get('요일B')} {safe_int(r.get('교시B'))}교시</td>"
            f"<td>{r.get('교사B','')}</td><td>{r.get('유형','')}</td></tr>"
            for _, r in sub.iterrows()
        )
    return f"""<!doctype html><html lang="ko"><head><meta charset="utf-8">
<title>{norm_date} 내역서</title>
<style>
body{{font-family:'맑은 고딕',sans-serif;font-size:12px}}
table{{width:100%;border-collapse:collapse}}
th,td{{border:1px solid #999;padding:5px;text-align:center}}
th{{background:#eef1f6}}
</style></head><body>
<h1 style="text-align:center">결강·보강 변경 내역서</h1>
<div style="text-align:center">{SCHOOL_YEAR} · {SCHOOL_NAME} · {norm_date} ({day})</div>
<h3>1. 결강 현황</h3>
<table><tr><th>교사</th><th>사유</th><th>교시</th><th>학급</th><th>과목</th><th>상세</th></tr>{rows_abs()}</table>
<h3>2. 보강 배정</h3>
<table><tr><th>교시</th><th>학급</th><th>과목</th><th>결강교사</th><th>보강교사</th><th>근거</th><th>비고</th></tr>{rows_sub()}</table>
<h3>3. 맞교환</h3>
<table><tr><th>교사A</th><th>내용</th><th>교사B</th><th>유형</th></tr>{rows_swap()}</table>
</body></html>"""
def build_weekly_schedule_excel_bytes(ref_date: date, *, use_test=False, title="전체 교사 시간표") -> bytes:
    monday = ref_date - timedelta(days=ref_date.weekday())
    ver = st.session_state.get("_data_version", 0)
    wb = Workbook(); ws = wb.active; ws.title = "전체 교사 시간표"
    title_font = Font(name="돋움", size=20, bold=True)
    small_font = Font(name="돋움", size=8)
    head_font = Font(name="돋움", size=9, bold=True)
    body_font = Font(name="돋움", size=8)
    thin = Side(style="hair", color="B7B7B7"); med = Side(style="medium", color="808080")
    fill_head = PatternFill("solid", fgColor="D9EAF7")
    fill_day = [PatternFill("solid", fgColor=x) for x in ("DDEBF7","E2F0D9","FFF2CC","E4DFEC","FCE4D6")]
    days = DAYS; periods = {d: list(range(1, PERIODS_PER_DAY.get(d,7)+1)) for d in days}
    col = 3
    day_ranges = {}
    for i,d in enumerate(days):
        start=col; end=col+len(periods[d])-1; day_ranges[d]=(start,end)
        col=end+1
    last_col=col-1
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=last_col)
    ws["A1"] = title; ws["A1"].font=title_font; ws["A1"].alignment=Alignment(horizontal="center",vertical="center")
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=2); ws["A2"]=f"{SCHOOL_YEAR} 학년도"; ws["A2"].font=small_font; ws["A2"].alignment=Alignment(horizontal="left",vertical="center")
    ws.merge_cells(start_row=2,start_column=last_col-4,end_row=2,end_column=last_col); ws.cell(2,last_col-4).value=SCHOOL_NAME; ws.cell(2,last_col-4).font=small_font; ws.cell(2,last_col-4).alignment=Alignment(horizontal="right",vertical="center")
    ws["A3"]="번호"; ws["B3"]="교사"
    for c in (1,2): ws.cell(3,c).font=head_font; ws.cell(3,c).alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); ws.cell(3,c).fill=fill_head
    ws.merge_cells("A3:A4"); ws.merge_cells("B3:B4")
    for i,d in enumerate(days):
        start,end=day_ranges[d]
        ws.merge_cells(start_row=3,start_column=start,end_row=3,end_column=end)
        cell=ws.cell(3,start); cell.value=f"{d}({(monday + timedelta(days=i)):%m/%d})"; cell.font=head_font; cell.alignment=Alignment(horizontal="center",vertical="center"); cell.fill=fill_day[i]
        for pno in periods[d]:
            cc=start+pno-1; pc=ws.cell(4,cc); pc.value=pno; pc.font=head_font; pc.alignment=Alignment(horizontal="center",vertical="center"); pc.fill=fill_day[i]
    teachers=sorted(st.session_state.timetable["교사명"].dropna().astype(str).str.strip().unique()) if not st.session_state.timetable.empty else []
    daily_indexes = {}
    for i, d in enumerate(days):
        ds=(monday + timedelta(days=i)).strftime("%Y-%m-%d")
        e=get_effective_timetable_for_date(ds,ver,use_test=use_test)
        daily_indexes[d] = {(str(r.교사명).strip(), safe_int(r.교시)): r for r in e.itertuples(index=False)} if not e.empty else {}
    row=5
    for n,t in enumerate(teachers,1):
        ws.merge_cells(start_row=row,start_column=1,end_row=row+1,end_column=1); ws.merge_cells(start_row=row,start_column=2,end_row=row+1,end_column=2)
        ws.cell(row,1).value=n; ws.cell(row,2).value=t
        for rr in (row,row+1):
            for cc in range(1,last_col+1):
                c=ws.cell(rr,cc); c.font=body_font; c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=Border(left=thin,right=thin,top=thin,bottom=thin)
        ws.cell(row,1).border=Border(left=med,right=thin,top=thin,bottom=thin); ws.cell(row,2).border=Border(left=thin,right=med,top=thin,bottom=thin)
        for i,d in enumerate(days):
            idx = daily_indexes[d]
            for pno in periods[d]:
                cc=day_ranges[d][0]+pno-1
                r=idx.get((t,pno))
                if r is None: continue
                subj=str(getattr(r,"과목","")).strip(); cls=str(getattr(r,"학급","")).strip(); typ=str(getattr(r,"변경유형","원본")).strip()
                ws.cell(row,cc).value=subj
                ws.cell(row+1,cc).value=cls
                if typ=="교환": fill=PatternFill("solid",fgColor="F4CCCC"); mark="🔄"
                elif typ=="보강": fill=PatternFill("solid",fgColor="D9EAD3"); mark="🟢"
                elif typ=="테스트교환": fill=PatternFill("solid",fgColor="EADCF8"); mark="🧪"
                elif typ=="시간강사": fill=PatternFill("solid",fgColor="FCE5CD"); mark="🟡"
                else: fill=fill_day[i]; mark=""
                ws.cell(row,cc).fill=fill; ws.cell(row+1,cc).fill=fill
                if mark: ws.cell(row,cc).value=f"{subj} {mark}".strip()
        row += 2
    ws.row_dimensions[1].height=22.5; ws.row_dimensions[2].height=14.25; ws.row_dimensions[3].height=18; ws.row_dimensions[4].height=16
    for rr in range(5,row): ws.row_dimensions[rr].height=18 if rr%2==1 else 16
    ws.column_dimensions["A"].width=5; ws.column_dimensions["B"].width=12
    for cc in range(3,last_col+1): ws.column_dimensions[get_column_letter(cc)].width=11
    ws.freeze_panes="C5"; ws.sheet_view.showGridLines=False
    ws.page_setup.orientation="landscape"; ws.page_setup.paperSize=ws.PAPERSIZE_A4; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_margins=PageMargins(left=0.25,right=0.25,top=0.4,bottom=0.4,header=0.2,footer=0.2)
    ws.print_title_rows="1:4"; ws.print_area=f"A1:{get_column_letter(last_col)}{row-1}"
    return _workbook_bytes(wb)
def build_weekly_class_schedule_excel_bytes(ref_date: date, *, use_test=False, title="전체 학급 시간표") -> bytes:
    monday = ref_date - timedelta(days=ref_date.weekday())
    ver = st.session_state.get("_data_version", 0)
    wb = Workbook(); ws = wb.active; ws.title = "전체 학급 시간표"
    title_font=Font(name="돋움",size=20,bold=True); small_font=Font(name="돋움",size=8); head_font=Font(name="돋움",size=9,bold=True); body_font=Font(name="돋움",size=8)
    thin=Side(style="hair",color="B7B7B7"); med=Side(style="medium",color="808080"); fill_head=PatternFill("solid",fgColor="D9EAF7")
    fill_day=[PatternFill("solid",fgColor=x) for x in ("DDEBF7","E2F0D9","FFF2CC","E4DFEC","FCE4D6")]
    periods={d:list(range(1,PERIODS_PER_DAY.get(d,7)+1)) for d in DAYS}; col=3; day_ranges={}
    for i,d in enumerate(DAYS):
        start=col; end=col+len(periods[d])-1; day_ranges[d]=(start,end); col=end+1
    last_col=col-1
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=last_col); ws["A1"]=title; ws["A1"].font=title_font; ws["A1"].alignment=Alignment(horizontal="center",vertical="center")
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=2); ws["A2"]=f"{SCHOOL_YEAR} 학년도"; ws["A2"].font=small_font
    ws.merge_cells(start_row=2,start_column=last_col-4,end_row=2,end_column=last_col); ws.cell(2,last_col-4).value=SCHOOL_NAME; ws.cell(2,last_col-4).font=small_font; ws.cell(2,last_col-4).alignment=Alignment(horizontal="right",vertical="center")
    ws["A3"]="번호"; ws["B3"]="학급"; ws.merge_cells("A3:A4"); ws.merge_cells("B3:B4")
    for c in (1,2): ws.cell(3,c).font=head_font; ws.cell(3,c).alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); ws.cell(3,c).fill=fill_head
    for i,d in enumerate(DAYS):
        start,end=day_ranges[d]; ws.merge_cells(start_row=3,start_column=start,end_row=3,end_column=end); cell=ws.cell(3,start); cell.value=f"{d}({(monday+timedelta(days=i)):%m/%d})"; cell.font=head_font; cell.alignment=Alignment(horizontal="center",vertical="center"); cell.fill=fill_day[i]
        for pno in periods[d]:
            pc=ws.cell(4,start+pno-1); pc.value=pno; pc.font=head_font; pc.alignment=Alignment(horizontal="center",vertical="center"); pc.fill=fill_day[i]
    daily_indexes={}
    classes=set()
    for i,d in enumerate(DAYS):
        ds=(monday+timedelta(days=i)).strftime("%Y-%m-%d"); e=get_effective_timetable_for_date(ds,ver,use_test=use_test); daily_indexes[d]={(str(r.학급).strip(),safe_int(r.교시)):r for r in e.itertuples(index=False) if str(getattr(r,"학급","")).strip()} if not e.empty else {}
        classes.update(k[0] for k in daily_indexes[d])
    row=5
    for n,cls in enumerate(sorted(classes),1):
        ws.merge_cells(start_row=row,start_column=1,end_row=row+1,end_column=1); ws.merge_cells(start_row=row,start_column=2,end_row=row+1,end_column=2); ws.cell(row,1).value=n; ws.cell(row,2).value=cls
        for rr in (row,row+1):
            for cc in range(1,last_col+1):
                c=ws.cell(rr,cc); c.font=body_font; c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=Border(left=thin,right=thin,top=thin,bottom=thin)
        for i,d in enumerate(DAYS):
            for pno in periods[d]:
                r=daily_indexes[d].get((cls,pno)); cc=day_ranges[d][0]+pno-1
                if r is None: continue
                subj=str(getattr(r,"과목","")).strip(); teacher=str(getattr(r,"교사명","")).strip(); typ=str(getattr(r,"변경유형","원본")).strip(); ws.cell(row,cc).value=subj; ws.cell(row+1,cc).value=teacher
                fill=fill_day[i]; mark=""
                if typ=="교환": fill=PatternFill("solid",fgColor="F4CCCC"); mark=" 🔄"
                elif typ=="보강": fill=PatternFill("solid",fgColor="D9EAD3"); mark=" 🟢"
                elif typ=="테스트교환": fill=PatternFill("solid",fgColor="EADCF8"); mark=" 🧪"
                elif typ=="시간강사": fill=PatternFill("solid",fgColor="FCE5CD"); mark=" 🟡"
                ws.cell(row,cc).value=f"{subj}{mark}".strip(); ws.cell(row,cc).fill=fill; ws.cell(row+1,cc).fill=fill
        row+=2
    ws.row_dimensions[1].height=22.5; ws.row_dimensions[2].height=14.25; ws.row_dimensions[3].height=18; ws.row_dimensions[4].height=16
    for rr in range(5,row): ws.row_dimensions[rr].height=18 if rr%2==1 else 16
    ws.column_dimensions["A"].width=5; ws.column_dimensions["B"].width=12
    for cc in range(3,last_col+1): ws.column_dimensions[get_column_letter(cc)].width=11
    ws.freeze_panes="C5"; ws.sheet_view.showGridLines=False; ws.page_setup.orientation="landscape"; ws.page_setup.paperSize=ws.PAPERSIZE_A4; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_margins=PageMargins(left=0.25,right=0.25,top=0.4,bottom=0.4,header=0.2,footer=0.2); ws.print_title_rows="1:4"; ws.print_area=f"A1:{get_column_letter(last_col)}{row-1}"
    return _workbook_bytes(wb)
def _workbook_bytes(wb):
    buf=io.BytesIO(); wb.save(buf); return buf.getvalue()
def to_excel_bytes(sheets: dict) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for name, df in sheets.items():
            (df if isinstance(df, pd.DataFrame) else pd.DataFrame(df)).to_excel(w, sheet_name=name[:31], index=False)
    return buf.getvalue()
SALARY_BASE_OPTIONS = [
    (9, "정교사(1급)"),
    (9, "전문상담교사(1급)"),
    (9, "사서교사(1급)"),
    (9, "보건교사(1급)"),
    (9, "영양교사(1급)"),
    (9, "교장·원장·교감·원감·교육장·장학(연구)직"),
    (8, "정교사(2급)"),
    (8, "전문상담교사(2급)"),
    (8, "사서교사(2급)"),
    (8, "보건교사(2급)"),
    (8, "영양교사(2급)"),
    (5, "준교사"),
    (5, "실기교사"),
]
SALARY_ACADEMIC_OPTIONS = [
    (0, "4년제 일반대학 졸업", "학령: 16년 → 학령가감: +0년"),
    (0, "교육대학·사범대학 졸업 (4년)", "학령: 16년 → 학령가감: +0년"),
    (2, "6년제 대학 졸업 (의대 등)", "학령: 18년 → 학령가감: +2년"),
]
SALARY_DEGREE_TYPES = [
    "동등 수준 추가 학사 학위 (2번째 대학교)",
    "석사학위 취득 수학기간",
    "박사학위 취득 수학기간",
]
SALARY_CAREER_TYPES = [
    "국·공립학교 교원 (기간제 포함, 자격 일치)",
    "사립학교 교원 (관할청 보고, 자격 일치)",
    "기간제교원 자격-학교급 불일치 (예:중등→초등)",
    "유·초·중등 강사 (전일제·종일제, 1일 8시간 ↑)",
    "유·초·중등 시간제 강사 (주 12시간 ↓ 또는 시수 불명)",
    "국가·지방공무원 (현역 군복무 포함)",
    "등록 학원 강사 / 신고 교습소 교습자",
    "회사 (상법상 합명·합자·주식·유한회사) 근무",
]
SALARY_CAREER_RATE_BY_TYPE = {
    "국·공립학교 교원 (기간제 포함, 자격 일치)": 100,
    "사립학교 교원 (관할청 보고, 자격 일치)": 100,
    "기간제교원 자격-학교급 불일치 (예:중등→초등)": 80,
    "유·초·중등 강사 (전일제·종일제, 1일 8시간 ↑)": 100,
    "유·초·중등 시간제 강사 (주 12시간 ↓ 또는 시수 불명)": 30,
    "국가·지방공무원 (현역 군복무 포함)": 100,
    "등록 학원 강사 / 신고 교습소 교습자": 50,
    "회사 (상법상 합명·합자·주식·유한회사) 근무": 40,
}
SALARY_DEGREE_RATE_BY_TYPE = {
    "동등 수준 추가 학사 학위 (2번째 대학교)": 80,
    "석사학위 취득 수학기간": 100,
    "박사학위 취득 수학기간": 100,
}
SALARY_GEMINI_PROMPT = """
당신은 대한민국 교육공무원 및 기간제교원 호봉 획정 서류 분석 전문가입니다.
첨부된 문서(경력증명서, 인사기록카드, 자격증 등)를 정확히 읽고, 아래 요청하는 형식의 JSON으로만 출력해 주세요.
[출력 구조 예시]
{
  "name": "홍길동",
  "baseSalary": 8,
  "qualLabel": "정교사(2급)",
  "academicValue": "0",
  "isSabom": true,
  "degrees": [
    {
      "type": "동등 수준 추가 학사 학위 (2번째 대학교)",
      "detail": "OO대학교 국어교육과 (학사)",
      "start": "2018-03-01",
      "end": "2020-02-28",
      "rate": 80
    }
  ],
  "careers": [
    {
      "type": "국·공립학교 교원 (기간제 포함, 자격 일치)",
      "detail": "OO고등학교 기간제교사",
      "start": "2022-03-01",
      "end": "2023-02-28",
      "inc": true,
      "rate": 100
    }
  ]
}
[경력 type 매칭 옵션]
- "국·공립학교 교원 (기간제 포함, 자격 일치)" (100%)
- "사립학교 교원 (관할청 보고, 자격 일치)" (100%)
- "기간제교원 자격-학교급 불일치 (예:중등→초등)" (80%)
- "유·초·중등 강사 (전일제·종일제, 1일 8시간 ↑)" (100%)
- "유·초·중등 시간제 강사 (주 12시간 ↓ 또는 시수 불명)" (30%)
- "국가·지방공무원 (현역 군복무 포함)" (100%)
- "등록 학원 강사 / 신고 교습소 교습자" (50%)
- "회사 (상법상 합명·합자·주식·유한회사) 근무" (40%)
[주의사항]
1. 날짜는 반드시 YYYY-MM-DD 형식이어야 합니다.
2. 경력사항이 여러 개일 경우 모두 careers 배열에 넣어주세요.
3. 마크다운 코드블록이나 기타 설명 없이 pure JSON 문자열만 반환해야 합니다.
""".strip()
def _salary_default_degree_df():
    return pd.DataFrame(columns=["학위 구분", "학교/전공 세부명", "입학일", "졸업일", "환산율"])
def _salary_default_career_df():
    return pd.DataFrame(columns=["경력 종류", "세부 근무처/직위", "시작일", "종료일", "종료일 산입", "환산율 (%)"])
def _salary_default_rate_for_career_type(career_type):
    return SALARY_CAREER_RATE_BY_TYPE.get(str(career_type or "").strip(), 100)
def _salary_default_rate_for_degree_type(degree_type):
    return SALARY_DEGREE_RATE_BY_TYPE.get(str(degree_type or "").strip(), 80)
def _salary_sync_rate_defaults(df, type_col, rate_col, rate_map, fallback_rate, previous_snapshot=None):
    columns = list(df.columns)
    out = df.copy().reset_index(drop=True)
    previous_snapshot = previous_snapshot or []
    changed = False
    for i in range(len(out)):
        current_type = str(out.at[i, type_col]).strip()
        mapped_rate = rate_map.get(current_type, fallback_rate)
        current_rate_raw = out.at[i, rate_col]
        current_rate = safe_int(current_rate_raw, mapped_rate)
        prev_type = ""
        if i < len(previous_snapshot):
            try:
                prev_type = str(previous_snapshot[i][0] or "").strip()
            except Exception:
                prev_type = ""
        if current_type and (not prev_type or prev_type != current_type):
            if current_rate != mapped_rate:
                out.at[i, rate_col] = mapped_rate
                changed = True
        elif pd.isna(current_rate_raw) or str(current_rate_raw).strip() in ("", "nan", "None", "NaT"):
            out.at[i, rate_col] = mapped_rate
            changed = True
        else:
            if current_rate < 0 or current_rate > 100:
                out.at[i, rate_col] = mapped_rate
                changed = True
    return out[columns], changed
def _salary_rate_snapshot(df, type_col, rate_col):
    if df is None or df.empty:
        return []
    snapshot = []
    for _, row in df.iterrows():
        snapshot.append((str(row.get(type_col, "")).strip(), safe_int(row.get(rate_col), -1)))
    return snapshot
def _salary_init_state():
    defaults = {
        "salary_name": "",
        "salary_calc_date": _today_kst(),
        "salary_org": "",
        "salary_position": "기간제교사",
        "salary_writer_name": "",
        "salary_writer_pos": "",
        "salary_base_idx": 7,
        "salary_academic_idx": 1,
        "salary_is_sabom": True,
        "salary_degrees": _salary_default_degree_df(),
        "salary_careers": _salary_default_career_df(),
        "salary_result": None,
        "salary_ai_message": "",
    }
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value.copy(deep=True) if isinstance(value, pd.DataFrame) else value
    if not st.session_state.get("_salary_rate_defaults_migrated", False):
        degrees = _salary_normalize_table(
            st.session_state.salary_degrees,
            ["학위 구분", "학교/전공 세부명", "입학일", "졸업일", "환산율"],
        )
        for i in range(len(degrees)):
            dtype = str(degrees.at[i, "학위 구분"] or "").strip()
            if dtype:
                degrees.at[i, "환산율"] = _salary_default_rate_for_degree_type(dtype)
        careers = _salary_normalize_table(
            st.session_state.salary_careers,
            ["경력 종류", "세부 근무처/직위", "시작일", "종료일", "종료일 산입", "환산율 (%)"],
        )
        for i in range(len(careers)):
            ctype = str(careers.at[i, "경력 종류"] or "").strip()
            if ctype:
                careers.at[i, "환산율 (%)"] = _salary_default_rate_for_career_type(ctype)
        st.session_state.salary_degrees = degrees
        st.session_state.salary_careers = careers
        st.session_state["_salary_rate_defaults_migrated"] = True
def _salary_safe_date(value):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    if not s or s.lower() in {"nan", "none", "nat"}:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None
def _salary_duration_days(start_value, end_value, include_end=True):
    start = _salary_safe_date(start_value)
    end = _salary_safe_date(end_value)
    if start is None or end is None:
        return 0
    if include_end:
        end = end + timedelta(days=1)
    y = end.year - start.year
    m = end.month - start.month
    d = end.day - start.day
    if d < 0:
        m -= 1
        first_of_month = end.replace(day=1)
        prev_month_last_day = first_of_month - timedelta(days=1)
        d += prev_month_last_day.day
    if m < 0:
        y -= 1
        m += 12
    return (y * 360) + (m * 30) + d
def _salary_ymd_from_360(days):
    days = max(0, int(days))
    return days // 360, (days % 360) // 30, days % 30
def _salary_duration_text(days):
    y, m, d = _salary_ymd_from_360(days)
    return f"{y}년 {m}월 {d}일"
def _salary_normalize_table(df, columns):
    if df is None or not isinstance(df, pd.DataFrame):
        return pd.DataFrame(columns=columns)
    out = df.copy()
    for c in columns:
        if c not in out.columns:
            out[c] = ""
    return out[columns].reset_index(drop=True)
def _salary_add_degree_row():
    _salary_init_state()
    df = _salary_normalize_table(
        st.session_state.salary_degrees,
        ["학위 구분", "학교/전공 세부명", "입학일", "졸업일", "환산율"],
    )
    degree_type = SALARY_DEGREE_TYPES[0]
    df.loc[len(df)] = {
        "학위 구분": degree_type,
        "학교/전공 세부명": "",
        "입학일": None,
        "졸업일": None,
        "환산율": _salary_default_rate_for_degree_type(degree_type),
    }
    st.session_state.salary_degrees = df
def _salary_add_career_row():
    _salary_init_state()
    df = _salary_normalize_table(
        st.session_state.salary_careers,
        ["경력 종류", "세부 근무처/직위", "시작일", "종료일", "종료일 산입", "환산율 (%)"],
    )
    career_type = SALARY_CAREER_TYPES[0]
    df.loc[len(df)] = {
        "경력 종류": career_type,
        "세부 근무처/직위": "",
        "시작일": None,
        "종료일": None,
        "종료일 산입": True,
        "환산율 (%)": _salary_default_rate_for_career_type(career_type),
    }
    st.session_state.salary_careers = df
def _salary_clean_ai_json(raw_text):
    value = (raw_text or "").strip()
    if value.startswith("```"):
        value = re.sub(r"^```(?:json)?\s*", "", value, flags=re.IGNORECASE)
        value = re.sub(r"\s*```$", "", value)
    return json.loads(value)
def _salary_gemini_config():
    api_key = ""
    model = "gemini-2.5-flash"
    try:
        api_key = str(st.secrets.get("GEMINI_API_KEY", "") or "").strip()
        model = str(st.secrets.get("GEMINI_MODEL", model) or model).strip()
    except Exception:
        pass
    api_key = api_key or os.getenv("GEMINI_API_KEY", "").strip()
    model = os.getenv("GEMINI_MODEL", model).strip() or model
    return api_key, model
def _salary_analyze_uploaded_file(uploaded_file):
    api_key, model = _salary_gemini_config()
    if not api_key:
        raise RuntimeError(
            "Gemini API 키가 없습니다. Streamlit Cloud Secrets에 GEMINI_API_KEY를 설정하거나 "
            "환경변수로 등록해 주세요."
        )
    raw = uploaded_file.getvalue()
    if not raw:
        raise ValueError("업로드된 파일이 비어 있습니다.")
    if len(raw) > 50 * 1024 * 1024:
        raise ValueError("PDF/이미지 파일은 50MB 이하로 사용해 주세요.")
    mime_type = getattr(uploaded_file, "type", None) or ""
    if not mime_type:
        suffix = Path(getattr(uploaded_file, "name", "")).suffix.lower()
        mime_type = "application/pdf" if suffix == ".pdf" else "image/jpeg"
    encoded = base64.b64encode(raw).decode("utf-8")
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
    payload = {
        "contents": [{
            "parts": [
                {"inlineData": {"mimeType": mime_type, "data": encoded}},
                {"text": SALARY_GEMINI_PROMPT},
            ]
        }],
        "generationConfig": {
            "responseMimeType": "application/json",
            "temperature": 0.1,
        },
    }
    response = requests.post(
        url,
        params={"key": api_key},
        json=payload,
        timeout=90,
    )
    if response.status_code >= 400:
        try:
            detail = response.json().get("error", {}).get("message", response.text)
        except Exception:
            detail = response.text
        raise RuntimeError(f"Gemini API 오류 ({response.status_code}): {detail}")
    try:
        data = response.json()
        ai_text = data["candidates"][0]["content"]["parts"][0]["text"]
        return _salary_clean_ai_json(ai_text)
    except Exception as exc:
        raise RuntimeError(f"Gemini 응답을 JSON으로 해석하지 못했습니다: {exc}") from exc
def _salary_apply_ai_result(data):
    if not data:
        raise ValueError("AI 분석 결과가 비어 있습니다.")
    if data.get("name"):
        st.session_state.salary_name = str(data["name"]).strip()
    if data.get("baseSalary") is not None:
        base = safe_int(data.get("baseSalary"), 0)
        idx = next((i for i, (v, _) in enumerate(SALARY_BASE_OPTIONS) if v == base), None)
        if idx is not None:
            st.session_state.salary_base_idx = idx
            qual = str(data.get("qualLabel") or SALARY_BASE_OPTIONS[idx][1]).strip()
            exact_label = next(
                (f"{v}호봉 · {label}" for v, label in SALARY_BASE_OPTIONS if v == base and label == qual),
                f"{base}호봉 · {SALARY_BASE_OPTIONS[idx][1]}",
            )
            st.session_state["salary_base_radio"] = exact_label
    if data.get("isSabom") is not None:
        is_sabom = bool(data.get("isSabom"))
        st.session_state.salary_is_sabom = is_sabom
        st.session_state["salary_sabom_checkbox"] = is_sabom
    degree_rows = []
    for deg in data.get("degrees") or []:
        degree_type = str(deg.get("type", SALARY_DEGREE_TYPES[0])).strip()
        degree_rows.append({
            "학위 구분": degree_type,
            "학교/전공 세부명": str(deg.get("detail", "")),
            "입학일": _salary_safe_date(deg.get("start")),
            "졸업일": _salary_safe_date(deg.get("end")),
            "환산율": _salary_default_rate_for_degree_type(degree_type),
        })
    st.session_state.salary_degrees = (
        pd.DataFrame(degree_rows, columns=["학위 구분", "학교/전공 세부명", "입학일", "졸업일", "환산율"])
        if degree_rows else _salary_default_degree_df()
    )
    career_rows = []
    for car in data.get("careers") or []:
        career_type = str(car.get("type", SALARY_CAREER_TYPES[0])).strip()
        career_rows.append({
            "경력 종류": career_type,
            "세부 근무처/직위": str(car.get("detail", "")),
            "시작일": _salary_safe_date(car.get("start")),
            "종료일": _salary_safe_date(car.get("end")),
            "종료일 산입": bool(car.get("inc", True)),
            "환산율 (%)": _salary_default_rate_for_career_type(career_type),
        })
    if career_rows:
        st.session_state.salary_careers = pd.DataFrame(
            career_rows,
            columns=["경력 종류", "세부 근무처/직위", "시작일", "종료일", "종료일 산입", "환산율 (%)"],
        )
    else:
        default_career_type = SALARY_CAREER_TYPES[0]
        st.session_state.salary_careers = pd.DataFrame([{
            "경력 종류": default_career_type,
            "세부 근무처/직위": "",
            "시작일": None,
            "종료일": None,
            "종료일 산입": True,
            "환산율 (%)": _salary_default_rate_for_career_type(default_career_type),
        }])
    st.session_state.salary_ai_message = (
        "AI 분석이 완료되었습니다. 성명·기산호봉·사범계 여부·학위·경력 정보를 자동 반영했습니다."
    )
def _salary_calculate(
    name, calc_date, org, position, writer_name, writer_pos,
    base_value, qual_label, academic_value, academic_label, is_sabom,
    degree_df, career_df
):
    total_converted_days = 0
    details = []
    degree_df = _salary_normalize_table(
        degree_df,
        ["학위 구분", "학교/전공 세부명", "입학일", "졸업일", "환산율"],
    )
    career_df = _salary_normalize_table(
        career_df,
        ["경력 종류", "세부 근무처/직위", "시작일", "종료일", "종료일 산입", "환산율 (%)"],
    )
    for _, row in degree_df.iterrows():
        dtype = str(row.get("학위 구분", "")).strip()
        detail = str(row.get("학교/전공 세부명", "")).strip()
        start = row.get("입학일")
        end = row.get("졸업일")
        rate = max(0, safe_int(row.get("환산율"), 0))
        raw_days = _salary_duration_days(start, end, True)
        converted = math.floor(raw_days * (rate / 100.0))
        total_converted_days += converted
        cy, cm, cd = _salary_ymd_from_360(converted)
        details.append({
            "구분": f"[학위] {dtype} ({detail})",
            "시작일": _salary_safe_date(start),
            "종료일": _salary_safe_date(end),
            "원기간": _salary_duration_text(raw_days),
            "환산율": f"{rate}%",
            "환산년": cy, "환산월": cm, "환산일": cd,
            "종류": "학위",
        })
    for _, row in career_df.iterrows():
        ctype = str(row.get("경력 종류", "")).strip()
        detail = str(row.get("세부 근무처/직위", "")).strip()
        start = row.get("시작일")
        end = row.get("종료일")
        inc = bool(row.get("종료일 산입", True))
        rate = max(0, safe_int(row.get("환산율 (%)"), 0))
        raw_days = _salary_duration_days(start, end, inc)
        converted = math.floor(raw_days * (rate / 100.0))
        total_converted_days += converted
        cy, cm, cd = _salary_ymd_from_360(converted)
        details.append({
            "구분": f"{ctype} ({detail})",
            "시작일": _salary_safe_date(start),
            "종료일": _salary_safe_date(end),
            "원기간": _salary_duration_text(raw_days),
            "환산율": f"{rate}%",
            "환산년": cy, "환산월": cm, "환산일": cd,
            "종류": "경력",
        })
    total_y, total_m, total_d = _salary_ymd_from_360(total_converted_days)
    add_years = 1 if is_sabom else 0
    calculated_step = base_value + academic_value + add_years + total_y
    final_step = max(1, min(40, calculated_step))
    return {
        "name": name,
        "calc_date": calc_date,
        "org": org,
        "position": position,
        "writer_name": writer_name,
        "writer_pos": writer_pos,
        "base_salary": base_value,
        "qual_label": qual_label,
        "academic_value": academic_value,
        "academic_label": academic_label,
        "add_years": add_years,
        "is_sabom": bool(is_sabom),
        "details": details,
        "total_y": total_y, "total_m": total_m, "total_d": total_d,
        "final_step": final_step,
        "rem_m": total_m, "rem_d": total_d,
    }
def _salary_report_html(result):
    def esc(value):
        if value is None:
            return ""
        if isinstance(value, (date, datetime)):
            return value.strftime("%Y-%m-%d")
        return html_lib.escape(str(value))
    rows = []
    for item in result.get("details", []):
        rows.append(
            "<tr>"
            f"<td class='left'>{esc(item['구분'])}</td>"
            f"<td>{esc(item['시작일']) or '-'}</td>"
            f"<td>{esc(item['종료일']) or '-'}</td>"
            f"<td>{esc(item['원기간'])}</td>"
            f"<td>{esc(item['환산율'])}</td>"
            f"<td>{item['환산년']}</td><td>{item['환산월']}</td><td>{item['환산일']}</td>"
            "</tr>"
        )
    if not rows:
        rows.append("<tr><td colspan='8'>학위 및 경력 사항이 없습니다.</td></tr>")
    calc_date = result.get("calc_date")
    calc_date_text = calc_date.strftime("%Y년 %m월 %d일") if isinstance(calc_date, date) else esc(calc_date)
    return f"""<!doctype html>
<html lang="ko">
<head>
<meta charset="utf-8">
<title>호봉 획정 (재획정) 조서 - {esc(result.get('name'))}</title>
<style>
@page {{ size: A4 portrait; margin: 12mm; }}
body {{ font-family: "Malgun Gothic","Noto Sans KR","Apple SD Gothic Neo",sans-serif; color:#111; font-size:11pt; line-height:1.45; }}
h1 {{ text-align:center; font-size:19pt; text-decoration:underline; margin:0 0 18px; }}
h2 {{ font-size:12pt; margin:18px 0 8px; }}
table {{ width:100%; border-collapse:collapse; table-layout:fixed; }}
th,td {{ border:1px solid #333; padding:6px 4px; text-align:center; vertical-align:middle; word-break:break-word; }}
th {{ background:#f1f1f1; font-weight:700; }}
td.left {{ text-align:left; }}
.meta th {{ width:14%; }} .meta td {{ width:36%; }}
.small {{ font-size:9.5pt; }}
.signature {{ display:flex; justify-content:space-between; margin-top:30px; }}
.no-print {{ margin-top:18px; text-align:center; }}
@media print {{ .no-print {{ display:none; }} }}
</style>
</head>
<body>
<h1>호봉 획정 (재획정) 조서</h1>
<table class="meta">
<tr><th>소 속</th><td>{esc(result.get('org'))}</td><th>직 위</th><td>{esc(result.get('position'))}</td></tr>
<tr><th>성 명</th><td>{esc(result.get('name'))}</td><th>교사 (자격)</th><td>{esc(result.get('qual_label'))}</td></tr>
</table>
<h2>■ 학위 수학기간 및 경력 내용 환산 상세</h2>
<table class="small">
<thead>
<tr><th rowspan="2">구분 (학위/경력 내용)</th><th colspan="3">기간</th><th rowspan="2">환산율</th><th colspan="3">환산 경력</th></tr>
<tr><th>시작일</th><th>종료일</th><th>원기간</th><th>년</th><th>월</th><th>일</th></tr>
</thead>
<tbody>{''.join(rows)}</tbody>
<tfoot>
<tr><th colspan="4" style="text-align:right">환산경력 합계 :</th><th colspan="4">{result['total_y']}년 {result['total_m']}월 {result['total_d']}일</th></tr>
</tfoot>
</table>
<h2>■ 호봉 산정 결과</h2>
<table>
<tr><th>기산호봉</th><th>학령가감</th><th>가산연수</th><th>환산경력 (학위+경력)</th><th>최종 초임호봉</th><th>잔여 기간</th></tr>
<tr>
<td>{result['base_salary']}호봉</td>
<td>{'+' if result['academic_value'] >= 0 else ''}{result['academic_value']}년</td>
<td>+{result['add_years']}년</td>
<td>{result['total_y']}년 {result['total_m']}월 {result['total_d']}일</td>
<td><strong>{result['final_step']}호봉</strong></td>
<td>{result['rem_m']}월 {result['rem_d']}일</td>
</tr>
</table>
<div class="signature">
<div><strong>위와 같이 호봉을 획정(재획정)함.</strong><br>일자: {calc_date_text}</div>
<div style="text-align:right">작성자 직급: {esc(result.get('writer_pos'))}<br>작성자 성명: {esc(result.get('writer_name'))} (인)</div>
</div>
<div class="no-print"><button onclick="window.print()">📄 A4 인쇄 / PDF 저장</button></div>
</body>
</html>"""
def render_salary_tab():
    _salary_init_state()
    st.markdown('<div class="sandbox-title">📋 교무호봉획정</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="apple-note"><strong>기간제교원 호봉(재)획정</strong> · '
        '원본 HTML/Apps Script의 입력·계산 흐름을 Streamlit용으로 재해석했습니다. '
        '계산 단위는 1년=360일, 1개월=30일입니다.</div>',
        unsafe_allow_html=True,
    )
    with st.container(border=True):
        st.markdown("### 🤖 AI 서류 자동 분석")
        st.caption("경력증명서·인사기록카드·자격증 등의 PDF/이미지를 올리면 Gemini가 성명, 기산호봉, 사범계 여부, 추가 학위, 경력 정보를 분석합니다.")
        ai_file = st.file_uploader(
            "분석 파일",
            type=["pdf", "png", "jpg", "jpeg", "webp"],
            key="salary_ai_file",
            label_visibility="collapsed",
        )
        c1, c2 = st.columns([3, 1])
        with c1:
            st.caption("PDF/이미지는 50MB 이하로 사용하세요. AI 결과는 최종 계산 전에 반드시 직접 확인하세요.")
        with c2:
            if st.button("⚡ AI 자동 입력 실행", type="primary", width="stretch", key="salary_ai_run"):
                if ai_file is None:
                    st.error("분석할 PDF 또는 이미지 파일을 선택해 주세요.")
                else:
                    try:
                        with st.spinner("Gemini AI가 서류 데이터를 분석하고 있습니다..."):
                            data = _salary_analyze_uploaded_file(ai_file)
                            _salary_apply_ai_result(data)
                        st.success("AI 분석이 완료되어 입력 항목에 자동 반영되었습니다.")
                        st.rerun()
                    except Exception as exc:
                        st.error(f"AI 분석 실패: {exc}")
        if st.session_state.get("salary_ai_message"):
            st.info(st.session_state.salary_ai_message)
    with st.container(border=True):
        st.markdown("### 👤 기본 인적사항 및 조서 출력 정보")
        r1, r2, r3, r4 = st.columns(4)
        with r1:
            st.session_state.salary_name = st.text_input("성명", value=st.session_state.salary_name, key="salary_name_input")
        with r2:
            st.session_state.salary_calc_date = st.date_input(
                "임용(호봉획정)일",
                value=_salary_safe_date(st.session_state.salary_calc_date) or _today_kst(),
                key="salary_calc_date_input",
            )
        with r3:
            st.session_state.salary_org = st.text_input("소속 기관", value=st.session_state.salary_org, key="salary_org_input")
        with r4:
            st.session_state.salary_position = st.text_input("직위/직급", value=st.session_state.salary_position, key="salary_position_input")
        r5, r6 = st.columns(2)
        with r5:
            st.session_state.salary_writer_name = st.text_input("작성자 성명", value=st.session_state.salary_writer_name, key="salary_writer_name_input")
        with r6:
            st.session_state.salary_writer_pos = st.text_input("작성자 직급", value=st.session_state.salary_writer_pos, key="salary_writer_pos_input")
    with st.container(border=True):
        st.markdown("### 🎓 교원자격증 (기산호봉 결정)")
        base_labels = [f"{value}호봉 · {label}" for value, label in SALARY_BASE_OPTIONS]
        current = max(0, min(int(st.session_state.get("salary_base_idx", 7)), len(base_labels) - 1))
        selected_base = st.radio("기산호봉", base_labels, index=current, key="salary_base_radio", label_visibility="collapsed")
        st.session_state.salary_base_idx = base_labels.index(selected_base)
        selected_base_value, selected_qual_label = SALARY_BASE_OPTIONS[st.session_state.salary_base_idx]
        st.caption("원본 분류: 9호봉 기산 / 8호봉 기산 / 5호봉 기산 교원자격증")
    with st.container(border=True):
        st.markdown("### 📚 최초(기본) 학력 (학부 기준 · 학령 계산)")
        academic_labels = [f"{label} · {desc}" for value, label, desc in SALARY_ACADEMIC_OPTIONS]
        current = max(0, min(int(st.session_state.get("salary_academic_idx", 1)), len(academic_labels) - 1))
        selected_acad = st.radio("최초 학력", academic_labels, index=current, key="salary_academic_radio", label_visibility="collapsed")
        st.session_state.salary_academic_idx = academic_labels.index(selected_acad)
        selected_acad_value, selected_acad_label, _ = SALARY_ACADEMIC_OPTIONS[st.session_state.salary_academic_idx]
    with st.container(border=True):
        head1, head2 = st.columns([4, 1])
        with head1:
            st.markdown("### 🎓 추가 학위 (복수 대학 졸업 · 석·박사 등)")
            st.caption("동등 수준 추가 학위는 원본 HTML과 같이 환산율 기본값 80%를 사용합니다.")
        with head2:
            if st.button("+ 추가 학위 등록", key="salary_add_degree", width="stretch"):
                _salary_add_degree_row()
                st.rerun()
        degrees = _salary_normalize_table(st.session_state.salary_degrees, ["학위 구분", "학교/전공 세부명", "입학일", "졸업일", "환산율"])
        degree_epoch = int(st.session_state.get("salary_degree_editor_epoch", 0) or 0)
        degree_snapshot = st.session_state.get("salary_degree_rate_snapshot", _salary_rate_snapshot(degrees, "학위 구분", "환산율"))
        edited_degrees = st.data_editor(
            degrees,
            num_rows="dynamic",
            width="stretch",
            hide_index=True,
            key=f"salary_degree_editor_{degree_epoch}",
            column_config={
                "학위 구분": st.column_config.SelectboxColumn("학위 구분", options=SALARY_DEGREE_TYPES, required=True),
                "학교/전공 세부명": st.column_config.TextColumn("학교/전공 세부명"),
                "입학일": st.column_config.DateColumn("입학일"),
                "졸업일": st.column_config.DateColumn("졸업일"),
                "환산율": st.column_config.NumberColumn("환산율", min_value=0, max_value=100, step=1),
            },
        )
        synced_degrees, degree_rate_changed = _salary_sync_rate_defaults(
            edited_degrees, "학위 구분", "환산율", SALARY_DEGREE_RATE_BY_TYPE, 80, degree_snapshot
        )
        st.session_state.salary_degrees = synced_degrees
        if degree_rate_changed:
            st.session_state.salary_degree_rate_snapshot = _salary_rate_snapshot(synced_degrees, "학위 구분", "환산율")
            st.session_state.salary_degree_editor_epoch = degree_epoch + 1
            st.rerun()
        else:
            st.session_state.salary_degree_rate_snapshot = _salary_rate_snapshot(synced_degrees, "학위 구분", "환산율")
    with st.container(border=True):
        st.markdown("### ➕ 가산연수 해당 여부")
        st.session_state.salary_is_sabom = st.checkbox(
            "사범계학교(대학에 설치된 교육학과 포함) 졸업자",
            value=bool(st.session_state.salary_is_sabom),
            key="salary_sabom_checkbox",
            help="수학연한 2년 이상의 사범계 학교 졸업 시 +1년 적용",
        )
        if st.session_state.salary_is_sabom:
            st.success("💡 사범계 가산연수 적용: +1년")
        else:
            st.info("💡 가산연수 미적용: +0년")
    with st.container(border=True):
        head1, head2 = st.columns([4, 1])
        with head1:
            st.markdown("### 📂 경력사항 입력")
        with head2:
            if st.button("+ 경력 추가", key="salary_add_career", width="stretch"):
                _salary_add_career_row()
                st.rerun()
        careers = _salary_normalize_table(
            st.session_state.salary_careers,
            ["경력 종류", "세부 근무처/직위", "시작일", "종료일", "종료일 산입", "환산율 (%)"],
        )
        career_epoch = int(st.session_state.get("salary_career_editor_epoch", 0) or 0)
        career_snapshot = st.session_state.get("salary_career_rate_snapshot", _salary_rate_snapshot(careers, "경력 종류", "환산율 (%)"))
        edited_careers = st.data_editor(
            careers,
            num_rows="dynamic",
            width="stretch",
            hide_index=True,
            key=f"salary_career_editor_{career_epoch}",
            column_config={
                "경력 종류": st.column_config.SelectboxColumn("경력 종류", options=SALARY_CAREER_TYPES, required=True),
                "세부 근무처/직위": st.column_config.TextColumn("세부 근무처/직위"),
                "시작일": st.column_config.DateColumn("시작일"),
                "종료일": st.column_config.DateColumn("종료일"),
                "종료일 산입": st.column_config.CheckboxColumn("종료일 산입", default=True),
                "환산율 (%)": st.column_config.NumberColumn("환산율 (%)", min_value=0, max_value=100, step=1),
            },
        )
        synced_careers, career_rate_changed = _salary_sync_rate_defaults(
            edited_careers, "경력 종류", "환산율 (%)", SALARY_CAREER_RATE_BY_TYPE, 100, career_snapshot
        )
        st.session_state.salary_careers = synced_careers
        if career_rate_changed:
            st.session_state.salary_career_rate_snapshot = _salary_rate_snapshot(synced_careers, "경력 종류", "환산율 (%)")
            st.session_state.salary_career_editor_epoch = career_epoch + 1
            st.rerun()
        else:
            st.session_state.salary_career_rate_snapshot = _salary_rate_snapshot(synced_careers, "경력 종류", "환산율 (%)")
        st.caption("원본 Index.html / Code.gs 기준 환산율: 국·공립 100% · 사립 100% · 자격 불일치 기간제 80% · 전일제 강사 100% · 시간제 강사 30% · 국가·지방공무원 100% · 학원/교습소 50% · 회사 40%")
        st.caption("경력 종류를 변경하면 해당 기준 환산율이 자동 적용되며, 마지막 '환산율 (%)' 값은 필요 시 직접 조정할 수 있습니다.")
        st.caption("※ 종료일 산입: 기간제 계약만료일·군 전역일 등 만료일을 포함하는 경력에만 체크하세요. 일반 퇴직일은 원칙적으로 제외합니다.")
    ccalc1, ccalc2 = st.columns([4, 1])
    with ccalc1:
        st.caption("입력된 학위·경력 중 시작일과 종료일이 모두 있는 항목만 실제 환산일수가 계산됩니다.")
    with ccalc2:
        do_calculate = st.button("📊 호봉 계산 및 조서 생성", type="primary", width="stretch", key="salary_calculate")
    if do_calculate:
        try:
            st.session_state.salary_result = _salary_calculate(
                st.session_state.salary_name,
                _salary_safe_date(st.session_state.salary_calc_date),
                st.session_state.salary_org,
                st.session_state.salary_position,
                st.session_state.salary_writer_name,
                st.session_state.salary_writer_pos,
                selected_base_value,
                selected_qual_label,
                selected_acad_value,
                selected_acad_label,
                bool(st.session_state.salary_is_sabom),
                st.session_state.salary_degrees,
                st.session_state.salary_careers,
            )
        except Exception as exc:
            st.error(f"호봉 계산 중 오류가 발생했습니다: {exc}")
    result = st.session_state.get("salary_result")
    if not result:
        return
    with st.container(border=True):
        st.markdown("## 호봉 획정 (재획정) 조서")
        meta1, meta2, meta3, meta4 = st.columns(4)
        meta1.metric("소속", result["org"] or "-")
        meta2.metric("직위", result["position"] or "-")
        meta3.metric("성명", result["name"] or "-")
        meta4.metric("교사(자격)", result["qual_label"] or "-")
        st.markdown("#### ■ 학위 수학기간 및 경력 내용 환산 상세")
        detail_rows = [{
            "구분 (학위/경력 내용)": item["구분"],
            "시작일": item["시작일"] or "-",
            "종료일": item["종료일"] or "-",
            "원기간": item["원기간"],
            "환산율": item["환산율"],
            "환산 경력(년)": item["환산년"],
            "환산 경력(월)": item["환산월"],
            "환산 경력(일)": item["환산일"],
        } for item in result["details"]]
        if detail_rows:
            st.dataframe(pd.DataFrame(detail_rows), width="stretch", hide_index=True)
        else:
            st.info("학위 및 경력 사항이 없습니다.")
        st.markdown("#### ■ 호봉 산정 결과")
        st.dataframe(pd.DataFrame([{
            "기산호봉": f"{result['base_salary']}호봉",
            "학령가감": f"{'+' if result['academic_value'] >= 0 else ''}{result['academic_value']}년",
            "가산연수": f"+{result['add_years']}년",
            "환산경력 (학위+경력)": f"{result['total_y']}년 {result['total_m']}월 {result['total_d']}일",
            "최종 초임호봉": f"{result['final_step']}호봉",
            "잔여 기간": f"{result['rem_m']}월 {result['rem_d']}일",
        }]), width="stretch", hide_index=True)
        calc_date_text = (
            result["calc_date"].strftime("%Y년 %m월 %d일")
            if isinstance(result.get("calc_date"), date)
            else str(result.get("calc_date", ""))
        )
        st.markdown(
            f'<div class="work-note"><strong>위와 같이 호봉을 획정(재획정)함.</strong><br>'
            f'일자: {html_lib.escape(calc_date_text)}'
            f'<span style="float:right">작성자 직급: {html_lib.escape(str(result["writer_pos"]))} · '
            f'작성자 성명: {html_lib.escape(str(result["writer_name"]))} (인)</span></div>',
            unsafe_allow_html=True,
        )
        report_html = _salary_report_html(result)
        file_date = result["calc_date"].strftime("%Y%m%d") if isinstance(result.get("calc_date"), date) else "일자미상"
        st.download_button(
            "📄 A4 인쇄용 HTML / PDF 저장",
            data=report_html.encode("utf-8"),
            file_name=f"호봉획정조서_{result.get('name') or '교원'}_{file_date}.html",
            mime="text/html",
            width="stretch",
            key="salary_report_download",
        )
        st.caption("다운로드한 HTML을 브라우저에서 열어 인쇄 → PDF로 저장하면 A4 조서로 사용할 수 있습니다.")
@st.cache_resource(show_spinner=False)
def _login_failure_store():
    return {"lock": threading.Lock(), "failures": {}}
def _login_client_key(uid):
    context = getattr(st, "context", None)
    headers = getattr(context, "headers", {}) if context is not None else {}
    forwarded = str(headers.get("X-Forwarded-For", "")).split(",")[0].strip()
    real_ip = str(headers.get("X-Real-IP", "")).strip()
    client = forwarded or real_ip
    nonce = st.session_state.setdefault("_login_nonce", py_secrets.token_hex(16))
    seed = f"{client or nonce}|{uid.strip().lower()}"
    return hashlib.sha256(seed.encode("utf-8")).hexdigest()
def _login_is_rate_limited(uid):
    store = _login_failure_store()
    key = _login_client_key(uid)
    now = time.monotonic()
    with store["lock"]:
        item = store["failures"].get(key)
        if not item:
            return False
        failures, blocked_until = item
        if blocked_until > now:
            return True
        if now - blocked_until > 900:
            store["failures"].pop(key, None)
    return False
def _record_login_failure(uid):
    store = _login_failure_store()
    key = _login_client_key(uid)
    now = time.monotonic()
    with store["lock"]:
        failures, _ = store["failures"].get(key, (0, 0.0))
        failures += 1
        blocked_until = now + (60.0 if failures >= MAX_LOGIN_ATTEMPTS else 0.0)
        store["failures"][key] = (failures, blocked_until)
def _clear_login_failures(uid):
    store = _login_failure_store()
    key = _login_client_key(uid)
    with store["lock"]:
        store["failures"].pop(key, None)
def _revalidate_authenticated_user(force=False):
    if not st.session_state.get("logged_in") or current_role() == ROLE_GUEST:
        return True
    now = time.monotonic()
    last = float(st.session_state.get("_auth_last_check", 0.0) or 0.0)
    if not force and now - last < AUTH_RECHECK_SECONDS:
        return True
    ids = load_id_sheet()
    if ids.empty or "아이디" not in ids.columns:
        st.session_state["_auth_stale"] = True
        st.session_state["_auth_last_check"] = now
        return True
    uid = str(st.session_state.get("user_id", "")).strip()
    match = ids[ids["아이디"].astype(str).str.strip() == uid]
    if match.empty:
        for key in list(st.session_state.keys()):
            del st.session_state[key]
        st.session_state.logged_in = False
        st.session_state.user_role = ROLE_GUEST
        return False
    row = match.iloc[0]
    role = _validated_role(row.get("권한", ROLE_TEACHER))
    st.session_state.user_role = role
    st.session_state.user_name = str(row.get("이름", "")).strip() or uid
    st.session_state.user_allowed_tabs = str(row.get("허용탭", "")).strip()
    st.session_state["_auth_last_check"] = now
    st.session_state["_auth_stale"] = False
    return True
def show_login_page():
    if "login_attempts" not in st.session_state:
        st.session_state.login_attempts = 0
    if "login_locked" not in st.session_state:
        st.session_state.login_locked = False
    IMAGE_URL = "https://i.imgur.com/Gl0YDO3.jpeg"
    st.markdown(
        f"""
        <div style="background:rgba(255,255,255,.035); padding:20px; border:1px solid rgba(255,255,255,.08); border-radius: 18px; text-align: center; margin-bottom: 20px;">
            <img src="{IMAGE_URL}" width="150" style="object-fit: contain;">
        </div>
        """,
        unsafe_allow_html=True
    )
    st.markdown(f"<div class='ai-hero-kicker'>{SCHOOL_YEAR} · SCHOOL OPERATIONS</div>", unsafe_allow_html=True)
    st.markdown(f"<div class='ai-hero-title'>{SCHOOL_NAME}</div>", unsafe_allow_html=True)
    st.markdown("<div class='ai-hero-sub'>시간표와 결보강 업무를 하나의 흐름으로 연결하는<br>학교 업무 운영 시스템</div>", unsafe_allow_html=True)
    st.markdown("<div class='ai-prompt-glow'>", unsafe_allow_html=True)
    if st.session_state.login_locked:
        st.error(f"🚫 로그인 시도가 {MAX_LOGIN_ATTEMPTS}회를 초과하여 차단되었습니다.")
        st.stop()
    id_input = st.text_input("아이디", placeholder="아이디를 입력하세요", key="login_id")
    st.markdown("</div>", unsafe_allow_html=True)
    col1, col2 = st.columns(2)
    with col1:
        if st.button("로그인", type="primary", width="stretch"):
            uid = _clean_user_text(id_input, MAX_ID_LENGTH)
            if not uid:
                st.error("아이디를 입력해주세요.")
            elif _login_is_rate_limited(uid):
                st.error("로그인 시도가 잠시 제한되었습니다. 1분 후 다시 시도해주세요.")
            else:
                ids = load_id_sheet()
                if ids.empty:
                    detail = st.session_state.get("_id_sheet_error") or st.session_state.get("_gsheet_last_error", "")
                    st.error("아이디 정보를 불러오지 못했습니다. Google Sheets 연결/권한을 확인해주세요.")
                    if detail:
                        st.caption(detail)
                else:
                    match = ids[ids["아이디"].astype(str).str.strip() == uid]
                if not ids.empty and not match.empty:
                    st.session_state.login_attempts = 0
                    _clear_login_failures(uid)
                    row = match.iloc[0]
                    role = _validated_role(row.get("권한", ROLE_TEACHER))
                    name = _clean_user_text(row.get("이름", ""), MAX_NAME_LENGTH) or uid
                    allowed = str(row.get("허용탭", "")).strip()
                    st.session_state.logged_in = True
                    st.session_state.user_id = uid
                    st.session_state.user_name = name
                    st.session_state.user_role = role
                    st.session_state.user_allowed_tabs = allowed
                    st.rerun()
                else:
                    st.session_state.login_attempts += 1
                    _record_login_failure(uid)
                    remaining = MAX_LOGIN_ATTEMPTS - st.session_state.login_attempts
                    if st.session_state.login_attempts >= MAX_LOGIN_ATTEMPTS:
                        st.session_state.login_locked = True
                        st.error(f"🚫 로그인 실패 {MAX_LOGIN_ATTEMPTS}회 초과로 차단되었습니다.")
                        st.rerun()
                    else:
                        st.error(f"등록되지 않은 아이디입니다. (남은 시도: {remaining}회)")
    with col2:
        if st.button("게스트로 입장", width="stretch"):
            st.session_state.logged_in = True
            st.session_state.user_id = "게스트"
            st.session_state.user_name = "게스트"
            st.session_state.user_role = ROLE_GUEST
            st.session_state.user_allowed_tabs = ""
            st.session_state.login_attempts = 0
            st.rerun()
    if st.session_state.login_attempts > 0:
        st.warning(f"현재 로그인 실패 횟수: {st.session_state.login_attempts} / {MAX_LOGIN_ATTEMPTS}")
    st.divider()
    st.markdown("#### 📝 아이디 추가 요청 (게스트용)")
    with st.form("id_request_form"):
        name = st.text_input("이름 *")
        email = st.text_input("이메일 *")
        desired = st.text_input("추가하고 싶은 아이디 *")
        memo = st.text_area("메모 (요청 사유 등)")
        submitted = st.form_submit_button("요청 제출", type="primary")
        if submitted:
            if not (name and email and desired):
                st.error("이름, 이메일, 아이디는 필수입니다.")
            else:
                try:
                    save_id_request(name, email, desired, memo)
                    st.success("요청이 정상적으로 접수되었습니다.")
                except Exception as exc:
                    st.error("아이디 추가 요청을 저장하지 못했습니다. Google Sheets 연결/권한을 확인해주세요.")
                    st.caption(str(exc))
def _week_anchor(ref_date=None):
    ref = ref_date or _today_kst()
    if not isinstance(ref, date):
        try:
            ref = datetime.strptime(normalize_date_str(ref), "%Y-%m-%d").date()
        except Exception:
            ref = _today_kst()
    return ref - timedelta(days=ref.weekday())
if "logged_in" not in st.session_state:
    st.session_state.logged_in = False
    st.session_state.user_id = None
    st.session_state.user_name = ""
    st.session_state.user_role = ROLE_GUEST
    st.session_state.user_allowed_tabs = ""
if not st.session_state.logged_in:
    show_login_page()
    st.stop()
if not _revalidate_authenticated_user():
    st.error("세션의 계정 권한이 변경되었거나 계정이 삭제되었습니다. 다시 로그인해주세요.")
    st.stop()
allowed_tabs_preview = get_user_allowed_tabs()
visible_tabs_preview = [t for t in ALL_APP_TABS if t in allowed_tabs_preview]
if current_role() != ROLE_GUEST and "교무호봉획정" not in visible_tabs_preview:
    visible_tabs_preview.append("교무호봉획정")
if can_manage_ids():
    for t in ["🔑 아이디·권한 관리", "📑 회원별 탭 권한 관리", "🛠️ 다중 출장·전체 조정 추천"]:
        if t not in visible_tabs_preview:
            visible_tabs_preview.append(t)
if "active_tab" not in st.session_state or st.session_state.active_tab not in visible_tabs_preview:
    if "교무호봉획정" in visible_tabs_preview and current_role() == ROLE_OFFICE:
        st.session_state.active_tab = "교무호봉획정"
        st.session_state.app_category = "🏫 교무"
    else:
        st.session_state.active_tab = visible_tabs_preview[0] if visible_tabs_preview else None
        st.session_state.app_category = "📚 수업" if st.session_state.active_tab in ALL_TABS else "🏫 교무"
if st.session_state.get("active_tab") in ALL_TABS and not init_state():
    st.error("⚠️ 시간표 초기화에 실패했습니다.")
    st.warning(st.session_state.get("_load_error", "시간표 데이터를 불러오지 못했습니다."))
    st.caption("수업 업무 실행을 중단했습니다. 교무행정 업무는 시간표와 독립적으로 사용할 수 있습니다.")
    if st.button("🔄 시간표 다시 연결", type="primary", key="fatal_reload_timetable"):
        _clear_gsheet_runtime_cache()
        st.session_state.pop("_load_error", None)
        st.rerun()
    st.stop()
if "ui_font" not in st.session_state or st.session_state.ui_font not in UI_FONT_OPTIONS:
    st.session_state.ui_font = UI_FONT_DEFAULT
st.markdown(
    f"""<style id="dialog-font-state">
:root {{ --app-font: {UI_FONT_OPTIONS[st.session_state.ui_font]}; }}
html,body,[data-testid="stAppViewContainer"],[data-testid="stAppViewContainer"] *{{font-family:var(--app-font)!important}}
body table,body table *,[data-testid="stDataFrame"],[data-testid="stDataFrame"] *,[data-testid="stDataEditor"],[data-testid="stDataEditor"] *{{font-family:var(--app-font)!important;--gdg-font-family:var(--app-font)!important}}
[class*="material-symbols"],[data-testid="stIconMaterial"],[data-testid="stIconMaterial"] *{{font-family:"Material Symbols Rounded","Material Symbols Outlined",sans-serif!important}}
</style>""",
    unsafe_allow_html=True,
)
st.markdown('<div class="streamlit-header-safe-space" aria-hidden="true"></div>', unsafe_allow_html=True)
@st.dialog("도구", width="small")
def render_tools_dialog():
    if current_role() == ROLE_GUEST:
        st.caption("게스트 모드에서는 사용할 수 있는 도구가 없습니다.")
        return
    st.markdown('<div class="tool-section-title">화면</div>', unsafe_allow_html=True)
    selected_font = st.selectbox(
        "글꼴",
        list(UI_FONT_OPTIONS.keys()),
        index=list(UI_FONT_OPTIONS.keys()).index(st.session_state.get("ui_font", UI_FONT_DEFAULT)),
        key="ui_font_selector",
        help="이 브라우저에서 사용할 수 있는 글꼴을 우선 적용합니다. 학교 PC에 해당 글꼴이 설치되어 있지 않으면 다음 대체 글꼴이 사용됩니다.",
    )
    st.session_state.ui_font = selected_font
    render_font_runtime_css(selected_font)
    if selected_font == "Noto Sans KR · GitHub (Light + Bold)":
        st.caption("GitHub의 NotoSansKR-Light.ttf(일반체)와 NotoSansKR-Bold.ttf(볼드체)를 사용합니다.")
        st.markdown(
            f"<div style=\"font-family:'{GITHUB_FONT_FAMILY}';font-weight:300;font-size:14px;line-height:1.7\">가나다라마바사 아자차카 · 일반체 <strong style=\"font-weight:700\">가나다라마바사 아자차카 · 볼드체</strong></div>",
            unsafe_allow_html=True,
        )
    elif selected_font == "학교안심 우주체 · GitHub":
        st.caption("GitHub의 HakgyoansimWoojuR.ttf를 사용합니다.")
        st.markdown(
            f"<div style=\"font-family:'{HAKYO_FONT_FAMILY}';font-size:16px;line-height:1.8\">학교안심 우주체 미리보기 · 가나다라마바사 아자차카 12345</div>",
            unsafe_allow_html=True,
        )
    else:
        st.caption("글꼴은 이 기기에 설치된 폰트를 우선 사용합니다.")
    st.divider()
    if can_full_data() or is_teacher():
        st.caption("① 새로고침 → ② 불러오기 → ③ 현재 작업 저장")
        st.divider()
        a, b = st.columns(2)
        if a.button("↩ Undo", width="stretch", key="top_undo"):
            if undo():
                if save_work_data_to_gsheet():
                    st.rerun()
                else:
                    redo()
                    st.error("Undo 저장에 실패하여 이전 상태로 복구했습니다.")
        if b.button("↪ Redo", width="stretch", key="top_redo"):
            if redo():
                if save_work_data_to_gsheet():
                    st.rerun()
                else:
                    undo()
                    st.error("Redo 저장에 실패하여 이전 상태로 복구했습니다.")
    st.divider()
    st.markdown('<div class="tool-section-title">출력</div>', unsafe_allow_html=True)
    if st.button("결보강 계획서", width="stretch", key="top_personal_plan_open"):
        st.session_state["top_output_mode"] = "personal_plan"
    if is_edu_or_master() and st.button("전체 일일 내역서", width="stretch", key="top_daily_report_open"):
        st.session_state["top_output_mode"] = "daily_report"
    output_mode = st.session_state.get("top_output_mode")
    if output_mode == "personal_plan":
        plan_date = st.date_input("기준일", value=_today_kst(), key="top_plan_date", label_visibility="collapsed")
        if st.button("파일 만들기", type="primary", width="stretch", key="top_personal_generate"):
            html = build_personal_plan_html(current_name(), plan_date.strftime("%Y-%m-%d"))
            st.download_button("HTML 다운로드", html.encode("utf-8"), f"결보강계획서_{current_name()}_{plan_date}.html", "text/html", key="top_dl_personal")
    elif output_mode == "daily_report" and is_edu_or_master():
        rd = st.date_input("기준일", value=_today_kst(), key="top_report_date", label_visibility="collapsed")
        if st.button("파일 만들기", type="primary", width="stretch", key="top_daily_generate"):
            day = rd.strftime("%Y-%m-%d")
            html = build_report_html(day)
            xls = to_excel_bytes({
                "결강": st.session_state.absences[st.session_state.absences["일자"] == day] if not st.session_state.absences.empty else pd.DataFrame(),
                "보강": st.session_state.subs[st.session_state.subs["일자"] == day] if not st.session_state.subs.empty else pd.DataFrame(),
                "맞교환": st.session_state.swaps,
            })
            st.download_button("HTML 다운로드", html.encode("utf-8"), f"내역서_{rd}.html", "text/html", key="top_dl_report_html")
            st.download_button("엑셀 다운로드", xls, f"내역서_{rd}.xlsx", key="top_dl_report_xlsx")
def _load_all_runtime_data_from_gsheet():
    _clear_gsheet_runtime_cache()
    ti, tt = load_timetable_from_gsheet()
    absences, subs, swaps, part_time, cumulative, duties = load_work_data_from_gsheet()
    st.session_state.teachers = ti
    st.session_state.timetable = tt
    st.session_state.absences = ensure_input_user(absences)
    st.session_state.subs = ensure_input_user(subs)
    st.session_state.swaps = ensure_input_user(swaps)
    st.session_state.part_time = ensure_part_time_columns(part_time)
    st.session_state.cumulative = cumulative
    st.session_state.duties = ensure_duty_columns(duties)
    _invalidate_all_caches()
def render_top_toolbar(visible_tabs):
    category_map = {
        "📚 수업": [t for t in ALL_TABS if t in visible_tabs],
        "🏫 교무": OFFICE_TABS[:] if current_role() != ROLE_GUEST else [],
    }
    available_categories = [c for c, tabs in category_map.items() if tabs]
    if not available_categories:
        return
    old_active = st.session_state.get("active_tab")
    if "app_category" not in st.session_state or st.session_state.app_category not in available_categories:
        if old_active in category_map.get("🏫 교무", []):
            st.session_state.app_category = "🏫 교무"
        else:
            st.session_state.app_category = available_categories[0]
    top_id, top_nav, top_quick, top_tools, top_user = st.columns([1.35, 4.65, 3.05, 0.72, 0.72], vertical_alignment="center")
    with top_id:
        st.markdown(
            f'<div class="app-identity"><strong>{html_lib.escape(str(current_name() or current_user()))}</strong> · {html_lib.escape(str(current_role()))}</div>',
            unsafe_allow_html=True,
        )
    with top_nav:
        category_cols = st.columns(2, gap="small")
        for category, col in zip(available_categories, category_cols):
            with col:
                is_active = st.session_state.app_category == category
                if st.button(
                    category,
                    key=f"app_category_btn_{category}",
                    width="stretch",
                    type="primary" if is_active else "secondary",
                    help="현재 선택된 업무 영역" if is_active else f"{category} 업무로 이동",
                ):
                    if category != st.session_state.app_category:
                        st.session_state.app_category = category
                        sub_tabs = category_map.get(category, [])
                        st.session_state.active_tab = sub_tabs[0] if sub_tabs else None
                        _clear_weekly_selection()
                        st.session_state.pop("weekly_dialog_use_test", None)
                        st.session_state.pop("weekly_dialog_title", None)
                        st.session_state.pop("weekly_dialog_instance", None)
                        st.rerun()
    with top_quick:
        q1, q2, q3 = st.columns(3, gap="small")
        with q1:
            if st.button("① 새로고침", width="stretch", key="top_quick_refresh", help="앱 화면과 캐시를 새로고침합니다."):
                _clear_gsheet_runtime_cache()
                _invalidate_all_caches()
                st.toast("화면을 새로고침했습니다.", icon="🔄")
                st.rerun()
        with q2:
            if st.button("② 불러오기", width="stretch", key="top_quick_load", help="Google Sheets의 최신 시간표와 업무 데이터를 불러옵니다."):
                try:
                    with st.spinner("최신 데이터를 불러오는 중..."):
                        _load_all_runtime_data_from_gsheet()
                    st.toast("최신 데이터를 불러왔습니다.", icon="⬇️")
                    st.rerun()
                except Exception as exc:
                    st.error(f"데이터를 불러오지 못했습니다: {exc}")
        with q3:
            if st.button("③ 작업 저장", width="stretch", type="primary", key="top_quick_save", help="현재 작업을 Google Sheets에 저장합니다."):
                try:
                    ok = save_work_data_to_gsheet()
                    if ok is not False:
                        st.toast("현재 작업을 저장했습니다.", icon="✅")
                except Exception as exc:
                    st.error(f"현재 작업 저장 중 오류가 발생했습니다: {exc}")
    with top_tools:
        if st.button("···", width="stretch", key="top_tools_open", help="화면·출력·기타 도구"):
            render_tools_dialog()
    with top_user:
        if st.button("↪", width="stretch", key="top_logout", help="로그아웃"):
            for k in list(st.session_state.keys()):
                del st.session_state[k]
            st.rerun()
    selected_category_tabs = category_map.get(st.session_state.app_category, [])
    if not selected_category_tabs:
        return
    if "active_tab" not in st.session_state or st.session_state.active_tab not in selected_category_tabs:
        st.session_state.active_tab = selected_category_tabs[0]
    sub_labels = [NAV_LABELS.get(t, t) for t in selected_category_tabs]
    sub_map = dict(zip(sub_labels, selected_category_tabs))
    current_label = NAV_LABELS.get(st.session_state.active_tab, st.session_state.active_tab)
    picked_sub = st.radio(
        "세부 업무",
        sub_labels,
        index=sub_labels.index(current_label) if current_label in sub_labels else 0,
        horizontal=True,
        key=f"app_sub_nav_{st.session_state.app_category}",
        label_visibility="collapsed",
    )
    picked_tab = sub_map.get(picked_sub, selected_category_tabs[0])
    if picked_tab != st.session_state.active_tab:
        st.session_state.active_tab = picked_tab
        _clear_weekly_selection()
        st.session_state.pop("weekly_dialog_use_test", None)
        st.session_state.pop("weekly_dialog_title", None)
        st.session_state.pop("weekly_dialog_instance", None)
        st.rerun()
if current_role() == ROLE_GUEST:
    st.title("게스트 모드")
    st.info("현재 게스트로 접속 중입니다. 아이디 추가 요청만 가능합니다.")
    with st.form("guest_request"):
        name = st.text_input("이름 *")
        email = st.text_input("이메일 *")
        desired = st.text_input("추가하고 싶은 아이디 *")
        memo = st.text_area("메모")
        if st.form_submit_button("요청 제출", type="primary"):
            if name and email and desired:
                save_id_request(name, email, desired, memo)
                st.success("요청이 접수되었습니다.")
            else:
                st.error("필수 항목을 입력해주세요.")
    st.stop()
allowed_tabs = get_user_allowed_tabs()
visible_tabs = [t for t in ALL_APP_TABS if t in allowed_tabs]
if current_role() != ROLE_GUEST and "교무호봉획정" not in visible_tabs:
    visible_tabs.append("교무호봉획정")
if can_manage_ids():
    for t in ["🔑 아이디·권한 관리", "📑 회원별 탭 권한 관리", "🛠️ 다중 출장·전체 조정 추천"]:
        if t not in visible_tabs:
            visible_tabs.append(t)
if not visible_tabs:
    st.warning("접근 가능한 탭이 없습니다.")
    st.stop()
render_top_toolbar(visible_tabs)
active_tab = st.session_state.active_tab
tab_map = {active_tab: st.container()}
PAGE_DESCRIPTIONS = {
    "시간표 조회": "오늘과 주간 시간표를 한 곳에서 빠르게 확인합니다.",
    "시간강사 관리": "시간강사 배정과 가용 시간을 관리합니다.",
    "결강·보강": "결강을 기록하고 보강 교사를 배정합니다.",
    "시간표 맞교환 & 변경 추천": "수업 교환과 가능한 대안을 확인합니다.",
    "통계": "결보강 및 시간표 변경 현황을 요약합니다.",
    "시간표 변경 테스트용": "실제 시간표에 반영하기 전 변경안을 검토합니다.",
    "변경된 교사 주간표": "이번 주 변경된 교사의 실제 주간 시간표를 확인합니다.",
    "📋 복무 관리 & 판단": "복무와 출장에 따른 조정 가능 수업을 확인합니다.",
    "🛠️ 다중 출장·전체 조정 추천": "여러 교사의 부재 상황을 함께 조정합니다.",
    "🔑 아이디·권한 관리": "사용자 계정과 역할을 관리합니다.",
    "📑 회원별 탭 권한 관리": "사용자별 업무 메뉴 접근 권한을 관리합니다.",
    "교무호봉획정": "기간제교원 호봉(재)획정과 조서 출력을 처리합니다.",
}
st.markdown(
    f'<div class="work-page-head"><div>'
    f'<div style="font-size:11px;color:#86868b;margin-bottom:6px;letter-spacing:.01em">{SCHOOL_NAME} · {SCHOOL_YEAR}</div>'
    f'<h1>{html_lib.escape(str(NAV_LABELS.get(active_tab, active_tab)))}</h1>'
    f'<p>{PAGE_DESCRIPTIONS.get(active_tab, "학교 시간표와 결보강 업무를 관리합니다.")}</p></div>'
    f'<div class="work-page-meta">{html_lib.escape(str(current_role()))} · {html_lib.escape(str(current_name() or current_user()))}</div></div>',
    unsafe_allow_html=True
)
if "시간표 조회" in tab_map:
    with tab_map["시간표 조회"]:
        view = st.radio("보기 방식", ["선택 날짜", "교사별 주간", "학급별 주간", "교사 1인"], horizontal=True, key="view_mode", label_visibility="collapsed")
        ver = st.session_state.get("_data_version", 0)
        if view == "선택 날짜":
            picked, matrix, selections = daily_schedule_picker(_today_kst(), key="view_daily", height=560,
                help_text="날짜를 달력에서 선택한 후 교사×교시 셀을 클릭하세요.")
            if selections:
                st.markdown("#### 선택 수업 상세")
                e=get_effective_timetable_for_date(picked.strftime("%Y-%m-%d"),ver)
                details=[]
                for sel in selections:
                    m=e[(e["교사명"]==sel["교사명"]) & (e["교시"].apply(safe_int)==sel["교시"])]
                    if not m.empty:
                        r=m.iloc[0]
                        details.append({"교사":r["교사명"],"교시":sel["교시"],"학급":r["학급"],"과목":r["과목"],"변경유형":r.get("변경유형","원본"),"변경상세":r.get("변경상세","")})
                st.dataframe(pd.DataFrame(details),width="stretch",hide_index=True)
        elif view == "교사별 주간":
            ref=calendar_picker("주간 기준일",_today_kst(),key="view_week_ref")
            st.markdown('<div class="matrix-legend"><span>↻ 교환</span><span>+ 보강</span><span>• 시간강사</span></div>', unsafe_allow_html=True)
            render_standard_weekly_matrix(effective_teacher_matrix(ref,ver,use_test=False), ref, row_label='교사명', key='view_teacher_week_matrix', title='교사별 주간 시간표', use_test=False, open_dialog=False)
            xlsx = build_weekly_schedule_excel_bytes(ref, use_test=False)
            st.download_button('📥 이 주간표 Excel 다운로드', xlsx, file_name=f'전체교사_주간시간표_{ref:%Y%m%d}.xlsx', mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', key='weekly_xlsx_teacher')
        elif view == "학급별 주간":
            ref=calendar_picker("주간 기준일",_today_kst(),key="view_class_ref")
            st.markdown('<div class="matrix-legend"><span>과목명 기준 · 조회 전용</span><span>↻ 교환</span><span>+ 보강</span></div>', unsafe_allow_html=True)
            render_standard_weekly_matrix(class_matrix(ver, ref_date=ref), ref, row_label='학급', key='view_class_week_matrix', title='학급별 주간 시간표', use_test=False, open_dialog=False)
            xlsx = build_weekly_class_schedule_excel_bytes(ref, use_test=False)
            st.download_button('📥 이 주간표 Excel 다운로드', xlsx, file_name=f'전체학급_주간시간표_{ref:%Y%m%d}.xlsx', mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', key='weekly_xlsx_class')
        else:
            tlist=get_all_teacher_names()
            t=st.selectbox("교사 선택",tlist,key="view_t")
            ref=calendar_picker("주간 기준일",_today_kst(),key="view_ref")
            teacher_week = effective_teacher_matrix(ref, ver, use_test=False)
            if not teacher_week.empty and t:
                teacher_week = teacher_week[teacher_week["교사명"].astype(str).str.strip() == str(t).strip()].reset_index(drop=True)
            render_standard_weekly_matrix(teacher_week, ref, row_label="교사명", key="view_single_teacher_week_matrix", title=f"{t} 주간 시간표", use_test=False, open_dialog=False)
if "시간강사 관리" in tab_map:
    with tab_map["시간강사 관리"]:
        st.subheader("시간강사 등록 및 가능 시간표")
        st.info("기본 입력은 달력과 가능 시간 매트릭스를 사용합니다. 아래 표는 기존 자료를 일괄 수정할 때 사용하는 고급 편집 영역입니다.")
        if can_full_data() or is_teacher():
            st.markdown("#### 📅 시간강사 빠른 등록")
            pt_names = st.session_state.part_time.get("시간강사명", pd.Series(dtype=str)).dropna().astype(str).str.strip().tolist() if not st.session_state.part_time.empty else []
            default_pt = pt_names[0] if pt_names else ""
            pt_name = st.text_input("시간강사명", value=default_pt, key="quick_pt_name")
            subjects = st.session_state.timetable.get("과목", pd.Series(dtype=str)).dropna().astype(str).unique().tolist()
            pt_subject = st.selectbox("담당과목", [""]+sorted(subjects), key="quick_pt_subject")
            orig_teachers = st.session_state.teachers["교사명"].dropna().astype(str).tolist() if not st.session_state.teachers.empty else []
            pt_orig = st.selectbox("대체교사", orig_teachers, key="quick_pt_orig")
            qstart,qend=calendar_range_picker(_today_kst(),_today_kst(),key="quick_pt_range",help_text="달력에서 대체 기간을 선택합니다.")
            st.markdown("가능한 교시를 클릭하세요")
            qcols=st.columns(5); qdays=DAYS
            avail={}
            for di,dayname in enumerate(qdays):
                with qcols[di]:
                    st.markdown(f"**{dayname}**")
                    for pp in range(1,8):
                        k=f"quick_pt_{dayname}_{pp}"
                        avail[(dayname,pp)]=st.checkbox(f"{pp}",key=k)
            if st.button("달력·매트릭스로 시간강사 추가",key="quick_pt_add",type="primary"):
                if not pt_name.strip() or not pt_orig.strip() or qstart>qend:
                    st.error("시간강사명·대체교사·기간을 확인하세요.")
                elif not any(avail.values()):
                    st.error("가능한 요일·교시를 하나 이상 선택하세요.")
                else:
                    row={c:"" for c in ensure_part_time_columns(pd.DataFrame()).columns}
                    row["시간강사명"]=pt_name.strip(); row["담당과목"]=pt_subject; row["대체교사"]=pt_orig; row["시작일"]=qstart.strftime("%Y-%m-%d"); row["종료일"]=qend.strftime("%Y-%m-%d")
                    for (dd,pp),flag in avail.items():
                        if flag: row[f"{dd}{pp}"]=1
                    candidate=ensure_part_time_columns(pd.concat([st.session_state.part_time,pd.DataFrame([row])],ignore_index=True))
                    ok,msg=validate_part_time_table(candidate)
                    if not ok: st.error(msg)
                    else:
                        before_pt = st.session_state.part_time.copy(deep=True)
                        st.session_state.part_time=candidate
                        if save_work_data_to_gsheet(["시간강사"]):
                            push_history("시간강사 추가"); st.success("추가 완료"); st.rerun()
                        st.session_state.part_time = before_pt
                        _invalidate_all_caches()
                        st.error("저장에 실패하여 이전 상태로 복구했습니다.")
            st.markdown("#### 🛠️ 기존 자료 고급 편집")
            if can_full_data():
                ed = st.data_editor(
                    st.session_state.part_time,
                    num_rows="dynamic",
                    width="stretch",
                    height=450,
                    key="pt_ed",
                    hide_index=True,
                    column_config={
                        "시작일": st.column_config.TextColumn("시작일 (YYYY-MM-DD)"),
                        "종료일": st.column_config.TextColumn("종료일 (YYYY-MM-DD)"),
                        "대체교사": st.column_config.TextColumn("대체할 정규교사명"),
                    }
                )
                if st.button("시간강사 정보 저장", type="primary"):
                    candidate_pt = ensure_part_time_columns(ed)
                    ok, msg = validate_part_time_table(candidate_pt)
                    if not ok:
                        st.error(msg)
                    else:
                        before_pt = st.session_state.part_time.copy(deep=True)
                        st.session_state.part_time = candidate_pt
                        if save_work_data_to_gsheet(["시간강사"]):
                            push_history("시간강사 수정")
                            st.success("저장 완료")
                            st.rerun()
                        st.session_state.part_time = before_pt
                        _invalidate_all_caches()
                        st.error("저장에 실패하여 이전 상태로 복구했습니다.")
            else:
                st.dataframe(filter_by_owner(st.session_state.part_time), width="stretch", hide_index=True)
        else:
            st.dataframe(st.session_state.part_time, width="stretch")
if "결강·보강" in tab_map:
    with tab_map["결강·보강"]:
        st.subheader("결강 등록 & 보강 배정")
        left, right = st.columns(2)
        with left:
            st.markdown("### 📌 결강 등록")
            d_sel, abs_matrix, abs_cells = daily_schedule_picker(_today_kst(), key="abs_daily", multi=True, height=430,
                help_text="달력에서 날짜를 선택하고 결강할 수업 셀을 하나 이상 클릭하세요. 교사와 교시가 자동으로 결정됩니다.")
            on_date=d_sel.strftime("%Y-%m-%d"); day=WEEKDAY_KR[d_sel.weekday()]
            reason=st.selectbox("사유",ABSENCE_REASONS,key="abs_reason")
            detail=st.text_input("상세 사유",key="abs_detail")
            ver=st.session_state.get("_data_version",0)
            selected_abs={}
            for x in abs_cells:
                selected_abs.setdefault(x["교사명"],[]).append(x["교시"])
            if selected_abs:
                st.markdown("#### 선택한 결강 수업")
                st.write(" · ".join(f"{t}: {', '.join(map(str,sorted(ps)))}교시" for t,ps in selected_abs.items()))
            if st.button("결강 등록",type="primary",key="btn_abs"):
                if not selected_abs:
                    st.error("시간표 매트릭스에서 결강 수업을 하나 이상 선택하세요.")
                else:
                    a=st.session_state.absences
                    new_rows=[]
                    for who,sel_p in selected_abs.items():
                        cid=f"{on_date}-{who}"
                        if not a.empty:
                            a=a[~((a["일자"]==on_date)&(a["교사명"]==who))]
                        e_tt=get_effective_timetable_for_date(on_date,ver)
                        todays=e_tt[(e_tt["교사명"]==who)&(e_tt["요일"]==day)].sort_values("교시")
                        for r in todays[todays["교시"].apply(safe_int).isin(sel_p)].itertuples():
                            new_rows.append({"결강ID":cid,"일자":on_date,"요일":day,"교사명":who,"사유":reason,"상세사유":detail,"교시":safe_int(r.교시),"학급":r.학급,"과목":r.과목,"등록시각":_now_text("%Y-%m-%d %H:%M"),"입력자":current_user()})
                    if new_rows:
                        before_absences = st.session_state.absences.copy(deep=True)
                        st.session_state.absences=pd.concat([a,pd.DataFrame(new_rows)],ignore_index=True)
                        if save_work_data_to_gsheet(["결강"]):
                            push_history("결강 등록")
                            st.success(f"{len(new_rows)}건 결강 등록 완료"); st.rerun()
                        st.session_state.absences = before_absences
                        _invalidate_all_caches()
                        st.error("결강 저장에 실패하여 작업을 원복했습니다.")
            show_abs=filter_by_owner(st.session_state.absences)
            st.dataframe(show_abs[show_abs["일자"]==on_date] if not show_abs.empty else pd.DataFrame(),height=250,hide_index=True)
        with right:
            st.markdown("### 📌 보강 배정")
            ab = filter_by_owner(st.session_state.absences)
            if ab.empty:
                st.info("먼저 결강을 등록하세요." if can_full_data() else "본인이 등록한 결강이 없습니다.")
            else:
                st.markdown("#### 📅 보강 대상 선택 — 달력 → 결강 매트릭스")
                sub_ref=calendar_picker("결강 날짜", _today_kst(), key="sub_ref")
                sub_date=sub_ref.strftime("%Y-%m-%d")
                day_abs=ab[ab["일자"].map(normalize_date_str)==sub_date].copy()
                if day_abs.empty:
                    st.info("선택한 날짜에 등록된 결강이 없습니다. 달력에서 다른 날짜를 선택하세요.")
                    cid=None; rows=pd.DataFrame(); head=None
                else:
                    teachers_for_day=sorted(day_abs["교사명"].dropna().astype(str).str.strip().unique())
                    idx={(str(r["교사명"]).strip(),safe_int(r["교시"])):r for _,r in day_abs.iterrows()}
                    sm_rows=[]
                    for teacher in teachers_for_day:
                        rr={"교사명":teacher}
                        for pp in range(1,MAX_PERIOD+1):
                            r=idx.get((teacher,pp))
                            rr[f"{pp}교시"]=(f"{r.get('학급','')} {r.get('과목','')} ❗".strip() if r is not None else "")
                        sm_rows.append(rr)
                    sub_matrix=pd.DataFrame(sm_rows,columns=["교사명"]+[f"{p}교시" for p in range(1,MAX_PERIOD+1)])
                    sub_event=st.dataframe(sub_matrix,hide_index=True,width="stretch",height=300,key=f"sub_abs_matrix_{sub_date}",on_select="rerun",selection_mode="single-cell")
                    sub_cells=getattr(getattr(sub_event,"selection",None),"cells",[]) or []
                    cid=None; rows=pd.DataFrame(); head=None
                    if sub_cells:
                        ri,cc=sub_cells[0]
                        if 0<=ri<len(sub_matrix) and cc!="교사명":
                            chosen_teacher=str(sub_matrix.iloc[ri]["교사명"]).strip(); chosen_period=safe_int(str(cc).replace("교시",""))
                            hit=day_abs[(day_abs["교사명"].astype(str).str.strip()==chosen_teacher)&(day_abs["교시"].apply(safe_int)==chosen_period)]
                            if not hit.empty:
                                cid=str(hit.iloc[0]["결강ID"]); rows=ab[ab["결강ID"]==cid].sort_values("교시"); head=rows.iloc[0]
                    if head is not None:
                        st.success(f"선택: **{head['일자']} · {head['교사명']} · {len(rows)}시간**")
                if head is not None:
                    include_pt = st.checkbox("시간강사 포함", key="sub_pt")
                    show_all = st.checkbox("전체 공강 교사 보기 (최대 20명)", value=True, key="show_all_free")
                    if st.button("전 교시 자동 배정", type="primary", key="auto_all"):
                        assignments = []
                        for r in rows.itertuples():
                            cand = recommend_substitutes(head["요일"], r.교시, r.과목, r.학급, head["교사명"], head["일자"], top_n=1, include_part_time=include_pt)
                            if not cand.empty:
                                assignments.append({"cid":cid,"on_date":head["일자"],"day":head["요일"],"period":r.교시,"class_name":r.학급,"subject":r.과목,"absent_teacher":head["교사명"],"sub_teacher":cand.iloc[0]["보강교사"],"method":"자동","priority":cand.iloc[0]["우선순위"],"memo":""})
                        accepted, errors = add_substitutes_batch(assignments, "자동 보강")
                        st.success(f"자동 보강 {accepted}건 처리 완료")
                        if errors: st.warning(" / ".join(errors[:5]))
                        st.rerun()
                    for r in rows.itertuples():
                        p = safe_int(r.교시)
                        with st.expander(f"{p}교시 · {r.학급} · {r.과목}", expanded=True):
                            cur = st.session_state.subs
                            assigned = None
                            if not cur.empty:
                                m = cur[(cur["결강ID"] == cid) & (cur["교시"] == p)]
                                if not m.empty:
                                    assigned = m.iloc[0]["보강교사"]
                            if assigned:
                                st.success(f"현재 배정: **{assigned}**")
                                if st.button("배정 취소", key=f"cancel_{cid}_{p}"):
                                    cancel_substitute(cid, p)
                                    st.rerun()
                            else:
                                cand = recommend_substitutes(
                                    head["요일"], p, r.과목, r.학급, head["교사명"], head["일자"],
                                    top_n=20 if show_all else 8,
                                    include_part_time=include_pt
                                )
                                if cand.empty:
                                    st.warning("해당 시간대에 비어있는 교사가 없습니다.")
                                else:
                                    st.dataframe(cand, hide_index=True, height=220)
                                    pick = st.selectbox("보강 교사 선택", cand["보강교사"], key=f"pick_{cid}_{p}")
                                    if st.button("배정", key=f"btn_{cid}_{p}"):
                                        pr = cand.loc[cand["보강교사"] == pick, "우선순위"].iloc[0]
                                        add_substitute(cid, head["일자"], head["요일"], p, r.학급, r.과목, head["교사명"], pick, "수동", pr, "")
                                        st.rerun()
if "시간표 맞교환 & 변경 추천" in tab_map:
    with tab_map["시간표 맞교환 & 변경 추천"]:
        week_anchor = calendar_picker(
            "교환 검색 기준 주", _today_kst(),
            key="exchange_week_anchor",
            help_text="선택한 날짜가 포함된 평일 주간 시간표를 표시합니다. 토·일은 표시하지 않습니다."
        )
        ver = st.session_state.get("_data_version", 0)
        lesson_matrix = effective_teacher_matrix(week_anchor, ver, use_test=False)
        render_standard_weekly_matrix(
            lesson_matrix, week_anchor, row_label="교사명",
            key="exchange_weekly_matrix",
            title="📅 현재 적용 주간 시간표 — 수업을 클릭해서 바로 작업",
            use_test=False
        )
        st.divider()
        st.markdown("#### 📜 실제 맞교환 이력")
        show_swaps = filter_by_owner(st.session_state.swaps)
        if show_swaps.empty:
            st.info("등록된 실제 맞교환 이력이 없습니다.")
        else:
            st.dataframe(show_swaps, width="stretch", hide_index=True)
if "통계" in tab_map:
    with tab_map["통계"]:
        st.subheader("보강 통계")
        period = st.selectbox("빠른 기간", ["직접 선택", "전체", "1학기", "2학기", "이번 달"], key="st_p")
        if period == "직접 선택":
            start, end = calendar_range_picker(date(2026, 3, 1), _today_kst(), key="stats_range", help_text="달력에서 시작일과 종료일을 선택합니다.")
        else:
            start, end = date(2026, 3, 1), _today_kst()
        if period == "1학기":
            start, end = date(2026, 3, 1), date(2026, 7, 31)
        elif period == "2학기":
            start, end = date(2026, 8, 1), date(2027, 2, 28)
        elif period == "이번 달":
            start = _today_kst().replace(day=1)
        cum = cumulative_sub_count(start.strftime("%Y-%m-%d"), end.strftime("%Y-%m-%d"), st.session_state.get("_data_version", 0))
        df = pd.DataFrame({"교사명": list(cum.keys()), "누적보강": list(cum.values())}).sort_values("누적보강", ascending=False)
        st.dataframe(df, width="stretch", hide_index=True)
        if not df.empty:
            st.bar_chart(df.set_index("교사명")["누적보강"])
if "시간표 변경 테스트용" in tab_map:
    with tab_map["시간표 변경 테스트용"]:
        st.markdown('<div class="sandbox-title">🧪 시간표 변경 테스트</div>', unsafe_allow_html=True)
        st.markdown('<div class="apple-note">저장되지 않는 테스트 공간 · 연계공강(순환) 발생 시 <strong>수업계에 확인해주세요.</strong></div>', unsafe_allow_html=True)
        if st.button("🔄 테스트 상태 초기화", type="secondary"):
            st.session_state.test_swaps = pd.DataFrame()
            st.session_state["test_has_cycle"] = False
            get_effective_timetable_for_date.clear()
            effective_teacher_matrix.clear()
            get_single_lesson_1to1_candidates.clear()
            get_single_lesson_linked_cycles.clear()
            st.success("테스트 상태가 초기화되었습니다.")
            st.rerun()
        st.markdown("#### 수업 선택 — **현재 적용 + 테스트 변경 결과**에서 수업 셀 하나를 클릭")
        st.caption("실제 변경과 현재까지의 테스트 변경을 모두 반영합니다. 선택한 현재 상태를 기준으로 다음 1:1 가능 위치를 계산합니다.")
        test_week_anchor = calendar_picker("테스트 검색 기준 주", _today_kst(), key="test_week_anchor", help_text="선택한 날짜가 포함된 주간 시간표를 테스트 기준으로 사용합니다.")
        ver = st.session_state.get("_data_version", 0)
        test_matrix = effective_teacher_matrix(test_week_anchor, ver, use_test=True)
        st.caption("표시 기준: 🔄 교환 변경 이력 · 🟢/ [보강] 보강 처리 이력 · 테스트 변경도 함께 반영")
        render_standard_weekly_matrix(
            test_matrix, test_week_anchor, row_label="교사명",
            key="test_week_preview", title="테스트 적용 주간표", use_test=True
        )
        if not st.session_state.get("test_swaps", pd.DataFrame()).empty:
            st.markdown("#### 현재 테스트 중인 맞교환 목록")
            st.dataframe(st.session_state.test_swaps, width="stretch", hide_index=True)
            col_btn1, col_btn2 = st.columns(2)
            with col_btn1:
                if st.session_state.get("test_has_cycle", False):
                    st.warning("연계 순환 교환은 교육과정부로 문의 바랍니다")
                else:
                    if st.button("현재 테스트 중인 맞교환 목록 결보강 계획서 출력하기", type="primary", key="test_report_btn"):
                        html = build_test_swaps_report_html(st.session_state.test_swaps)
                        st.download_button(
                            "HTML 다운로드 (테스트 결보강 계획서)",
                            html.encode("utf-8"),
                            f"테스트_결보강계획서_{datetime.now().strftime('%Y%m%d_%H%M%S')}.html",
                            "text/html",
                            key="test_html_dl"
                        )
                        st.info("이 HTML을 브라우저에서 열고 Ctrl+P → PDF로 저장하시면 양식과 거의 동일한 PDF가 생성됩니다.")
            with col_btn2:
                if st.button("현재 테스트 중인 맞교환 목록 전체 삭제하기", key="test_clear_btn"):
                    st.session_state.test_swaps = pd.DataFrame()
                    st.session_state["test_has_cycle"] = False
                    get_effective_timetable_for_date.clear()
                    effective_teacher_matrix.clear()
                    get_single_lesson_1to1_candidates.clear()
                    get_single_lesson_linked_cycles.clear()
                    st.success("테스트 중인 맞교환 목록이 전체 삭제되었습니다.")
                    st.rerun()
if "변경된 교사 주간표" in tab_map:
    with tab_map["변경된 교사 주간표"]:
        st.markdown("### 변경된 교사 주간표")
        st.caption("원본 교사별 주간표 형식을 기준으로 전체 교사를 한 번에 표시합니다. 변경 이력이 있는 교사는 이름과 상태에 표시됩니다.")
        ref = calendar_picker("주간 기준일", _today_kst(), key="chg_ref")
        changed = set(str(x).strip() for x in get_changed_teachers_for_week(ref) if str(x).strip())
        teacher_names = []
        teachers_df = st.session_state.get("teachers", pd.DataFrame())
        if isinstance(teachers_df, pd.DataFrame) and not teachers_df.empty and "교사명" in teachers_df.columns:
            teacher_names.extend(teachers_df["교사명"].dropna().astype(str).str.strip().tolist())
        timetable_df = st.session_state.get("timetable", pd.DataFrame())
        if isinstance(timetable_df, pd.DataFrame) and not timetable_df.empty and "교사명" in timetable_df.columns:
            teacher_names.extend(timetable_df["교사명"].dropna().astype(str).str.strip().tolist())
        all_teachers = sorted({x for x in teacher_names if x})
        all_teachers = sorted(all_teachers, key=lambda x: (x not in changed, x))
        st.markdown(
            f'<div class="work-note"><strong>전체 {len(all_teachers)}명</strong> · '
            f'이번 주 변경 이력 {len(changed)}명 · 🔄 변경 교사 우선 표시 · 조회 전용</div>',
            unsafe_allow_html=True,
        )
        if not all_teachers:
            st.info("교사 목록이 없습니다. 시간표/교사 시트의 교사명 데이터를 확인해 주세요.")
        else:
            ver = st.session_state.get("_data_version", 0)
            monday = ref - timedelta(days=ref.weekday())
            week_dates = [monday + timedelta(days=i) for i in range(5)]
            daily_indexes = {}
            for d in week_dates:
                ds = d.strftime("%Y-%m-%d")
                e = get_effective_timetable_for_date(ds, ver, use_test=False)
                if isinstance(e, pd.DataFrame) and not e.empty:
                    daily_indexes[ds] = {(str(r.교사명).strip(), safe_int(r.교시)): r for r in e.itertuples(index=False)}
                else:
                    daily_indexes[ds] = {}
            abs_df = st.session_state.get("absences", pd.DataFrame())
            abs_lookup = set()
            if isinstance(abs_df, pd.DataFrame) and not abs_df.empty and {"일자", "교사명", "교시"}.issubset(abs_df.columns):
                for _, ar in abs_df.iterrows():
                    abs_lookup.add((str(ar.get("일자", "")).strip(), str(ar.get("교사명", "")).strip(), safe_int(ar.get("교시"))))
            for idx, teacher in enumerate(all_teachers, 1):
                is_changed = teacher in changed
                subject = get_teacher_subject(teacher)
                badge = "<span class='changed-teacher-badge'>🔄 변경 있음</span>" if is_changed else ""
                st.markdown(
                    f'<div class="changed-teacher-block">'
                    f'<div class="changed-teacher-head">'
                    f'<div><span class="changed-teacher-index">{idx:02d}</span>'
                    f'<span class="changed-teacher-name">{html_lib.escape(str(teacher))}</span>{badge}'
                    f'<span class="changed-teacher-subject">{html_lib.escape((" · " + subject) if subject else "")}</span></div>'
                    f'<div class="changed-teacher-status">주간 조회 전용</div>'
                    f'</div></div>',
                    unsafe_allow_html=True,
                )
                rows = []
                for period in range(1, MAX_PERIOD + 1):
                    row = {"교사명": teacher, "교시": period}
                    for day_idx, d in enumerate(week_dates):
                        ds = d.strftime("%Y-%m-%d")
                        r = daily_indexes[ds].get((teacher, period))
                        cell = ""
                        if r is not None:
                            class_name = str(getattr(r, "학급", "") or "").strip()
                            subj = str(getattr(r, "과목", "") or "").strip()
                            cell = f"{class_name} {subj}".strip()
                            typ = str(getattr(r, "변경유형", "원본") or "원본").strip()
                            src = str(getattr(r, "변경출처", "") or "").strip()
                            if typ == "교환":
                                cell += f" 🔄 {src or '교환'}"
                            elif typ == "테스트교환":
                                cell += f" 🧪 {src or '테스트교환'}"
                            elif typ == "보강":
                                cell += f" 🟢 {src or '보강'}"
                            elif typ == "시간강사":
                                original = str(getattr(r, "원본교사", "") or "").strip()
                                cell += f" 🟡 {original}→시간강사" if original else " 🟡 시간강사"
                            if (ds, teacher, period) in abs_lookup:
                                cell = f"[결강] {cell}"
                        row[DAYS[day_idx]] = cell
                    rows.append(row)
                teacher_grid = pd.DataFrame(rows, columns=["교사명", "교시"] + DAYS)
                render_standard_weekly_matrix(
                    teacher_grid, ref, row_label="교시", key=f"changed_teacher_week_{idx}",
                    title=None, use_test=False, height=335, open_dialog=False
                )
                if idx < len(all_teachers):
                    st.markdown('<div class="changed-teacher-divider"></div>', unsafe_allow_html=True)
if "📋 복무 관리 & 판단" in tab_map:
    with tab_map["📋 복무 관리 & 판단"]:
        st.subheader("📋 복무 관리 & 판단")
        left, right = st.columns([1, 1.4])
        with left:
            st.markdown("### 복무 등록")
            if can_full_data():
                teacher_options = st.session_state.teachers["교사명"].tolist()
            else:
                my_name = current_name()
                if my_name and my_name in st.session_state.teachers["교사명"].tolist():
                    teacher_options = [my_name]
                else:
                    teacher_options = [my_name] if my_name else []
                    st.warning("등록된 이름이 교사 목록에 없습니다. 관리자에게 문의하세요.")
            t = st.selectbox("교사", teacher_options, key="duty_t")
            d, duty_matrix, duty_cells = daily_schedule_picker(_today_kst(), key="duty_daily", teacher_filter=t, multi=True, height=360,
                help_text="달력에서 날짜를 선택하고 해당 교사의 수업 셀을 클릭하세요. 빈 셀도 복무 시간 선택에 사용할 수 있도록 아래 교시 버튼을 제공합니다.")
            reason = st.selectbox("사유", ABSENCE_REASONS, key="duty_r")
            detail = st.text_input("상세", key="duty_det")
            all_day = st.checkbox("하루 전체", key="duty_all")
            if all_day:
                periods=[]
            else:
                periods=sorted({x["교시"] for x in duty_cells if x["교사명"]==t})
                if periods:
                    st.info("선택된 교시: " + ", ".join(f"{p}교시" for p in periods))
                else:
                    periods=period_matrix_picker("복무 교시", "duty_p_matrix", st.session_state.get("duty_p_matrix_sel", []))
                    st.session_state["duty_p_matrix_sel"]=periods
            if st.button("복무 등록", type="primary", key="duty_reg"):
                if all_day or periods:
                    duties = ensure_duty_columns(st.session_state.duties)
                    duties = duties[~((duties["교사명"] == t) & (duties["일자"] == d.strftime("%Y-%m-%d")))]
                    news = []
                    if all_day:
                        news.append({"교사명": t, "일자": d.strftime("%Y-%m-%d"), "교시": 0,
                                     "사유": reason, "상세사유": detail,
                                     "등록시각": _now_text("%Y-%m-%d %H:%M"),
                                     "입력자": current_user()})
                    else:
                        for p in periods:
                            news.append({"교사명": t, "일자": d.strftime("%Y-%m-%d"), "교시": p,
                                         "사유": reason, "상세사유": detail,
                                         "등록시각": _now_text("%Y-%m-%d %H:%M"),
                                         "입력자": current_user()})
                    before_duties = duties.copy(deep=True)
                    st.session_state.duties = pd.concat([before_duties, pd.DataFrame(news)], ignore_index=True)
                    if save_work_data_to_gsheet(["복무"]):
                        push_history(f"복무 ({t})")
                        st.success("등록 완료")
                        st.rerun()
                    st.session_state.duties = before_duties
                    _invalidate_all_caches()
                    st.error("저장에 실패하여 이전 상태로 복구했습니다.")
            st.markdown("#### 등록된 복무")
            duties = ensure_duty_columns(st.session_state.duties)
            show_duties = filter_by_owner(duties)
            if not show_duties.empty:
                grouped = show_duties.groupby(["교사명", "일자", "사유"]).agg({"교시": list}).reset_index()
                grouped["교시표시"] = grouped["교시"].apply(format_periods)
                show_df = grouped[["교사명", "일자", "교시표시", "사유"]]
                st.dataframe(show_df, height=220, hide_index=True, width="stretch")
                options = [f"{row['교사명']} · {row['일자']} · {row['교시표시']} · {row['사유']}" for _, row in show_df.iterrows()]
                selected_label = st.selectbox("복무 선택 (이름·일자 기준)", options, key="duty_select_name")
                if st.button("이 복무 선택", key="duty_sel"):
                    sel_idx = options.index(selected_label)
                    sel_row = show_df.iloc[sel_idx]
                    st.session_state["_sel_duty"] = {
                        "교사명": sel_row["교사명"], "일자": sel_row["일자"],
                        "교시표시": sel_row["교시표시"], "사유": sel_row["사유"]
                    }
                    st.session_state.pop("_duty_search_cache", None)
                    st.rerun()
                if st.button("선택 행 삭제", key="duty_del"):
                    sel_idx = options.index(selected_label)
                    sel_row = show_df.iloc[sel_idx]
                    before_duties = duties.copy(deep=True)
                    st.session_state.duties = duties[~((duties["교사명"] == sel_row["교사명"]) &
                                                       (duties["일자"] == sel_row["일자"]))].reset_index(drop=True)
                    if save_work_data_to_gsheet(["복무"]):
                        st.session_state.pop("_sel_duty", None)
                        st.session_state.pop("_duty_search_cache", None)
                        st.success("삭제 완료")
                        st.rerun()
                    st.session_state.duties = before_duties
                    _invalidate_all_caches()
                    st.error("삭제 저장에 실패하여 이전 상태로 복구했습니다.")
            else:
                st.info("등록된 복무 없음")
        with right:
            st.markdown("### 선택 교사 처리 + 스마트 교환 검색")
            sel = st.session_state.get("_sel_duty")
            if not sel:
                st.info("왼쪽에서 복무를 선택하세요.")
            else:
                t_name = sel["교사명"]
                d_str = sel["일자"]
                day_kr = WEEKDAY_KR[datetime.strptime(d_str, "%Y-%m-%d").weekday()]
                st.markdown(f"**{t_name}** · {d_str} ({day_kr}) · {sel.get('교시표시', '')}")
                future_days = st.slider("미래 검색 일수", 7, 21, 14, key="future_days")
                if st.button("🔍 동일 학급 최우선 1:1 / 연계 검색 실행", type="primary", key="run_smart_search"):
                    with st.spinner("검색 중..."):
                        base_date = datetime.strptime(d_str, "%Y-%m-%d").date()
                        date_list = [base_date + timedelta(days=i) for i in range(0, future_days+1) if (base_date + timedelta(days=i)).weekday() < 5]
                        ver = st.session_state.get("_data_version", 0)
                        e_cache = {td.strftime("%Y-%m-%d"): get_effective_timetable_for_date(td.strftime("%Y-%m-%d"), ver) for td in date_list}
                        e_today = e_cache.get(d_str, pd.DataFrame())
                        my_lessons = e_today[e_today["교사명"] == t_name] if not e_today.empty else pd.DataFrame()
                        duty_periods = set()
                        if "전체" in sel.get("교시표시", ""):
                            duty_periods = set(range(1, 8))
                        else:
                            raw = st.session_state.duties
                            mask = (raw["교사명"] == t_name) & (raw["일자"] == d_str)
                            duty_periods = set(raw.loc[mask, "교시"].tolist())
                            if 0 in duty_periods:
                                duty_periods = set(range(1, 8))
                        results = {}
                        for _, lesson in my_lessons.iterrows():
                            p = safe_int(lesson["교시"])
                            if p not in duty_periods:
                                continue
                            my_class = lesson["학급"]
                            my_grade = grade_of(my_class)
                            my_group = subject_group(lesson["과목"])
                            my_subject = lesson["과목"]
                            candidates_1to1 = []
                            candidates_linked = []
                            for td in date_list:
                                tds = td.strftime("%Y-%m-%d")
                                tday = WEEKDAY_KR[td.weekday()]
                                e_tt = e_cache[tds]
                                others = e_tt[(e_tt["교시"] == p) & (e_tt["교사명"] != t_name)] if not e_tt.empty else pd.DataFrame()
                                for o in others.itertuples():
                                    if is_free(t_name, tday, p, tds, e_tt) and is_free(o.교사명, day_kr, p, d_str, e_today):
                                        other_class = str(o.학급).strip()
                                        if other_class != str(my_class).strip():
                                            continue
                                        other_grade = grade_of(other_class)
                                        same_class = True
                                        same_grade = (other_grade == my_grade)
                                        score = 200
                                        if subject_group(o.과목) == my_group: score += 40
                                        if tds == d_str: score += 15
                                        candidates_1to1.append({
                                            "type": "1:1", "date": tds, "day": tday, "period": p,
                                            "teacher": o.교사명, "class": str(other_class), "subject": str(o.과목),
                                            "lesson": f"{other_class} {o.과목}",
                                            "same_class": same_class, "same_grade": same_grade, "score": score
                                        })
                                if is_free(t_name, tday, p, tds, e_tt):
                                    for ot in st.session_state.teachers["교사명"].tolist()[:30]:
                                        if ot != t_name and is_free(ot, day_kr, p, d_str, e_today):
                                            candidates_linked.append({
                                                "type": "연계", "date": tds, "day": tday, "period": p,
                                                "teacher": ot, "lesson": "공강", "score": 30
                                            })
                                            break
                            candidates_1to1 = sorted(candidates_1to1, key=lambda x: (-x["same_class"], -x["same_grade"], -x["score"]))[:6]
                            candidates_linked = sorted(candidates_linked, key=lambda x: -x["score"])[:4]
                            results[p] = {
                                "my_class": my_class, "my_subject": my_subject, "my_grade": my_grade,
                                "one_to_one": candidates_1to1, "linked": candidates_linked
                            }
                        st.session_state["_duty_search_cache"] = results
                        st.success("검색 완료!")
                cache = st.session_state.get("_duty_search_cache")
                if cache:
                    for p, data in cache.items():
                        with st.expander(f"{p}교시 · {data['my_class']} {data['my_subject']} (학년 {data['my_grade']})", expanded=True):
                            st.markdown("#### 🏆 1:1 맞교환 (동일 학급 최우선)")
                            if not data["one_to_one"]:
                                st.info("조건에 맞는 1:1 대상 없음")
                            else:
                                for i, c in enumerate(data["one_to_one"]):
                                    mark = "🏆 동일학급"
                                    st.write(f"{mark} | {c['date']} ({c['day']}) {c['period']}교시 - **{c['teacher']}** ({c['lesson']})")
                                    if st.button("이 수업과 1:1 맞교환 실행", key=f"o2o_{p}_{i}"):
                                        a_info = {"교사명": t_name, "일자": d_str, "요일": day_kr, "교시": p,
                                                  "학급": data["my_class"], "과목": data["my_subject"]}
                                        b_info = {"교사명": c["teacher"], "일자": c["date"], "요일": c["day"], "교시": c["period"],
                                                  "학급": c.get("class", ""), "과목": c.get("subject", "")}
                                        if do_swap(a_info, b_info, d_str, c["date"]):
                                            st.success("1:1 맞교환 등록 완료!")
                                        st.session_state.pop("_duty_search_cache", None)
                                        st.rerun()
                            st.markdown("#### 🔗 연계 교환")
                            if not data["linked"]:
                                st.info("연계 교환 대상 없음")
                            else:
                                for i, c in enumerate(data["linked"]):
                                    st.write(f"• {c['date']} ({c['day']}) {c['period']}교시 - **{c['teacher']}**")
                                    if st.button("연계 교환 실행", key=f"lnk_{p}_{i}"):
                                        a_info = {"교사명": t_name, "일자": d_str, "요일": day_kr, "교시": p,
                                                  "학급": data["my_class"], "과목": data["my_subject"]}
                                        do_linked_swap(a_info, c["teacher"], d_str, c["date"], c["day"], c["period"])
                                        st.success("연계 교환 등록 완료!")
                                        st.session_state.pop("_duty_search_cache", None)
                                        st.rerun()
                else:
                    st.info("위에서 **검색 실행** 버튼을 눌러주세요.")
if "🛠️ 다중 출장·전체 조정 추천" in tab_map:
    with tab_map["🛠️ 다중 출장·전체 조정 추천"]:
        st.subheader("🛠️ 다중 출장·전체 조정 추천 (마스터/교육과정부 전용)")
        st.info("여러 교사가 동시에 출장·복무일 때 사용합니다. **시작일(교시) ~ 종료일(교시)** 범위를 지정할 수 있습니다.")
        remaining_budget = get_current_budget()
        st.metric("현재 보강비 잔액", f"{remaining_budget:,.0f}원")
        if remaining_budget > 1000000:
            budget_factor = 0.7
            budget_msg = "예산 충분 → 보강 허용 비중 높음"
        elif remaining_budget > 300000:
            budget_factor = 1.0
            budget_msg = "예산 보통 → 균형 추천"
        else:
            budget_factor = 1.8
            budget_msg = "예산 부족 → 맞교환 우선 추천"
        st.caption(f"추천 모드: {budget_msg} (교환 가중치 ×{budget_factor})")
        with st.expander("예산 변경 이력 보기"):
            budget_df = load_budget_df()
            st.dataframe(budget_df.tail(20), width="stretch", hide_index=True)
        st.markdown("### 1. 불가능한 교사·기간 선택 (시작일 ~ 종료일)")
        if "multi_absent" not in st.session_state:
            st.session_state.multi_absent = []
        abs_teacher = st.selectbox("교사", st.session_state.teachers["교사명"].tolist(), key="multi_t")
        start_date, end_date, start_periods, end_periods = range_calendar_matrix_picker(
            _today_kst(), _today_kst(), key="multi_range",
        )
        start_all = 0 in start_periods
        end_all = 0 in end_periods
        st.markdown("#### 🗓️ 선택 범위 미리보기")
        st.caption(f"{start_date:%Y-%m-%d} ({WEEKDAY_KR[start_date.weekday()]}) → {end_date:%Y-%m-%d} ({WEEKDAY_KR[end_date.weekday()]}) · 시작 {format_periods(start_periods)} / 종료 {format_periods(end_periods)}")
        if st.button("기간 추가 (시작일~종료일)", type="primary"):
            if start_date > end_date:
                st.error("시작일이 종료일보다 늦을 수 없습니다.")
            else:
                current = start_date
                added_count = 0
                while current <= end_date:
                    if current.weekday() < 5:
                        d_str = current.strftime("%Y-%m-%d")
                        if current == start_date:
                            periods_to_add = start_periods if start_periods else [0]
                        elif current == end_date:
                            periods_to_add = end_periods if end_periods else [0]
                        else:
                            periods_to_add = [0]
                        for p in periods_to_add:
                            item = {"교사": abs_teacher, "일자": d_str, "교시": p}
                            if item not in st.session_state.multi_absent:
                                st.session_state.multi_absent.append(item)
                                added_count += 1
                    current += timedelta(days=1)
                st.success(f"{added_count}건이 추가되었습니다.")
                st.rerun()
        if st.session_state.multi_absent:
            st.write("**현재 선택된 불가능 목록**")
            st.dataframe(pd.DataFrame(st.session_state.multi_absent), width="stretch", hide_index=True)
            if st.button("목록 초기화"):
                st.session_state.multi_absent = []
                st.rerun()
            if st.button("🔍 전체 조정 추천 실행", type="primary"):
                with st.spinner("다중 조건으로 추천 검색 중..."):
                    all_recs = []
                    for item in st.session_state.multi_absent:
                        t_name = item["교사"]
                        d_str = item["일자"]
                        p = item["교시"]
                        day_kr = WEEKDAY_KR[datetime.strptime(d_str, "%Y-%m-%d").weekday()]
                        e_tt = get_effective_timetable_for_date(d_str, st.session_state.get("_data_version", 0))
                        if p == 0:
                            lessons = e_tt[e_tt["교사명"] == t_name]
                        else:
                            lessons = e_tt[(e_tt["교사명"] == t_name) & (e_tt["교시"] == p)]
                        if lessons.empty:
                            continue
                        for _, lesson in lessons.iterrows():
                            actual_p = safe_int(lesson["교시"])
                            base_date = datetime.strptime(d_str, "%Y-%m-%d").date()
                            for i in range(0, 15):
                                td = base_date + timedelta(days=i)
                                if td.weekday() >= 5:
                                    continue
                                tds = td.strftime("%Y-%m-%d")
                                tday = WEEKDAY_KR[td.weekday()]
                                df_swap, cycles, _ = get_target_time_recommendations(
                                    t_name, d_str, actual_p, lesson["학급"], lesson["과목"],
                                    tds, actual_p, budget_factor=budget_factor, version=st.session_state.get("_data_version", 0), use_test=False
                                )
                                for _, row in df_swap.head(3).iterrows():
                                    all_recs.append({
                                        "원본교사": t_name, "원본일자": d_str, "원본교시": actual_p,
                                        "유형": "1:1", "추천교사": row["교사B"],
                                        "추천내용": row["현재 수업"], "점수": row["점수"],
                                        "동일학급": row.get("same_class", False)
                                    })
                                for cyc in cycles[:2]:
                                    all_recs.append({
                                        "원본교사": t_name, "원본일자": d_str, "원본교시": actual_p,
                                        "유형": f"{cyc['length']}인순환", "추천교사": cyc["path_desc"][:40]+"...",
                                        "추천내용": cyc["path_desc"], "점수": cyc["score"],
                                        "동일학급": True
                                    })
                    if all_recs:
                        rec_df = pd.DataFrame(all_recs).sort_values(["동일학급", "점수"], ascending=[False, False])
                        st.session_state["_multi_recs"] = rec_df
                        st.success(f"총 {len(rec_df)}건의 추천이 생성되었습니다.")
                    else:
                        st.warning("조건에 맞는 추천이 없습니다.")
                        st.session_state["_multi_recs"] = pd.DataFrame()
            if "_multi_recs" in st.session_state and not st.session_state["_multi_recs"].empty:
                st.markdown("### 추천 결과 (동일학급 우선 + 예산 반영)")
                st.dataframe(st.session_state["_multi_recs"], width="stretch", hide_index=True)
                st.caption("실제 적용은 「시간표 맞교환 & 변경 추천」 탭이나 「복무 관리」 탭에서 진행하세요.")
if "🔑 아이디·권한 관리" in tab_map:
    with tab_map["🔑 아이디·권한 관리"]:
        st.subheader("🔑 아이디 · 권한 관리")
        st.caption("마스터와 교육과정부만 접근 가능합니다.")
        ids_df = load_id_sheet()
        st.markdown("### 현재 등록 아이디")
        edited = st.data_editor(
            ids_df,
            num_rows="dynamic",
            width="stretch",
            key="id_editor",
            column_config={
                "아이디": st.column_config.TextColumn("아이디", required=True),
                "이름": st.column_config.TextColumn("등록 이름", required=True),
                "권한": st.column_config.SelectboxColumn(
                    "권한",
                    options=[ROLE_MASTER, ROLE_EDU, ROLE_OFFICE, ROLE_TEACHER],
                    required=True
                ),
                "허용탭": st.column_config.TextColumn("허용탭 (쉼표로 구분, 비워두면 기본값)")
            }
        )
        col1, col2, col3 = st.columns(3)
        with col1:
            if st.button("아이디 목록 저장", type="primary"):
                masters = edited[edited["권한"] == ROLE_MASTER]
                if len(masters) == 0:
                    st.error("마스터 권한 아이디가 최소 1명은 있어야 합니다.")
                else:
                    try:
                        save_id_sheet(edited)
                        st.success("저장 완료")
                        _revalidate_authenticated_user(force=True)
                        st.rerun()
                    except Exception as exc:
                        st.error(f"저장할 수 없습니다: {exc}")
        with col2:
            if is_master():
                st.info("마스터 양도: 원하는 아이디의 권한을 '마스터'로 변경 후 저장")
        with col3:
            if st.button("아이디 추가 요청 목록 보기"):
                req_ws = get_worksheet(WORK_SHEET_ID, "아이디추가요청")
                req_df = df_from_worksheet(req_ws)
                st.dataframe(req_df, width="stretch")
        st.divider()
        st.markdown("### 권한 설명")
        st.markdown("""
        | 권한 | 설명 |
        |------|------|
        | **마스터** | 모든 권한 + 아이디 관리/삭제/양도 + 마스터 양도 + 전체 데이터 관리 |
        | **교육과정부** | 아이디 저장/삭제 + 전체 탭/데이터 관리 |
        | **교무계원** | 교무행정 중심 업무 및 교무호봉획정 |
        | **일반교사** | 본인이 입력한 데이터만 조회·저장·삭제 + 본인 이름으로만 복무 등록 |
        | **게스트** | 아이디 추가요청만 가능 |
        """)
if "📑 회원별 탭 권한 관리" in tab_map:
    with tab_map["📑 회원별 탭 권한 관리"]:
        st.subheader("📑 회원별 탭 권한 관리")
        st.caption("각 회원에게 보여줄 탭을 개별적으로 설정할 수 있습니다.")
        ids_df = load_id_sheet()
        if ids_df.empty:
            st.warning("등록된 아이디가 없습니다.")
        else:
            selected_user = st.selectbox(
                "회원 선택",
                options=ids_df["아이디"].tolist(),
                format_func=lambda x: f"{x} ({ids_df.loc[ids_df['아이디']==x, '이름'].values[0] if len(ids_df.loc[ids_df['아이디']==x])>0 else ''})"
            )
            user_row = ids_df[ids_df["아이디"] == selected_user].iloc[0]
            current_allowed = str(user_row.get("허용탭", "")).strip()
            if current_allowed:
                current_list = [t.strip() for t in current_allowed.split(",") if t.strip()]
            else:
                role = str(user_row.get("권한", ROLE_TEACHER))
                current_list = DEFAULT_TABS.get(role, [])
            st.markdown(f"**현재 선택 회원**: `{selected_user}` / 이름: `{user_row.get('이름','')}` / 권한: `{user_row.get('권한','')}`")
            st.markdown("#### 허용할 탭 선택")
            new_allowed = []
            cols = st.columns(2)
            for idx, tab_name in enumerate(ALL_APP_TABS):
                disabled = False
                if tab_name in ("🔑 아이디·권한 관리", "📑 회원별 탭 권한 관리", "🛠️ 다중 출장·전체 조정 추천") and not can_manage_ids():
                    disabled = True
                with cols[idx % 2]:
                    checked = st.checkbox(
                        tab_name,
                        value=(tab_name in current_list),
                        key=f"tab_check_{selected_user}_{idx}",
                        disabled=disabled
                    )
                    if checked:
                        new_allowed.append(tab_name)
            if st.button("이 회원의 탭 권한 저장", type="primary"):
                ids_df.loc[ids_df["아이디"] == selected_user, "허용탭"] = ",".join(new_allowed)
                save_id_sheet(ids_df)
                st.success(f"{selected_user} 님의 탭 권한이 저장되었습니다.")
                if selected_user == current_user():
                    st.session_state.user_allowed_tabs = ",".join(new_allowed)
                st.rerun()
            st.divider()
            st.markdown("#### 빠른 설정")
            c1, c2, c3 = st.columns(3)
            with c1:
                if st.button("기본 탭으로 초기화"):
                    role = str(user_row.get("권한", ROLE_TEACHER))
                    default = DEFAULT_TABS.get(role, [])
                    ids_df.loc[ids_df["아이디"] == selected_user, "허용탭"] = ",".join(default)
                    save_id_sheet(ids_df)
                    st.success("기본값으로 초기화됨")
                    st.rerun()
            with c2:
                if st.button("모든 탭 허용"):
                    ids_df.loc[ids_df["아이디"] == selected_user, "허용탭"] = ",".join(ALL_APP_TABS)
                    save_id_sheet(ids_df)
                    st.success("모든 탭 허용됨")
                    st.rerun()
            with c3:
                if st.button("모든 탭 차단"):
                    ids_df.loc[ids_df["아이디"] == selected_user, "허용탭"] = ""
                    save_id_sheet(ids_df)
                    st.success("모든 탭 차단됨")
                    st.rerun()
if "교무호봉획정" in tab_map:
    with tab_map["교무호봉획정"]:
        render_salary_tab()
